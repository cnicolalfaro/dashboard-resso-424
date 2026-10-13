import json
import os
import sqlite3
import tempfile
import unittest
import uuid
from pathlib import Path
from unittest.mock import patch

import pymupdf as fitz
from openpyxl import Workbook
from werkzeug.security import generate_password_hash

from .app import create_app
from .catalog import import_tarja, load_catalog, rut
from .documents import HEIGHT, WIDTH, VERSION, make_document, valid_payload
from .scanning import MAX_PAGES, process_scan, suggestion
from .storage import connect, private_root


class ServiceTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.fixture_directory = tempfile.TemporaryDirectory(
            prefix="nami-unittest-", dir="/tmp"
        )
        cls.addClassCleanup(cls.fixture_directory.cleanup)
        cls.root = Path(cls.fixture_directory.name).resolve()
        cls.tarja = cls.root / "tarja.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "TARJA"
        sheet.append(["Nombre", "RUT", "Cargo", "Especialidad", "Rol", "Estado"])
        sheet.append(["Persona Uno", "12.345.678-5", "Operador", "Mecánica", "Titular", "Activo"])
        sheet.append(["Persona Dos", "9.876.543-3", "Operador", "Eléctrica", "Relevo", "Activo"])
        sheet.append(["Persona Terminada", "11.111.111-1", "Operador", "Otra", "Titular", "Finiquitado"])
        book.save(cls.tarja)
        cls.catalog = {
            "crews": [{"id": str(index), "name": f"Cuadrilla {index}",
                       "ruts": ["12345678-5", "9876543-3"]} for index in range(1, 5)],
            "courses": [{"id": "rf1", "code": "RF-01", "title": "Difusión de seguridad",
                         "content": "Contenido privado de difusión, no evaluación."},
                        {"id": "rf2", "code": "RF-02", "title": "Segundo curso",
                         "content": "Segundo contenido."},
                        {"id": "rf3", "code": "RF-03", "title": "Tercer curso",
                         "content": "Tercer contenido."},
                        {"id": "rf4", "code": "RF-04", "title": "Cuarto curso",
                         "content": "Cuarto contenido."}],
        }
        (cls.root / "catalog.json").write_text(json.dumps(cls.catalog))
        os.chmod(cls.root / "catalog.json", 0o600)
        cls.config = {
            "NAMI_DATA_DIR": str(cls.root), "NAMI_TARJA_PATH": str(cls.tarja),
            "NAMI_SECRET_KEY": "unittest-session-signing-secret-not-for-production",
            "NAMI_USERNAME": "test-operator",
            "NAMI_PASSWORD_HASH": generate_password_hash("test-password-not-for-production"),
            "NAMI_ORIGIN": "https://nami.example",
        }
        cls.app = create_app(cls.config)
        cls.app.config["TESTING"] = True

    def setUp(self):
        self.client = self.app.test_client()
        session = self.client.get("/api/nami/session")
        self.csrf = session.json["csrf"]
        login = self.post("/login", {"username": self.config["NAMI_USERNAME"],
                                    "password": "test-password-not-for-production"})
        # Use a distinct address per test; the persistent login limiter is deliberate.
        self.assertEqual(login.status_code, 200)
        self.csrf = login.json["csrf"]

    def headers(self):
        return {"Origin": self.config["NAMI_ORIGIN"], "X-CSRF-Token": self.csrf}

    def post(self, route, data):
        return self.client.post("/api/nami" + route, json=data, headers=self.headers(),
                                environ_overrides={"REMOTE_ADDR": self.id()})

    def record(self):
        response = self.post("/records", {"crew_id": "1", "date": "2026-10-12",
                                          "courses": ["rf1", "rf2"], "format_accepted": True})
        self.assertEqual(response.status_code, 201, response.json)
        data = self.client.get(response.json["pdf_url"])
        self.assertEqual(data.status_code, 200)
        pdf = data.data
        data.close()
        return response.json, pdf

    def upload(self, data):
        boundary = uuid.uuid4().hex
        encoded = (f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="scan.pdf"\r\n'
                   'Content-Type: application/pdf\r\n\r\n').encode() + data + f"\r\n--{boundary}--\r\n".encode()
        return self.client.post("/api/nami/scans", data=encoded, headers=self.headers(),
                                content_type=f"multipart/form-data; boundary={boundary}")

    def manifest(self, identifier):
        db = connect(self.root)
        try:
            return json.loads(db.execute("SELECT manifest FROM records WHERE id=?", (identifier,)).fetchone()[0])
        finally:
            db.close()

    def test_auth_csrf_cookie_and_logout(self):
        anonymous = self.app.test_client()
        self.assertEqual(anonymous.get("/api/nami/catalog").status_code, 401)
        self.assertEqual(anonymous.get("/api/nami/results").status_code, 401)
        session = anonymous.get("/api/nami/session")
        cookie = session.headers["Set-Cookie"]
        for word in ("Secure", "HttpOnly", "SameSite=Strict"):
            self.assertIn(word, cookie)
        self.assertEqual(anonymous.post("/api/nami/login", json={}).status_code, 403)
        self.assertEqual(self.client.post("/api/nami/logout", headers={
            **self.headers(), "Origin": "https://evil.example"}).status_code, 403)
        self.assertEqual(self.client.post("/api/nami/logout", headers={
            **self.headers(), "X-CSRF-Token": "wrong"}).status_code, 403)
        self.assertEqual(self.post("/logout", {}).status_code, 200)
        self.assertEqual(self.client.get("/api/nami/catalog").status_code, 401)

    def test_login_without_pre_authentication_csrf_matches_frontend_contract(self):
        anonymous = self.app.test_client()
        credentials = {"username": self.config["NAMI_USERNAME"],
                       "password": "test-password-not-for-production"}
        for headers in ({}, {"Origin": "https://evil.example"},
                        {"Origin": self.config["NAMI_ORIGIN"], "Sec-Fetch-Site": "cross-site"}):
            response = anonymous.post("/api/nami/login", json=credentials, headers=headers,
                                      environ_overrides={"REMOTE_ADDR": "frontend-login-test"})
            self.assertEqual(response.status_code, 403)
        response = anonymous.post("/api/nami/login", json=credentials,
                                  headers={"Origin": self.config["NAMI_ORIGIN"]},
                                  environ_overrides={"REMOTE_ADDR": "frontend-login-test"})
        self.assertEqual(response.status_code, 200, response.json)
        self.assertTrue(response.json["authenticated"])
        self.assertTrue(response.json["csrf"])
        self.assertEqual(anonymous.post("/api/nami/logout",
                                       headers={"Origin": self.config["NAMI_ORIGIN"]}).status_code, 403)
        self.assertEqual(anonymous.post("/api/nami/logout",
                                       headers={"Origin": self.config["NAMI_ORIGIN"],
                                                "X-CSRF-Token": response.json["csrf"]}).status_code, 200)

    def test_catalog_and_read_only_import(self):
        before = self.tarja.read_bytes()
        catalog = self.client.get("/api/nami/catalog")
        self.assertEqual(catalog.status_code, 200)
        self.assertEqual([p["rut"] for p in catalog.json["crews"][0]["participants"]],
                         ["12345678-5", "9876543-3"])
        self.assertNotIn("11111111-1", import_tarja(self.tarja))
        self.assertEqual(before, self.tarja.read_bytes())
        self.assertEqual(rut("12.345.678-5"), "12345678-5")
        with self.assertRaises(ValueError):
            rut("12345678-0")
        book = Workbook()
        active = book.active
        active.title = "TARJA"
        active.append(["Nombre", "RUT", "Cargo", "Especialidad", "Rol", "Estado", "Fecha ingreso"])
        active.append(["Persona", "12345678-5", "Cargo", "Especialidad", "Rol", "Activo", "2020-01-01"])
        terminated = book.create_sheet("Finiquitados")
        terminated.append(["RUT", "Fe_Fin"])
        terminated.append(["12345678-5", "2021-01-01"])
        path = self.root / "with-finiquitados.xlsm"
        book.save(path)
        self.assertEqual(import_tarja(path), {})

    def test_record_manifest_pdf_immutable_and_private(self):
        record, pdf = self.record()
        manifest = self.manifest(record["id"])
        self.assertEqual(manifest["version"], VERSION)
        self.assertEqual(manifest["participants"][0]["rut"], "12345678-5")
        self.assertEqual(len(manifest["pages"][0]["rows"]), 2)
        with fitz.open(stream=pdf, filetype="pdf") as document:
            self.assertEqual(len(document), record["pages"])
            text = "".join(page.get_text() for page in document)
            self.assertIn("SGR-R-006.DCH", text)
            self.assertIn("NO equivale", text)
            self.assertIn("Segundo contenido", text)
        payload = manifest["pages"][0]["qr"]
        self.assertIsNotNone(valid_payload(payload, self.config["NAMI_SECRET_KEY"]))
        self.assertIsNone(valid_payload(payload + "0", self.config["NAMI_SECRET_KEY"]))
        db = connect(self.root)
        try:
            with self.assertRaises(sqlite3.IntegrityError):
                db.execute("UPDATE records SET manifest='{}' WHERE id=?", (record["id"],))
            with self.assertRaises(sqlite3.IntegrityError):
                db.execute("DELETE FROM records WHERE id=?", (record["id"],))
        finally:
            db.close()
        self.assertEqual(self.app.test_client().get(record["pdf_url"]).status_code, 401)
        response = self.client.get(record["pdf_url"])
        self.assertFalse(response.headers.get("Access-Control-Allow-Origin"))
        response.close()

    def test_scan_alignment_red_detection_previews_and_confirmation(self):
        record, pdf = self.record()
        manifest = self.manifest(record["id"])
        row = manifest["pages"][0]["rows"][0]
        with fitz.open(stream=pdf, filetype="pdf") as document:
            page = document[0]
            y = row["box"][1] + row["box"][3] / 2
            page.draw_line((row["rut_x"], y), (WIDTH - 48, y), color=(1, 0, 0), width=2)
            marked = document.tobytes()
        scan = self.upload(marked)
        self.assertEqual(scan.status_code, 201, scan.json)
        self.assertTrue(all(page["record_id"] == record["id"] for page in scan.json["pages"]), scan.json)
        rows = [row for page in scan.json["pages"] for row in page["rows"]]
        self.assertEqual([row["suggested"] for row in rows], ["NO_FIRMO", "FIRMO"], scan.json)
        preview = self.client.get(rows[0]["image_url"])
        self.assertEqual(preview.status_code, 200)
        self.assertEqual(preview.mimetype, "image/png")
        preview.close()
        self.assertEqual(self.app.test_client().get(rows[0]["image_url"]).status_code, 401)
        duplicate = self.upload(marked)
        self.assertTrue(duplicate.json["duplicate"])
        self.assertEqual(duplicate.json["id"], scan.json["id"])
        self.assertFalse(any(result["scan_id"] == scan.json["id"]
                             for result in self.client.get("/api/nami/results").json["results"]))
        self.assertEqual(self.post(f'/scans/{scan.json["id"]}/confirm',
                                  {"decisions": [{"row_id": rows[0]["id"], "status": "NO_FIRMO"}]}).status_code, 400)
        self.assertEqual(self.post(f'/scans/{scan.json["id"]}/confirm', {
            "decisions": [{"row_id": rows[0]["id"], "status": "FIRMO"}] * 2}).status_code, 400)
        decisions = [{"row_id": row["id"], "status": row["suggested"]} for row in rows]
        route = f'/scans/{scan.json["id"]}/confirm'
        confirmed = self.post(route, {"decisions": decisions})
        self.assertEqual(confirmed.status_code, 200, confirmed.json)
        self.assertEqual(len(confirmed.json["results"]), 4)
        self.assertFalse(confirmed.json["duplicate"])
        repeated = self.post(route, {"decisions": list(reversed(decisions))})
        self.assertTrue(repeated.json["duplicate"])
        decisions[0]["status"] = "FIRMO"
        self.assertEqual(self.post(route, {"decisions": decisions}).status_code, 409)
        evidence = self.client.get(confirmed.json["results"][0]["evidence_url"])
        self.assertEqual(evidence.data, marked)
        evidence.close()

    def test_unidentified_and_cropped_scans_block_confirmation(self):
        record, pdf = self.record()
        for cropped in (False, True):
            with fitz.open(stream=pdf, filetype="pdf") as document:
                if cropped:
                    document[0].set_cropbox(fitz.Rect(35, 35, WIDTH - 35, HEIGHT - 35))
                else:
                    document[0].draw_rect(fitz.Rect(WIDTH - 145, 35, WIDTH - 25, 155),
                                          color=(1, 1, 1), fill=(1, 1, 1))
                data = document.tobytes()
            scan = self.upload(data)
            self.assertEqual(scan.status_code, 201, scan.json)
            self.assertIsNone(scan.json["pages"][0]["record_id"], scan.json)
            self.assertEqual(self.post(f'/scans/{scan.json["id"]}/confirm', {
                "decisions": [{"row_id": "unknown", "status": "FIRMO"}]}).status_code, 409)

    def test_grayscale_and_ambiguous_red_review(self):
        import numpy as np
        row = {"box": [42, 225, WIDTH - 84, 36], "rut_x": 48, "signature_x": 440}
        gray = np.full((75, 1065, 3), 255, np.uint8)
        self.assertEqual(suggestion(gray, row, False)[0], "REVISAR")
        partial = gray.copy()
        partial[35:38, 20:200] = [255, 0, 0]
        self.assertEqual(suggestion(partial, row, True)[0], "REVISAR")
        record, pdf = self.record()
        with fitz.open(stream=pdf, filetype="pdf") as source, fitz.open() as gray_pdf:
            for page in source:
                pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), colorspace=fitz.csGRAY)
                target = gray_pdf.new_page(width=WIDTH, height=HEIGHT)
                target.insert_image(target.rect, stream=pix.tobytes("png"))
            data = gray_pdf.tobytes()
        scan = self.upload(data)
        self.assertEqual(scan.status_code, 201, scan.json)
        self.assertTrue(all(page["record_id"] == record["id"] for page in scan.json["pages"]), scan.json)
        self.assertTrue(all(row["suggested"] == "REVISAR" for page in scan.json["pages"] for row in page["rows"]))
        decisions = [{"row_id": row["id"], "status": "FIRMO"} for page in scan.json["pages"] for row in page["rows"]]
        self.assertEqual(self.post(f'/scans/{scan.json["id"]}/confirm',
                                   {"decisions": decisions}).status_code, 200)

    def test_invalid_inputs_and_request_bounds(self):
        self.assertEqual(self.post("/records", {"format_accepted": False}).status_code, 400)
        self.assertEqual(self.post("/records", {"crew_id": "1", "courses": ["invalid"],
                                               "date": "2026-10-12", "format_accepted": True}).status_code, 400)
        self.assertEqual(self.upload(b"not-pdf").status_code, 400)
        self.assertEqual(self.upload(b"%PDF-invalid").status_code, 422)
        with patch("nami_service.app.MAX_SCAN_UPLOAD", 1024 * 1024):
            self.assertEqual(self.upload(b"%PDF-" + b"0" * (1024 * 1024)).status_code, 413)
        with fitz.open() as document:
            for _ in range(MAX_PAGES + 1):
                document.new_page()
            self.assertEqual(self.upload(document.tobytes()).status_code, 422)
        with fitz.open() as document:
            document.new_page(width=5000, height=5000)
            self.assertEqual(self.upload(document.tobytes()).status_code, 422)
        oversized = self.client.post("/api/nami/login", data=b" " * (128 * 1024 + 1),
                                     content_type="application/json", headers=self.headers())
        self.assertEqual(oversized.status_code, 413)
        for path in (str(Path(__file__).parent), "relative", str(self.root.parent)):
            with self.assertRaises(ValueError):
                private_root(path)

    def test_login_failure_rate_limit_generic_errors(self):
        anonymous = self.app.test_client()
        csrf = anonymous.get("/api/nami/session").json["csrf"]
        for index in range(11):
            response = anonymous.post("/api/nami/login", json={"username": "unknown",
                                                               "password": "invalid"},
                                      headers={"Origin": self.config["NAMI_ORIGIN"], "X-CSRF-Token": csrf},
                                      environ_overrides={"REMOTE_ADDR": "rate-limit-test"})
            self.assertEqual(response.status_code, 401 if index < 10 else 429)
            self.assertNotIn("password", str(response.json).lower())

    def test_pagination_preserves_manifest_order_and_one_signature(self):
        person = self.app.extensions["nami_catalog"]["crews"][0]["participants"][0]
        people = [{**person, "nombre": f"Persona {index}"} for index in range(30)]
        manifest, pdf = make_document({"id": "4", "name": "Orden explícito", "participants": people},
                                      self.catalog["courses"], "2026-10-12",
                                      self.config["NAMI_SECRET_KEY"])
        rows = [row for page in manifest["pages"] for row in page["rows"]]
        self.assertEqual([row["nombre"] for row in rows], [person["nombre"] for person in people])
        self.assertEqual(len(rows), 30)
        with fitz.open(stream=pdf, filetype="pdf") as document:
            self.assertEqual(len(document), 4)

    def test_one_to_four_courses_and_crews_keep_one_signature_per_person(self):
        for count in range(1, 5):
            response = self.post("/records", {
                "crew_id": str(count), "date": "2026-10-12", "format_accepted": True,
                "courses": [f"rf{index}" for index in range(1, count + 1)],
            })
            self.assertEqual(response.status_code, 201, response.json)
            manifest = self.manifest(response.json["id"])
            self.assertEqual(len(manifest["courses"]), count)
            self.assertEqual(len(manifest["participants"]), 2)
            self.assertEqual(sum(len(page["rows"]) for page in manifest["pages"]), 2)

    def test_five_configured_courses_are_rejected_by_backend(self):
        courses = self.app.extensions["nami_catalog"]["courses"]
        courses.append({"id": "rf5", "code": "RF-05", "title": "Quinto curso", "content": "Contenido."})
        db = connect(self.root)
        try:
            before = db.execute("SELECT COUNT(*) FROM records").fetchone()[0]
            response = self.post("/records", {
                "crew_id": "1", "date": "2026-10-12", "format_accepted": True,
                "courses": ["rf1", "rf2", "rf3", "rf4", "rf5"],
            })
            self.assertEqual(response.status_code, 400, response.json)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM records").fetchone()[0], before)
        finally:
            courses.pop()
            db.close()

    def test_only_selected_roster_numeric_ruts_and_empty_crews_not_selectable(self):
        root = self.root / "selected-roster"
        root.mkdir(mode=0o700)
        workbook = root / "synthetic-realistic.xlsx"
        book = Workbook()
        summary = book.active
        summary.title = "Resumen"
        summary.append(["RUT", "NOMBRE", "CATEGORIA", "ESPECIALIDAD"])
        summary.append(["12345678-5", "Incorrecto Resumen", "Incorrecto", "Incorrecto"])
        summary.append(["44444444-4", "Solo Resumen", "Incorrecto", "Incorrecto"])
        roster = book.create_sheet("TARJA")
        roster.append(["Plantilla sintética privada"])
        roster.append(["RUT", "metadatos, no encabezado de personal"])
        roster.append(["NOMBRE", "A PATERNO", "A MATERNO", "CATEGORIA", "ESPECIALIDAD", "TURNO", "RUT"])
        roster.append(["Persona", "Paterno", "Materno", "Operador", "Mecánica", "Turno A", 123456785.0])
        maestra = book.create_sheet("Maestra")
        maestra.append(["RUT", "ESTADO"])
        maestra.append([123456785.0, "Activo"])
        departures = book.create_sheet("Finiquitados")
        departures.append(["RUT", "Fe_Fin"])
        departures.append(["11111111-1", "2025-01-01"])
        book.save(workbook)
        before = workbook.read_bytes()
        configured = {"courses": self.catalog["courses"], "crews": [
            {"id": "active-crew", "name": "Personal activo", "ruts": ["12345678-5", "11111111-1"]},
            {"id": "inactive-crew", "name": "Personal no vigente", "ruts": ["11111111-1"]},
        ]}
        (root / "catalog.json").write_text(json.dumps(configured))
        data = load_catalog(root, str(workbook))
        self.assertEqual([crew["id"] for crew in data["crews"]], ["active-crew"])
        person = data["crews"][0]["participants"][0]
        self.assertEqual(person["nombre"], "Persona Paterno Materno")
        self.assertEqual(person["cargo"], "Operador")
        self.assertEqual(person["rol"], "")
        self.assertEqual(person["rut"], "12345678-5")
        self.assertEqual(rut(123456785.0), "12345678-5")
        with self.assertRaises(ValueError):
            rut(12345678.5)
        app = create_app({**self.config, "NAMI_DATA_DIR": str(root), "NAMI_TARJA_PATH": str(workbook)})
        client = app.test_client()
        login = client.post("/api/nami/login", json={"username": self.config["NAMI_USERNAME"],
                                                    "password": "test-password-not-for-production"},
                            headers={"Origin": self.config["NAMI_ORIGIN"]})
        self.assertEqual(login.status_code, 200)
        self.assertEqual(client.get("/api/nami/catalog").json["crews"], data["crews"])
        headers = {"Origin": self.config["NAMI_ORIGIN"], "X-CSRF-Token": login.json["csrf"]}
        response = client.post("/api/nami/records", json={"crew_id": "inactive-crew", "date": "2026-10-12",
                                                         "courses": ["rf1"], "format_accepted": True},
                               headers=headers)
        self.assertEqual(response.status_code, 400)
        self.assertEqual(before, workbook.read_bytes())
        configured["crews"][0]["ruts"].append("44444444-4")
        (root / "catalog.json").write_text(json.dumps(configured))
        with self.assertRaises(ValueError):
            load_catalog(root, str(workbook))
        configured["crews"][0]["ruts"].pop()
        (root / "catalog.json").write_text(json.dumps(configured))
        roster.title = "Personal base"
        book.save(workbook)
        with self.assertRaises(ValueError):
            load_catalog(root, str(workbook))
        self.assertEqual(load_catalog(root, str(workbook), roster_sheet="Personal base"), data)

    def test_dated_departure_beats_undated_active_but_later_dated_rehire_wins(self):
        root = self.root / "effective-departure"
        root.mkdir(mode=0o700)
        workbook = root / "synthetic-effective-dates.xlsx"
        book = Workbook()
        roster = book.active
        roster.title = "TARJA"
        roster.append(["NOMBRE", "RUT", "CATEGORIA", "ESPECIALIDAD", "ESTADO", "Fecha ingreso"])
        roster.append(["Persona Finiquitada", "12345678-5", "Operador", "Mecánica", "Activo", None])
        roster.append(["Persona Recontratada", "9876543-3", "Operador", "Mecánica", "Activo", "2026-08-01"])
        maestra = book.create_sheet("Maestra")
        maestra.append(["RUT", "ESTADO", "Fecha ingreso"])
        maestra.append(["12345678-5", "Activo", None])
        maestra.append(["9876543-3", "Activo", "2026-08-01"])
        departures = book.create_sheet("Finiquitados")
        departures.append(["RUT", "Fe_Fin"])
        departures.append(["12345678-5", "2026-06-01"])
        departures.append(["9876543-3", "2026-06-01"])
        book.save(workbook)
        before = workbook.read_bytes()
        people, known = import_tarja(workbook, include_known=True)
        self.assertEqual(set(people), {"9876543-3"})
        self.assertEqual(known, {"12345678-5", "9876543-3"})
        configured = {"courses": self.catalog["courses"], "crews": [
            {"id": "mixed-crew", "name": "Cuadrilla privada", "ruts": ["12345678-5", "9876543-3"]},
            {"id": "departed-crew", "name": "No vigente", "ruts": ["12345678-5"]},
        ]}
        (root / "catalog.json").write_text(json.dumps(configured))
        app = create_app({**self.config, "NAMI_DATA_DIR": str(root), "NAMI_TARJA_PATH": str(workbook)})
        client = app.test_client()
        login = client.post("/api/nami/login", json={"username": self.config["NAMI_USERNAME"],
                                                    "password": "test-password-not-for-production"},
                            headers={"Origin": self.config["NAMI_ORIGIN"]})
        self.assertEqual(login.status_code, 200)
        crews = client.get("/api/nami/catalog").json["crews"]
        self.assertEqual([crew["id"] for crew in crews], ["mixed-crew"])
        self.assertEqual([person["rut"] for person in crews[0]["participants"]], ["9876543-3"])
        record = client.post("/api/nami/records", json={"crew_id": "mixed-crew", "date": "2026-10-12",
                                                       "courses": ["rf1"], "format_accepted": True},
                             headers={"Origin": self.config["NAMI_ORIGIN"],
                                      "X-CSRF-Token": login.json["csrf"]})
        self.assertEqual(record.status_code, 201, record.json)
        self.assertEqual([person["rut"] for person in record.json["participants"]], ["9876543-3"])
        response = client.get(record.json["pdf_url"])
        with fitz.open(stream=response.data, filetype="pdf") as document:
            text = " ".join(page.get_text() for page in document)
            self.assertIn("Persona Recontratada", text)
            self.assertNotIn("Persona Finiquitada", text)
        response.close()
        self.assertEqual(before, workbook.read_bytes())

    def test_actual_tarja_aliases_rehire_exclusion_arbitrary_crew_and_full_long_roster(self):
        root = self.root / "actual-format"
        root.mkdir(mode=0o700)
        workbook = root / "actual-tarja.xlsx"
        book = Workbook()
        sheet = book.active
        sheet.title = "TARJA"
        sheet.append(["Private workbook"])
        sheet.append(["Current roster"])
        sheet.append(["NOMBRE", "A PATERNO", "A MATERNO", "CATEGORIA", "ESPECIALIDAD",
                      "TURNO", "RUT", "FE_ING", "ROL"])
        first_name = "María Alejandra " + "NombreExtraordinariamenteLargo " * 8
        cargo = "Operadora de equipos de elevación y transporte " * 5
        specialty = "Especialista en mantenimiento electromecánico " * 5
        role = "Coordinación de actividades y supervisión operacional " * 6
        sheet.append([first_name, "ApellidoPaternoCompleto", "ApellidoMaternoCompleto",
                      cargo, specialty, "Turno A", "12345678-5", "2026-01-01", role])
        sheet.append(["Persona", "Recontratada", "Vigente", "Operador", "Mecánica",
                      "Turno B", "9876543-3", "2026-01-01"])
        sheet.append(["Persona", "Baja", "Actual", "Operador", "Mecánica",
                      "Turno B", "11111111-1", "2023-01-01"])
        sheet.append(["Persona", "Baja", "Futura", "Operador", "Mecánica",
                      "Turno B", "22222222-2", "2026-01-01"])
        maestra = book.create_sheet("Maestra")
        maestra.append(["RUT", "ESTADO", "FE_ING"])
        maestra.append(["9876543-3", "Activo", "2026-01-01"])
        terminated = book.create_sheet("Finiquitados")
        terminated.append(["RUT", "Fe_Fin"])
        terminated.append(["9876543-3", "2025-01-01"])
        terminated.append(["11111111-1", "2026-06-01"])
        terminated.append(["22222222-2", "2027-01-01"])
        terminated.append(["33333333-3", "2026-01-01"])
        book.save(workbook)
        before = workbook.read_bytes()
        configured = {"courses": self.catalog["courses"], "crews": [{
            "id": "crew-A_12", "name": "Cuadrilla con identificación privada",
            "ruts": ["12345678-5", "11111111-1", "9876543-3", "33333333-3", "22222222-2"],
        }]}
        (root / "catalog.json").write_text(json.dumps(configured))
        app = create_app({**self.config, "NAMI_DATA_DIR": str(root), "NAMI_TARJA_PATH": str(workbook)})
        client = app.test_client()
        login = client.post("/api/nami/login", json={"username": self.config["NAMI_USERNAME"],
                                                    "password": "test-password-not-for-production"},
                            headers={"Origin": self.config["NAMI_ORIGIN"]})
        self.assertEqual(login.status_code, 200)
        catalog = client.get("/api/nami/catalog").json
        people = catalog["crews"][0]["participants"]
        self.assertEqual([person["rut"] for person in people], ["12345678-5", "9876543-3", "22222222-2"])
        self.assertEqual(people[0]["nombre"], " ".join((first_name.strip(),
                                                     "ApellidoPaternoCompleto", "ApellidoMaternoCompleto")))
        self.assertEqual(people[0]["cargo"], cargo.strip())
        self.assertEqual(people[1]["rol"], "")
        headers = {"Origin": self.config["NAMI_ORIGIN"], "X-CSRF-Token": login.json["csrf"]}
        record = client.post("/api/nami/records", json={
            "crew_id": "crew-A_12", "date": "2026-10-12", "courses": ["rf1", "rf2", "rf3", "rf4"],
            "format_accepted": True}, headers=headers)
        self.assertEqual(record.status_code, 201, record.json)
        response = client.get(record.json["pdf_url"])
        pdf = response.data
        response.close()
        with fitz.open(stream=pdf, filetype="pdf") as document:
            text = " ".join(" ".join(page.get_text() for page in document).split())
            for value in (people[0]["nombre"], cargo, specialty, role):
                self.assertIn(" ".join(value.split()), text)
        db = connect(root)
        try:
            manifest = json.loads(db.execute("SELECT manifest FROM records WHERE id=?",
                                            (record.json["id"],)).fetchone()[0])
        finally:
            db.close()
        rows = [row for page in manifest["pages"] for row in page["rows"]]
        self.assertEqual(len(rows), 3)
        self.assertGreater(rows[0]["box"][3], 36)
        long_people = [{**people[0], "rut": identifier, "nombre": people[0]["nombre"] + f" Persona{index}"}
                       for index, identifier in enumerate(("12345678-5", "9876543-3", "11111111-1",
                                                           "22222222-2", "33333333-3", "44444444-4"))]
        paginated, _ = make_document({"id": "crew-A_12", "name": "Cuadrilla privada",
                                     "participants": long_people},
                                    self.catalog["courses"], "2026-10-12", self.config["NAMI_SECRET_KEY"])
        signature_pages = [page for page in paginated["pages"] if page["rows"]]
        self.assertGreater(len(signature_pages), 1)
        self.assertEqual([row["nombre"] for page in signature_pages for row in page["rows"]],
                         [person["nombre"] for person in long_people])
        for page in signature_pages:
            for row in page["rows"]:
                self.assertLessEqual(row["box"][1] + row["box"][3], HEIGHT - 100)
        with patch.object(self, "client", client), patch.object(self, "csrf", login.json["csrf"]):
            scan = self.upload(pdf)
            self.assertEqual(scan.status_code, 201, scan.json)
            self.assertTrue(all(page["record_id"] == record.json["id"] for page in scan.json["pages"]), scan.json)
            decisions = [{"row_id": row["id"], "status": "FIRMO"}
                         for page in scan.json["pages"] for row in page["rows"]]
            confirmed = self.post(f'/scans/{scan.json["id"]}/confirm', {"decisions": decisions})
            self.assertEqual(confirmed.status_code, 200, confirmed.json)
            self.assertEqual(len(confirmed.json["results"]), 12)
        self.assertEqual(before, workbook.read_bytes())
        configured["crews"][0]["ruts"].append("44444444-4")
        (root / "catalog.json").write_text(json.dumps(configured))
        with self.assertRaises(ValueError):
            load_catalog(root, str(workbook))

    def test_scanner_scaling_and_perspective_use_registered_row_crops(self):
        import cv2
        import numpy as np
        record, pdf = self.record()
        with fitz.open(stream=pdf, filetype="pdf") as source, fitz.open() as document:
            for page in source:
                pix = page.get_pixmap(matrix=fitz.Matrix(2.5, 2.5), colorspace=fitz.csRGB)
                image = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, 3)
                height, width = image.shape[:2]
                original = np.array([[0, 0], [width - 1, 0], [width - 1, height - 1],
                                     [0, height - 1]], dtype=np.float32)
                anchors = np.array([[70, 55], [width - 85, 35], [width - 45, height - 65],
                                    [50, height - 95]], dtype=np.float32)
                transform = cv2.getPerspectiveTransform(original, anchors)
                warped = cv2.warpPerspective(image, transform, (width, height),
                                              borderValue=(255, 255, 255))
                ok, png = cv2.imencode(".png", cv2.cvtColor(warped, cv2.COLOR_RGB2BGR))
                self.assertTrue(ok)
                target = document.new_page(width=WIDTH * 1.1, height=HEIGHT * 1.1)
                target.insert_image(target.rect, stream=png.tobytes())
            data = document.tobytes()
        scan = self.upload(data)
        self.assertEqual(scan.status_code, 201, scan.json)
        self.assertTrue(all(page["record_id"] == record["id"] for page in scan.json["pages"]), scan.json)
        rows = [row for page in scan.json["pages"] for row in page["rows"]]
        self.assertEqual([row["rut"] for row in rows], ["12345678-5", "9876543-3"])
        self.assertTrue(all(row["suggested"] == "FIRMO" for row in rows), scan.json)
        for row in rows:
            response = self.client.get(row["image_url"])
            crop = cv2.imdecode(np.frombuffer(response.data, np.uint8), cv2.IMREAD_COLOR)
            response.close()
            self.assertEqual(crop.shape[:2], (75, 1065))

    def test_rotated_and_duplicate_pages(self):
        record, pdf = self.record()
        with fitz.open(stream=pdf, filetype="pdf") as document:
            for page in document:
                page.set_rotation(90)
            rotated = document.tobytes()
        scan = self.upload(rotated)
        self.assertEqual(scan.status_code, 201, scan.json)
        self.assertTrue(all(page["record_id"] == record["id"] for page in scan.json["pages"]), scan.json)
        with fitz.open(stream=pdf, filetype="pdf") as source, fitz.open() as document:
            document.insert_pdf(source, from_page=0, to_page=0)
            document.insert_pdf(source, from_page=0, to_page=0)
            duplicated = document.tobytes()
        scan = self.upload(duplicated)
        self.assertEqual(scan.status_code, 201)
        self.assertIsNone(scan.json["pages"][1]["record_id"])

    def test_worker_timeout_is_hard_bounded(self):
        with patch("nami_service.scanning.SCAN_SECONDS", 0.02):
            with self.assertRaises(ValueError):
                process_scan(self.root, str(uuid.uuid4()), self.tarja, {}, self.config["NAMI_SECRET_KEY"])

    def test_role_protection_and_restart_persistence(self):
        alternate = create_app({**self.config, "NAMI_ROLE": "registrar"})
        client = alternate.test_client()
        csrf = client.get("/api/nami/session").json["csrf"]
        login = client.post("/api/nami/login", json={"username": self.config["NAMI_USERNAME"],
                                                   "password": "test-password-not-for-production"},
                            headers={"Origin": self.config["NAMI_ORIGIN"], "X-CSRF-Token": csrf},
                            environ_overrides={"REMOTE_ADDR": "role-test"})
        self.assertEqual(login.status_code, 200)
        forbidden = client.post("/api/nami/scans", headers={
            "Origin": self.config["NAMI_ORIGIN"], "X-CSRF-Token": login.json["csrf"]})
        self.assertEqual(forbidden.status_code, 403)
        record, pdf = self.record()
        response = client.get(record["pdf_url"])
        self.assertEqual(response.data, pdf)
        response.close()

    def test_z_confirmation_transaction_rolls_back_on_failure(self):
        record, pdf = self.record()
        scan = self.upload(pdf)
        rows = [row for page in scan.json["pages"] for row in page["rows"]]
        self.assertEqual(len(rows), 2, scan.json)
        decisions = [{"row_id": row["id"], "status": "FIRMO"} for row in rows]
        db = connect(self.root)
        try:
            db.execute("""CREATE TRIGGER test_abort_result BEFORE INSERT ON results
                          WHEN NEW.course_id = 'rf2'
                          BEGIN SELECT RAISE(ABORT, 'test rollback'); END""")
            response = self.post(f'/scans/{scan.json["id"]}/confirm', {"decisions": decisions})
            self.assertEqual(response.status_code, 500)
            self.assertEqual(db.execute("SELECT COUNT(*) FROM results WHERE scan_id=?",
                                        (scan.json["id"],)).fetchone()[0], 0)
            self.assertEqual(db.execute("SELECT confirmed FROM scans WHERE id=?",
                                        (scan.json["id"],)).fetchone()[0], 0)
        finally:
            db.execute("DROP TRIGGER test_abort_result")
            db.close()
        response = self.post(f'/scans/{scan.json["id"]}/confirm', {"decisions": decisions})
        self.assertEqual(response.status_code, 200, response.json)


if __name__ == "__main__":
    unittest.main()
