import datetime
import re
import shutil
import uuid
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.shared import Pt

from .templates import as_templates, course_slug
from .word_register import RenderError, convert_to_pdf


PLACEHOLDER = "NOMBRETRABAJADOR"


def evaluation_template(templates, course):
    return as_templates(templates).evaluation(course)


def available_courses(templates, courses):
    return as_templates(templates).available(courses)


def stamp_name(document, name, foreman=""):
    found = False
    for section in document.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            for paragraph in footer.paragraphs:
                text = "".join(run.text for run in paragraph.runs)
                if PLACEHOLDER not in text:
                    continue
                for index, run in enumerate(paragraph.runs):
                    run.text = f"[{name}]" if index == 0 else ""
                if foreman:
                    tail = f"        Capataz: {foreman}"
                    if len(paragraph.runs) > 1:
                        paragraph.runs[1].text = tail
                    else:
                        paragraph.add_run(tail)
                found = True
    return found


def stamp_date(document, text):
    """Write the date in the cell right after the 'Fecha' label of the header table."""
    for table in document.tables:
        for row in table.rows:
            cells, seen = [], set()
            for cell in row.cells:
                if id(cell._tc) not in seen:
                    seen.add(id(cell._tc))
                    cells.append(cell)
            for index, cell in enumerate(cells[:-1]):
                if " ".join(cell.text.split()).lower() == "fecha":
                    target = cells[index + 1]
                    paragraph = target.paragraphs[0]
                    for run in paragraph.runs:
                        run.text = ""
                    run = paragraph.runs[0] if paragraph.runs else paragraph.add_run()
                    run.text = text
                    run.font.size = Pt(9)
                    return True
    return False


def make_evaluations(templates, manifest, workdir):
    """One PDF per crew: for each person, one sheet per selected course they still lack."""
    templates = as_templates(templates)
    work = Path(workdir) / uuid.uuid4().hex
    work.mkdir(mode=0o700)
    foreman = manifest.get("crew", {}).get("foreman", "")
    date_text = datetime.date.fromisoformat(manifest["date"]).strftime("%d-%m-%Y") if manifest.get("date") else ""
    try:
        jobs, missing = [], []
        for course in manifest["courses"]:
            if templates.evaluation(course) is None:
                missing.append(course["code"])
        for person in manifest["participants"]:
            for course in manifest["courses"]:
                template = templates.evaluation(course)
                if template is None or course["id"] not in person.get("pending", [course["id"]]):
                    continue
                source = work / f"e{len(jobs):04d}.docx"
                document = Document(str(template))
                if not stamp_name(document, person["nombre"], person.get("foreman") or foreman):
                    raise RenderError(f"La evaluación {course['code']} no tiene [NOMBRETRABAJADOR] en el pie")
                if date_text:
                    stamp_date(document, date_text)
                document.save(str(source))
                jobs.append((person, course, source))
        if not jobs:
            raise RenderError("No hay evaluaciones disponibles para los cursos y personas seleccionados")
        results = convert_to_pdf([job[2] for job in jobs], work, work)
        output = fitz.open()
        pages = 0
        for (person, course, _), result in zip(jobs, results):
            with fitz.open(result) as sheet:
                output.insert_pdf(sheet)
                pages += len(sheet)
        output.set_metadata({"title": "Evaluaciones de la cuadrilla"})
        data = output.tobytes(garbage=3, deflate=True)
        output.close()
        return data, {"sheets": len(jobs), "pages": pages, "missing_templates": missing}
    finally:
        shutil.rmtree(work, ignore_errors=True)


def has_placeholder(document):
    """True when a footer carries [NOMBRETRABAJADOR], where each worker's name is stamped."""
    for section in document.sections:
        for footer in (section.footer, section.first_page_footer, section.even_page_footer):
            for paragraph in footer.paragraphs:
                if PLACEHOLDER in "".join(run.text for run in paragraph.runs):
                    return True
    return False