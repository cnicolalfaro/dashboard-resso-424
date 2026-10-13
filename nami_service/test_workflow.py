import datetime
import io
import json
import os
import re
import shutil
import tempfile
import time
import unittest
import uuid
import zipfile
from pathlib import Path

import pymupdf as fitz
from docx import Document
from docx.shared import Pt
from openpyxl import Workbook, load_workbook
from werkzeug.security import generate_password_hash

from .app import create_app
from .catalog import import_tarja, rut
from .documents import valid_evaluation_payload, valid_payload
from .evaluation_scan import find_person, normalize, tesseract_path
from .planillas import (crew_shift, find_foreman, is_night, matrix_courses, plan_groups, read_crews, read_matrix,
                        read_shifts, select_pending, tokens)
from .prepare import CARGO_PREFIX, cargo_key, compose_unit, list_cargos, resolve_courses, safe_name
from .storage import connect

SECRET = "unittest-session-signing-secret-not-for-production"
PASSWORD = "test-password-not-for-production"
ORIGIN = "https://nami.example"
LIBREOFFICE = shutil.which("soffice") is not None


def make_rut(body, dotted=False):
    total = sum(int(char) * (2 + index % 6) for index, char in enumerate(reversed(str(body))))
    expected = 11 - total % 11
    check = "K" if expected == 10 else "0" if expected == 11 else str(expected)
    return f"{body // 1000000}.{body // 1000 % 1000:03d}.{body % 1000:03d}-{check}" if dotted else f"{body}-{check}"


PEOPLE = {
    "foreman_day": (10000001, "CAPATAZ", "DIA", "UNO", "10X10 B"),
    "foreman_night": (10000002, "CAPATAZ", "NOCHE", "DOS", "10X10 B-N"),
    "a1": (10000003, "ANA", "PRIMERA", "SINTETICA", "10X10 B"),
    "a2": (10000004, "BEA", "SEGUNDA", "SINTETICA", "10X10 B"),
    "a3": (10000005, "CARLA", "TERCERA", "SINTETICA", "10X10 B"),
    "b1": (10000006, "DANA", "CUARTA", "SINTETICA", "10X10 B-N"),
    "b2": (10000007, "ELSA", "QUINTA", "SINTETICA", "10X10 B-N"),
    "foreman_day3": (10000008, "CAPATAZ", "DIA", "TRES", "10X10 B"),
    "c1": (10000009, "FANY", "SEXTA", "SINTETICA", "10X10 B"),
}
RUTS = {key: make_rut(value[0]) for key, value in PEOPLE.items()}


def workbook_bytes(build):
    book = Workbook()
    build(book)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


CARGOS = {"c1": "."}
SPECIALTIES = {"c1": "CARPINTERO"}


def cargo_of(key):
    return CARGOS.get(key) or ("CAPATAZ" if key.startswith("foreman") else "MAESTRO PRIMERA")


def specialty_of(key):
    return SPECIALTIES.get(key, "MONTAJE")


def build_tarja(book):
    sheet = book.active
    sheet.title = "TARJA"
    sheet.append([])
    sheet.append(["resumen"])
    sheet.append(["ROL", "RUT", "NOMBRE", "A PATERNO", "A MATERNO", "ESPECIALIDAD", "CATEGORIA", "TURNO"])
    for index, (key, (body, first, paternal, maternal, shift)) in enumerate(PEOPLE.items(), 1):
        sheet.append([f"90000{index}", make_rut(body, dotted=True), first, paternal, maternal,
                      specialty_of(key), cargo_of(key), shift])


def build_crews(book):
    sheet = book.active
    sheet.title = "Cuadrillas"
    sheet.append(["RUT", "NOMBRE", "CARGO", "ESPECIALIDAD", "CUADRILLA"])
    for key in ("a1", "a2", "a3"):
        sheet.append([RUTS[key], key, "MAESTRO", "MONTAJE", "CUADRILLA 1 - CAPATAZ DIA UNO"])
    sheet.append([make_rut(19999999), "NO EXISTE", "MAESTRO", "MONTAJE", "CUADRILLA 1 - CAPATAZ DIA UNO"])
    for key in ("b1", "b2"):
        sheet.append([RUTS[key], key, "MAESTRO", "MONTAJE", "CUADRILLA 2 - CAPATAZ NOCHE DOS"])
    sheet.append([RUTS["c1"], "c1", "MAESTRO", "MONTAJE", "CUADRILLA 3 - CAPATAZ DIA TRES"])


def build_matrix(book):
    sheet = book.active
    sheet.title = "Matriz"
    for _ in range(3):
        sheet.append([])
    sheet.append(["Trabajador", "RUT", "Cargo", "RF-034", "RF-020", "RF-033"])
    rows = {"a1": ("Aprobado", "Pendiente"), "a2": ("Pendiente", "Pendiente"), "a3": ("Aprobado", "Aprobado"),
            "foreman_day": ("Pendiente", "Aprobado"), "foreman_night": ("Aprobado", "Aprobado"),
            "b2": ("aprobado", "APROBADO"), "c1": ("Pendiente", "Pendiente")}
    for key, (first, second) in rows.items():
        sheet.append([key, RUTS[key], "MAESTRO", first, second, "Pendiente"])
    legend = book.create_sheet("Leyenda módulos")
    legend.append(["Código", "Módulo", "Curso", "N°"])
    legend.append(["RF-034", "Difusión, entrega y evaluación de RF-01 Energía Eléctrica vs 03", "RF", "034"])
    legend.append(["RF-020", "Difusión, entrega y evaluación de RF-02 Trabajo en altura vs 03", "RF", "020"])
    legend.append(["RF-033", "Difusión, entrega y evaluación de RF-03 Maniobras de Izajes vs 02", "RF", "033"])


def build_planning_crews(book):
    """Daily roster: each CAPATAZ row opens a crew; only 'Trabajando' people count; latest date wins."""
    sheet = book.active
    sheet.title = "Dia 11-10-26"
    sheet.append(["ITEM", "FECHA", "ROL", "RUT", "NOMBRE", "A. PATERNO", "A. MATERNO", "ESPECIALIDAD",
                  "CATEGORIA", "TURNO", "ESTADO PERSONAL"])
    old, day = datetime.datetime(2026, 10, 10), datetime.datetime(2026, 10, 11)
    rows = [(old, "foreman_day", "CAPATAZ", "Trabajando"), (old, "b1", "MAESTRO PRIMERA", "Trabajando"),
            (day, "a3", "MAESTRO PRIMERA", "Trabajando"),
            (day, "foreman_day", "CAPATAZ", "Trabajando"), (day, "a1", "MAESTRO PRIMERA", "Trabajando"),
            (day, "a2", "MAESTRO PRIMERA", "Inducción"),
            (day, "foreman_night", "CAPATAZ", "Trabajando"), (day, "b1", "MAESTRO PRIMERA", "Trabajando"),
            (day, "b2", "MAESTRO PRIMERA", "Trabajando"),
            (day, "foreman_day3", "CAPATAZ", "Trabajando"), (day, "c1", "MAESTRO PRIMERA", "Trabajando"),
            (day, "ghost", "MAESTRO PRIMERA", "Trabajando")]
    for item, (date, key, category, state) in enumerate(rows, 1):
        if key == "ghost":
            sheet.append([item, date, "", make_rut(19999999), "NO", "EXISTE", "", "MONTAJE", category, "10X10 B", state])
            continue
        first, paternal, maternal = PEOPLE[key][1:4]
        sheet.append([item, date, "", RUTS[key], first, paternal, maternal, "MONTAJE", category, PEOPLE[key][4], state])
    decoy = book.create_sheet("Cuadrillas")
    decoy.append(["ITEM", "ROL", "RUT", "NOMBRE", "CATEGORIA", "ESTADO PERSONAL"])
    decoy.append([1, "", RUTS["a1"], "DECOY", "CAPATAZ", "Trabajando"])


def build_matrix_titled(book):
    """Matrix whose headers carry the course name and whose approvals are check marks."""
    sheet = book.active
    sheet.title = "Matriz"
    for _ in range(3):
        sheet.append([])
    titles = ["RF-01 Evaluación RF Energía Eléctrica vs 03", "RF-02 Evaluación RF Trabajo en altura vs 03",
              "RF-03 Evaluación RF Maniobras de Izajes vs 02"]
    sheet.append(["N°", "RUT", "Trabajador", "Cargo", "Estado", *titles])
    marks = {"a1": ("✓", "✗", "✗"), "a2": ("✓*", "✓", "✗"), "a3": ("✓", "✓", "✓"),
             "foreman_day": ("✗", "✗", "✗"), "b2": ("✓", "✓", "✗")}
    for number, (key, values) in enumerate(marks.items(), 1):
        sheet.append([number, RUTS[key], key, "MAESTRO", "Vigente", *values])
    data = book.create_sheet("Datos")
    data.append(["N", "RUT", "NOMBRE", "RF-01", "RF-02", "RF-03"])
    data.append([1, RUTS["a1"], "decoy", 0, 0, 0])
    courses = book.create_sheet("Cursos")
    courses.append(["CODIGO", "FAMILIA", "CURSO", "FUENTES"])
    for code, title in zip(("RF-01", "RF-02", "RF-03"), titles):
        courses.append([code, "Riesgos de Fatalidad", title, "Forms"])
    courses.append(["IRL-024", "Inducción", "IRL-024 Curso que no es RF", "Forms"])
    legend = book.create_sheet("Leyenda")
    legend.append(["Símbolo", "Significado"])
    legend.append(["✓", "Aprobado"])
    legend.append(["✓*", "Certificado encontrado, aún no figura aprobado"])


def build_register_template(path):
    document = Document()
    section = document.sections[0]
    section.left_margin, section.right_margin = Pt(44), Pt(21)
    section.top_margin, section.bottom_margin = Pt(20), Pt(14)
    for text in ("CORPORACION NACIONAL DEL COBRE", "SISTEMA DE GESTION DE RIESGOS", "REGISTRO DE COMUNICACIONES",
                 "CODIGO: SGR-R-006.DCH", "", ""):
        document.add_paragraph(text)
    info = document.add_table(rows=9, cols=14)
    for row in (0, 1, 2, 5, 6, 7, 8):
        info.cell(row, 0).merge(info.cell(row, 3))
        info.cell(row, 4).merge(info.cell(row, 13))
    info.cell(3, 3).merge(info.cell(3, 13))
    info.cell(4, 2).merge(info.cell(4, 13))
    info.cell(0, 0).text = "NOMBRE RELATOR (ES):"
    info.cell(1, 0).text = "FECHA:"
    info.cell(3, 3).text = "OBJETIVOS: texto"
    info.cell(4, 2).text = "MATERIA(S) TRATADA(S):-"
    info.cell(4, 2).add_paragraph("1. materia de ejemplo")
    info.cell(8, 0).text = "PARTICIPANTES:"
    people = document.add_table(rows=16, cols=6)
    people.style = "Table Grid"
    people.autofit = False
    widths = (22, 250, 65, 65, 114, 56)
    for row in people.rows:
        for index, cell in enumerate(row.cells):
            cell.width = Pt(widths[index])
    for index, label in enumerate(("N", "NOMBRE", "RUT", "CARGO", "ESPECIALIDAD", "FIRMA")):
        people.cell(0, index).text = label
    for row in range(1, 16):
        people.cell(row, 0).text = str(row)
        people.cell(row, 1).text = "nombre"
        people.cell(row, 1).add_paragraph("ROL:[    ]")
    document.add_paragraph("")
    document.save(path)


def build_evaluation_template(path, title):
    document = Document()
    document.add_paragraph(f"EVALUACIÓN DE ENTENDIMIENTO {title}")
    table = document.add_table(rows=4, cols=4)
    table.style = "Table Grid"
    for row, labels in enumerate((("Nombre", "", "RUT", ""), ("Cargo", "", "ROL", ""),
                                  ("Área / Especialidad", "", "Fecha", ""), ("FIRMA", "", "Resultado", ""))):
        for column, label in enumerate(labels):
            table.cell(row, column).text = label
    document.add_paragraph("1. ¿Pregunta de ejemplo?")
    footer = document.sections[0].footer.paragraphs[0]
    footer.add_run("[NOMBRETRABAJADOR")
    footer.add_run(" ]")
    document.save(path)


class PlanillasTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.directory = tempfile.TemporaryDirectory(prefix="nami-unittest-")
        cls.addClassCleanup(cls.directory.cleanup)
        root = Path(cls.directory.name)
        cls.tarja, cls.crews_file, cls.matrix_file = root / "t.xlsx", root / "c.xlsx", root / "m.xlsx"
        cls.tarja.write_bytes(workbook_bytes(build_tarja))
        cls.crews_file.write_bytes(workbook_bytes(build_crews))
        cls.matrix_file.write_bytes(workbook_bytes(build_matrix))
        cls.people, cls.known = import_tarja(cls.tarja, include_known=True)
        cls.shifts = read_shifts(cls.tarja)
        cls.crews = {crew["id"]: crew for crew in
                     read_crews(cls.crews_file, cls.people, cls.known, cls.shifts)["crews"]}
        cls.matrix = read_matrix(cls.matrix_file)

    def test_night_is_marked_by_a_trailing_n_in_turno(self):
        for value in ("10X10 A-N", "10X10 B-N", "10X10 N", "10x10 b-n"):
            self.assertTrue(is_night(value), value)
        for value in ("10X10 A", "8X6 B", "4X3", "5X2", "10X10 D", "", None):
            self.assertFalse(is_night(value), value)

    def test_crews_foreman_label_shift_and_exclusions(self):
        day, night = self.crews["1"], self.crews["2"]
        self.assertEqual(day["foreman"], "CAPATAZ DIA UNO")
        self.assertEqual((day["shift"], night["shift"]), ("DIA", "NOCHE"))
        self.assertEqual(day["shift_source"], "capataz")
        self.assertEqual([item["motivo"] for item in day["excluded"]], ["No figura en la TARJA"])
        self.assertEqual({p["rut"] for p in day["participants"]},
                         {RUTS[key] for key in ("a1", "a2", "a3", "foreman_day")})
        self.assertTrue(next(p for p in day["participants"] if p["rut"] == RUTS["foreman_day"])["capataz"])
        names = [p["nombre"] for p in day["participants"]]
        self.assertEqual(names, sorted(names))

    def test_foreman_shift_outranks_members_and_majority_is_the_fallback(self):
        shifts = {"1": "10X10 B", "2": "10X10 B-N", "3": "10X10 B-N", "boss": "10X10 B"}
        members = [{"rut": "1"}, {"rut": "2"}, {"rut": "3"}]
        self.assertEqual(crew_shift(members, shifts, "boss"), ("DIA", True, "capataz"))
        self.assertEqual(crew_shift(members, shifts, None), ("NOCHE", True, "mayoría"))
        self.assertEqual(crew_shift(members, {}, None), (None, False, None))

    def test_foreman_lookup_ignores_word_order_and_rejects_ambiguity(self):
        by_name = {tokens("UNO DIA CAPATAZ"): ["a"], tokens("REPETIDO NOMBRE"): ["x", "y"]}
        self.assertEqual(find_foreman("capataz dia uno", [], by_name), "a")
        self.assertIsNone(find_foreman("REPETIDO NOMBRE", [], by_name))
        self.assertIsNone(find_foreman("", [], by_name))

    def test_pending_rule_only_fully_approved_people_are_left_out(self):
        courses = {course["code"]: course for course in matrix_courses(self.matrix["legend"])}
        self.assertEqual(courses["RF-01"]["id"], "RF-034")
        day, night = self.crews["1"], self.crews["2"]
        one = [courses["RF-01"]]
        pending, approved = select_pending(day, one, self.matrix)
        self.assertEqual({p["rut"] for p in pending}, {RUTS["a2"], RUTS["foreman_day"]})
        self.assertEqual(approved, 2)
        both = [courses["RF-01"], courses["RF-02"]]
        pending, approved = select_pending(day, both, self.matrix)
        self.assertEqual({p["rut"] for p in pending}, {RUTS["a1"], RUTS["a2"], RUTS["foreman_day"]})
        self.assertEqual({p["rut"]: p["pending"] for p in pending}[RUTS["a1"]], ["RF-020"])
        self.assertEqual(approved, 1)
        pending, approved = select_pending(night, one, self.matrix)
        self.assertEqual([(p["rut"], p["matrix"]) for p in pending], [(RUTS["b1"], "sin_matriz")])
        self.assertEqual(approved, 2)

    def test_matrix_only_counts_the_word_aprobado_in_any_case(self):
        self.assertTrue(self.matrix["workers"][RUTS["b2"]]["RF-034"])
        self.assertFalse(self.matrix["workers"][RUTS["a2"]]["RF-034"])
        self.assertFalse(self.matrix["workers"][RUTS["a3"]]["RF-033"])

    def test_small_crews_are_grouped_by_shift_within_the_limits(self):
        items = [{"id": "1", "shift": "DIA", "pending": 2}, {"id": "2", "shift": "NOCHE", "pending": 1},
                 {"id": "3", "shift": "DIA", "pending": 2}, {"id": "4", "shift": "DIA", "pending": 5},
                 {"id": "5", "shift": "DIA", "pending": 0}, {"id": "6", "shift": None, "pending": 1},
                 {"id": "7", "shift": "DIA", "pending": 4}]
        self.assertEqual(plan_groups(items), [["1", "3", "7"], ["2"], ["4"], ["5"], ["6"]])
        fours = [{"id": str(n), "shift": "DIA", "pending": 4} for n in range(1, 5)]
        self.assertEqual(plan_groups(fours), [["1", "2", "3"], ["4"]])
        ones = [{"id": str(n), "shift": "DIA", "pending": 1} for n in range(1, 8)]
        self.assertEqual(plan_groups(ones), [["1", "2", "3", "4", "5", "6"], ["7"]])

    def test_composed_unit_keeps_each_person_tied_to_their_own_crew_and_foreman(self):
        sources = {"crews": {"crews": list(self.crews.values())}, "matrix": self.matrix}
        courses = resolve_courses(self.matrix, ["RF-01"])
        crew, pending, approved = compose_unit(sources, ["1", "3"], courses)
        self.assertEqual(crew["name"], "CUADRILLAS 1, 3")
        self.assertEqual(crew["foreman"], "CAPATAZ DIA UNO / CAPATAZ DIA TRES")
        self.assertEqual([(p["crew_id"], p["foreman"]) for p in pending],
                         [("1", "CAPATAZ DIA UNO")] * 2 + [("3", "CAPATAZ DIA TRES")] * 2)
        self.assertEqual({p["crew_label"] for p in pending}, {"Cuadrilla 1", "Cuadrilla 3"})
        self.assertEqual(approved, 2)
        single, people, _ = compose_unit(sources, "1", courses)
        self.assertEqual(single["id"], "1")
        self.assertFalse(any("crew_label" in p for p in people))
        with self.assertRaises(ValueError):
            compose_unit(sources, ["1", "2"], courses)

    def test_daily_roster_opens_a_crew_at_each_capataz_and_keeps_only_working_people(self):
        path = Path(self.directory.name) / "planificacion.xlsx"
        path.write_bytes(workbook_bytes(build_planning_crews))
        result = read_crews(path, self.people, self.known, self.shifts)
        crews = {crew["id"]: crew for crew in result["crews"]}
        self.assertEqual(list(crews), ["SIN-CUADRILLA", "1", "2", "3"])
        self.assertEqual(crews["1"]["name"], "CUADRILLA 1 - CAPATAZ DIA UNO")
        self.assertEqual(crews["1"]["foreman"], "CAPATAZ DIA UNO")
        self.assertEqual({p["rut"] for p in crews["1"]["participants"]}, {RUTS["foreman_day"], RUTS["a1"]})
        self.assertEqual([item["motivo"] for item in crews["1"]["excluded"]], ["No está trabajando (Inducción)"])
        self.assertEqual({p["rut"] for p in crews["2"]["participants"]},
                         {RUTS["foreman_night"], RUTS["b1"], RUTS["b2"]})
        self.assertEqual((crews["1"]["shift"], crews["2"]["shift"]), ("DIA", "NOCHE"))
        self.assertEqual([item["motivo"] for item in crews["3"]["excluded"]], ["No figura en la TARJA"])
        self.assertEqual(crews["SIN-CUADRILLA"]["foreman"], "")
        self.assertEqual([p["rut"] for p in crews["SIN-CUADRILLA"]["participants"]], [RUTS["a3"]])
        warnings = result["warnings"]
        self.assertEqual((warnings["hoja"], warnings["fecha_planificacion"]), ("Dia 11-10-26", "2026-10-11"))
        self.assertEqual((warnings["sin_capataz"], warnings["no_trabajando"], warnings["excluidos"]), (1, 1, 1))
        self.assertTrue(all(RUTS["b1"] != p["rut"] for p in crews["1"]["participants"]))

    def test_matrix_with_course_names_in_headers_and_check_marks(self):
        path = Path(self.directory.name) / "matriz-nueva.xlsx"
        path.write_bytes(workbook_bytes(build_matrix_titled))
        matrix = read_matrix(path)
        self.assertEqual(matrix["modules"], ["RF-01", "RF-02", "RF-03"])
        self.assertTrue(matrix["workers"][RUTS["a1"]]["RF-01"])
        self.assertFalse(matrix["workers"][RUTS["a1"]]["RF-02"])
        self.assertFalse(matrix["workers"][RUTS["a2"]]["RF-01"], "✓* is not an approval")
        self.assertTrue(matrix["workers"][RUTS["a2"]]["RF-02"])
        courses = matrix_courses(matrix["legend"])
        self.assertEqual([(c["code"], c["id"]) for c in courses], [("RF-01", "RF-01"), ("RF-02", "RF-02"), ("RF-03", "RF-03")])
        self.assertEqual(courses[0]["title"], "RF-01 Evaluación RF Energía Eléctrica vs 03")
        self.assertNotIn("IRL-024", [c["id"] for c in courses])

    def test_matrix_headers_alone_name_the_courses_without_a_courses_sheet(self):
        def only_matrix(book):
            build_matrix_titled(book)
            del book["Cursos"]
        path = Path(self.directory.name) / "matriz-sin-cursos.xlsx"
        path.write_bytes(workbook_bytes(only_matrix))
        courses = matrix_courses(read_matrix(path)["legend"])
        self.assertEqual(courses[1]["title"], "RF-02 Evaluación RF Trabajo en altura vs 03")

    def test_numeric_approvals_count_only_when_exactly_one(self):
        from .planillas import is_approved
        for value in ("Aprobado", "aprobado", "✓", 1, 1.0, True, "1"):
            self.assertTrue(is_approved(value), value)
        for value in ("✓*", "✗", "Pendiente", 0, 2, False, None, "", "0"):
            self.assertFalse(is_approved(value), value)

    def test_cargo_units_cover_the_whole_tarja_and_name_the_missing_cargo(self):
        sources = {"people": self.people, "crews": {"crews": list(self.crews.values())}, "matrix": self.matrix}
        self.assertEqual([(c["key"], c["label"], c["total"]) for c in list_cargos(sources)],
                         [("CAPATAZ", "CAPATAZ", 3), ("CARPINTERO", "CARPINTERO", 1),
                          ("MAESTRO PRIMERA", "MAESTRO PRIMERA", 5)])
        courses = resolve_courses(self.matrix, ["RF-01"])
        unit, pending, approved = compose_unit(sources, CARGO_PREFIX + "CAPATAZ", courses)
        self.assertEqual((unit["name"], unit["foreman"], unit["shift"], unit["cargo"]),
                         ("CARGO CAPATAZ", "", None, "CAPATAZ"))
        self.assertEqual({p["rut"] for p in pending}, {RUTS["foreman_day"], RUTS["foreman_day3"]})
        self.assertEqual(approved, 1)
        names = [p["nombre"] for p in pending]
        self.assertEqual(names, sorted(names))
        self.assertEqual({p["rut"] for p in compose_unit(sources, CARGO_PREFIX + "MAESTRO PRIMERA", courses)[1]},
                         {RUTS["a2"], RUTS["b1"]})
        self.assertEqual([p["rut"] for p in compose_unit(sources, CARGO_PREFIX + "CARPINTERO", courses)[1]],
                         [RUTS["c1"]])
        with self.assertRaises(ValueError):
            compose_unit(sources, CARGO_PREFIX + "NO EXISTE", courses)
        self.assertEqual(cargo_key({"cargo": " maestro   primera ", "especialidad": "X"}), "MAESTRO PRIMERA")
        self.assertEqual(cargo_key({"cargo": ".", "especialidad": " carpintero  "}), "CARPINTERO")
        self.assertEqual(cargo_key({"cargo": "", "especialidad": "."}), "")
        self.assertEqual(cargo_key({"cargo": None}), "")
        self.assertEqual(safe_name("OOCC MOLDAJE/ENFIERRADURA"), "OOCC MOLDAJE-ENFIERRADURA")
    def test_invalid_workbooks_are_rejected_with_clear_messages(self):
        empty = Path(self.directory.name) / "empty.xlsx"
        empty.write_bytes(workbook_bytes(lambda book: book.active.append(["nada"])))
        with self.assertRaises(ValueError):
            read_matrix(empty)
        with self.assertRaises(ValueError):
            read_crews(empty, self.people, self.known)


class EvaluationMatchTests(unittest.TestCase):
    ONE = str(uuid.uuid4())
    TWO = str(uuid.uuid4())

    def manifests(self):
        courses = [{"id": "c1", "code": "RF-01"}, {"id": "c2", "code": "RF-02"}]
        people = [{"nombre": "CARLOS ANDRES PRUEBA SINTETICO", "rut": "1", "pending": ["c1", "c2"]},
                  {"nombre": "ÁNGELA MARÍA SOTO PÉREZ", "rut": "2", "pending": ["c2"]}]
        return {self.ONE: {"courses": courses, "participants": people},
                self.TWO: {"courses": courses, "participants": [dict(people[0])]}}

    def test_text_read_with_noise_and_accents_matches_the_pending_person(self):
        match, _ = find_person(2, normalize("[Angela Maria Soto Perez]"), self.manifests(), [self.ONE])
        self.assertEqual((match["rut"], match["code"], match["record_id"]), ("2", "RF-02", self.ONE))
        match, _ = find_person(1, normalize("CARLOS ANDRES PRUEBA SINTETIC0"), self.manifests(), [self.TWO])
        self.assertEqual(match["record_id"], self.TWO)

    def test_a_course_the_person_was_not_pending_for_is_not_matched(self):
        match, reason = find_person(1, normalize("ANGELA MARIA SOTO PEREZ"), self.manifests(), [self.ONE])
        self.assertIsNone(match)
        self.assertTrue(reason)

    def test_same_person_in_several_records_is_ambiguous_without_a_scanned_register(self):
        match, reason = find_person(1, normalize("CARLOS ANDRES PRUEBA SINTETICO"), self.manifests(), [])
        self.assertIsNone(match)
        self.assertIn("más de una", reason)

    def test_tampered_evaluation_qr_is_rejected(self):
        from .documents import evaluation_payload
        payload = evaluation_payload(self.ONE, 3, SECRET)
        self.assertEqual(valid_evaluation_payload(payload, SECRET), ("evaluation", self.ONE, 3))
        self.assertIsNone(valid_evaluation_payload(payload.replace(":3:", ":4:"), SECRET))
        self.assertIsNone(valid_evaluation_payload(payload, SECRET + "x"))
        self.assertIsNone(valid_payload(payload, SECRET))


class WorkflowTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.base = tempfile.TemporaryDirectory(prefix="nami-unittest-")
        cls.addClassCleanup(cls.base.cleanup)
        cls.files = {"tarja": ("tarja.xlsx", workbook_bytes(build_tarja)),
                     "crews": ("cuadrillas.xlsx", workbook_bytes(build_crews)),
                     "matrix": ("matriz.xlsx", workbook_bytes(build_matrix))}

    def make_app(self, role="operator", prepare_templates=True):
        root = Path(self.base.name) / uuid.uuid4().hex
        root.mkdir(mode=0o700)
        tarja = Path(self.base.name) / f"{uuid.uuid4().hex}.xlsx"
        tarja.write_bytes(self.files["tarja"][1])
        (root / "catalog.json").write_text(json.dumps({
            "crews": [{"id": "1", "name": "Cuadrilla 1", "ruts": [RUTS["a1"]]}],
            "courses": [{"id": "rf1", "code": "RF-01", "title": "Curso", "content": "Contenido"}]}))
        os.chmod(root / "catalog.json", 0o600)
        if prepare_templates:
            (root / "templates").mkdir(mode=0o700)
            build_register_template(root / "templates" / "registro.docx")
            build_evaluation_template(root / "templates" / "eval-rf01.docx", "RF-01")
            build_evaluation_template(root / "templates" / "eval-rf02.docx", "RF-02")
        config = {"NAMI_DATA_DIR": str(root), "NAMI_TARJA_PATH": str(tarja), "NAMI_SECRET_KEY": SECRET,
                  "NAMI_USERNAME": "test-operator", "NAMI_ROLE": role, "NAMI_ORIGIN": ORIGIN,
                  "NAMI_PASSWORD_HASH": generate_password_hash(PASSWORD)}
        app = create_app(config)
        app.config["TESTING"] = True
        return app, root

    def setUp(self):
        self.app, self.root = self.make_app()
        self.client = self.login(self.app)

    def login(self, app):
        client = app.test_client()
        self.csrf = client.get("/api/nami/session").json["csrf"]
        response = client.post("/api/nami/login", json={"username": "test-operator", "password": PASSWORD},
                               headers={"Origin": ORIGIN, "X-CSRF-Token": self.csrf},
                               environ_overrides={"REMOTE_ADDR": self.id()})
        self.assertEqual(response.status_code, 200)
        self.csrf = response.json["csrf"]
        return client

    def headers(self):
        return {"Origin": ORIGIN, "X-CSRF-Token": self.csrf}

    def upload(self, kind, name=None, data=None, client=None):
        default_name, default_data = self.files.get(kind, self.files["tarja"])
        boundary = uuid.uuid4().hex
        body = (f'--{boundary}\r\nContent-Disposition: form-data; name="kind"\r\n\r\n{kind}\r\n'
                f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{name or default_name}"\r\n'
                'Content-Type: application/octet-stream\r\n\r\n').encode() + (
                    default_data if data is None else data) + f"\r\n--{boundary}--\r\n".encode()
        return (client or self.client).post("/api/nami/sources", data=body, headers=self.headers(),
                                            content_type=f"multipart/form-data; boundary={boundary}")

    def upload_all(self):
        for kind in ("tarja", "crews", "matrix"):
            self.assertEqual(self.upload(kind).status_code, 201)

    def post(self, route, data):
        return self.client.post("/api/nami" + route, json=data, headers=self.headers())

    def upload_template(self, kind, data, course=None, name="plantilla.docx", client=None):
        boundary = uuid.uuid4().hex
        fields = f'--{boundary}\r\nContent-Disposition: form-data; name="kind"\r\n\r\n{kind}\r\n'
        if course is not None:
            fields += f'--{boundary}\r\nContent-Disposition: form-data; name="course"\r\n\r\n{course}\r\n'
        body = (fields + f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="{name}"\r\n'
                'Content-Type: application/octet-stream\r\n\r\n').encode() + data + f"\r\n--{boundary}--\r\n".encode()
        return (client or self.client).post("/api/nami/templates", data=body, headers=self.headers(),
                                            content_type=f"multipart/form-data; boundary={boundary}")

    @staticmethod
    def docx_bytes(build):
        stream = io.BytesIO()
        directory = tempfile.mkdtemp(prefix="nami-unittest-")
        try:
            path = Path(directory) / "plantilla.docx"
            build(path)
            return path.read_bytes()
        finally:
            shutil.rmtree(directory, ignore_errors=True)

    def test_evaluation_word_is_uploaded_per_course_and_validated(self):
        valid = self.docx_bytes(lambda path: build_evaluation_template(path, "RF-03"))
        self.assertEqual(self.upload_template("evaluation", valid, "RF-03").status_code, 409)
        self.upload_all()
        courses = {c["code"]: c["evaluation"] for c in self.client.get("/api/nami/plan").json["courses"]}
        self.assertFalse(courses["RF-03"])
        response = self.upload_template("evaluation", valid, "rf-03")
        self.assertEqual((response.status_code, response.json["course"], response.json["replaced"]), (201, "RF-03", False))
        courses = {c["code"]: c["evaluation"] for c in self.client.get("/api/nami/plan").json["courses"]}
        self.assertTrue(courses["RF-03"])
        stored = self.root / "templates" / "eval-rf03.docx"
        self.assertEqual(stored.read_bytes(), valid)
        self.assertEqual(stored.stat().st_mode & 0o777, 0o600)
        self.assertEqual(os.listdir(self.root / "work"), [name for name in os.listdir(self.root / "work") if name == "profile"])

    def test_replacing_an_evaluation_keeps_the_previous_version(self):
        self.upload_all()
        first = self.docx_bytes(lambda path: build_evaluation_template(path, "primera"))
        second = self.docx_bytes(lambda path: build_evaluation_template(path, "segunda"))
        self.assertEqual(self.upload_template("evaluation", first, "RF-01").status_code, 201)
        result = self.upload_template("evaluation", second, "RF-01")
        self.assertTrue(result.json["replaced"])
        self.assertEqual((self.root / "templates" / "eval-rf01.docx").read_bytes(), second)
        archived = list((self.root / "templates" / "anteriores").iterdir())
        self.assertEqual(len(archived), 2)
        self.assertEqual({item.stat().st_mode & 0o777 for item in archived}, {0o600})
        self.assertEqual((self.root / "templates" / "anteriores").stat().st_mode & 0o777, 0o700)

    def test_template_upload_rejects_invalid_files_and_courses(self):
        self.upload_all()
        valid = self.docx_bytes(lambda path: build_evaluation_template(path, "RF-03"))
        no_footer = self.docx_bytes(lambda path: Document().save(path))
        self.assertEqual(self.upload_template("evaluation", valid, "RF-99").status_code, 422)
        self.assertEqual(self.upload_template("evaluation", valid, "").status_code, 422)
        self.assertEqual(self.upload_template("otra", valid, "RF-03").status_code, 400)
        self.assertEqual(self.upload_template("evaluation", valid, "RF-03", name="x.pdf").status_code, 400)
        self.assertEqual(self.upload_template("evaluation", b"no es word", "RF-03").status_code, 400)
        self.assertEqual(self.upload_template("evaluation", b"PK\x03\x04" + b"roto" * 30, "RF-03").status_code, 422)
        rejected = self.upload_template("evaluation", no_footer, "RF-03")
        self.assertEqual(rejected.status_code, 422)
        self.assertIn("NOMBRETRABAJADOR", rejected.json["error"])
        self.assertEqual(self.upload_template("register", no_footer).status_code, 422)
        self.assertFalse((self.root / "templates" / "eval-rf03.docx").exists())
        self.assertEqual(os.listdir(self.root / "work"), [name for name in os.listdir(self.root / "work") if name == "profile"])
        before = (self.root / "templates" / "registro.docx").read_bytes()
        self.assertEqual(self.upload_template("register", no_footer).status_code, 422)
        self.assertEqual((self.root / "templates" / "registro.docx").read_bytes(), before)

    def test_register_word_can_be_replaced_when_it_has_the_expected_structure(self):
        self.upload_all()
        good = self.docx_bytes(build_register_template)
        response = self.upload_template("register", good)
        self.assertEqual((response.status_code, response.json["replaced"]), (201, True))
        self.assertEqual(len(os.listdir(self.root / "templates" / "anteriores")), 1)

    def test_only_roles_that_can_create_may_upload_templates(self):
        app, _ = self.make_app(role="reviewer")
        client = self.login(app)
        valid = self.docx_bytes(lambda path: build_evaluation_template(path, "RF-03"))
        self.assertEqual(self.upload_template("evaluation", valid, "RF-03", client=client).status_code, 403)
        self.assertEqual(self.app.test_client().post("/api/nami/templates", data={}).status_code, 403)

    @unittest.skipUnless(LIBREOFFICE, "LibreOffice requerido")
    def test_uploaded_evaluation_is_used_when_generating(self):
        self.upload_all()
        valid = self.docx_bytes(lambda path: build_evaluation_template(path, "RF-03"))
        self.assertEqual(self.upload_template("evaluation", valid, "RF-03").status_code, 201)
        response = self.post("/records/word", {"crew_id": "1", "date": "2026-10-12", "courses": ["RF-03"]})
        self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(response.json["summary"]["evaluation_sheets"], 4)
        self.assertEqual(response.json["summary"]["courses_without_evaluation"], [])

    def test_plan_with_daily_roster_and_titled_matrix_through_the_page_api(self):
        self.assertEqual(self.upload("tarja").status_code, 201)
        crews = self.upload("crews", name="cuadrillas dia 11_10.xlsx", data=workbook_bytes(build_planning_crews))
        self.assertEqual(crews.status_code, 201, crews.json)
        self.assertEqual(self.upload("matrix", name="Matriz_RESSO.xlsx", data=workbook_bytes(build_matrix_titled)).status_code, 201)
        plan = self.client.get("/api/nami/plan?courses=RF-01").json
        self.assertEqual([(c["code"], c["title"]) for c in plan["courses"]][:2],
                         [("RF-01", "RF-01 Evaluación RF Energía Eléctrica vs 03"),
                          ("RF-02", "RF-02 Evaluación RF Trabajo en altura vs 03")])
        by_id = {crew["id"]: crew for crew in plan["crews"]}
        self.assertEqual(by_id["1"]["foreman"], "CAPATAZ DIA UNO")
        self.assertEqual((by_id["1"]["pending"], by_id["2"]["pending"], by_id["3"]["pending"]), (1, 2, 2))
        self.assertEqual(plan["warnings"]["fecha_planificacion"], "2026-10-11")

    def test_requires_authentication_and_csrf(self):
        anonymous = self.app.test_client()
        self.assertEqual(anonymous.get("/api/nami/plan").status_code, 401)
        self.assertEqual(anonymous.get("/api/nami/sources").status_code, 401)
        self.assertEqual(anonymous.get("/api/nami/results/export.xlsx").status_code, 401)
        self.assertEqual(self.client.post("/api/nami/sources", data={},
                                          headers={"Origin": ORIGIN}).status_code, 403)

    def test_reviewer_role_cannot_upload_or_generate(self):
        app, _ = self.make_app(role="reviewer")
        client = self.login(app)
        self.assertEqual(self.upload("tarja", client=client).status_code, 403)
        self.assertEqual(client.post("/api/nami/records/word", json={}, headers=self.headers()).status_code, 403)
        self.assertEqual(client.get("/api/nami/sources").status_code, 200)

    def test_upload_validation(self):
        self.assertEqual(self.upload("tarja", name="tarja.txt").status_code, 400)
        self.assertEqual(self.upload("matrix", name="m.xlsb").status_code, 400)
        self.assertEqual(self.upload("otra").status_code, 400)
        self.assertEqual(self.upload("tarja", data=b"no es una planilla").status_code, 400)
        broken = self.upload("tarja", data=b"PK\x03\x04" + b"corrupto" * 20)
        self.assertEqual(broken.status_code, 422)
        self.assertEqual(self.upload("crews").status_code, 422)
        self.assertEqual(os.listdir(self.root / "sources"), [])
        empty = workbook_bytes(lambda book: book.active.append(["nada"]))
        self.assertEqual(self.upload("matrix", data=empty).status_code, 422)

    def test_workbooks_larger_than_the_multipart_chunk_can_be_uploaded_through_the_real_server(self):
        import http.client
        import random
        import threading
        import waitress
        random.seed(5)

        def large_matrix(book):
            build_matrix(book)
            filler = book.create_sheet("Relleno")
            for _ in range(6000):
                filler.append([uuid.UUID(int=random.getrandbits(128)).hex for _ in range(4)])

        data = workbook_bytes(large_matrix)
        self.assertGreater(len(data), 200 * 1024)
        server = waitress.create_server(self.app, host="127.0.0.1", port=0, threads=2,
                                        max_request_body_size=20 * 1024 * 1024)
        threading.Thread(target=server.run, daemon=True).start()
        self.addCleanup(server.close)

        def call(method, path, body=None, headers=None):
            connection = http.client.HTTPConnection("127.0.0.1", server.effective_port, timeout=60)
            connection.request(method, path, body=body, headers=headers or {})
            response = connection.getresponse()
            return response.status, response.getheader("Set-Cookie"), response.read()

        _, cookie, content = call("GET", "/api/nami/session")
        cookie = cookie.split(";")[0]
        csrf = json.loads(content)["csrf"]
        status, renewed, content = call("POST", "/api/nami/login", json.dumps(
            {"username": "test-operator", "password": PASSWORD}),
            {"Origin": ORIGIN, "Content-Type": "application/json", "X-CSRF-Token": csrf, "Cookie": cookie})
        self.assertEqual(status, 200)
        cookie, csrf = renewed.split(";")[0], json.loads(content)["csrf"]
        boundary = uuid.uuid4().hex
        body = (f'--{boundary}\r\nContent-Disposition: form-data; name="kind"\r\n\r\ntarja\r\n'
                f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="t.xlsx"\r\n'
                'Content-Type: application/octet-stream\r\n\r\n').encode() + self.files["tarja"][1] + \
            f"\r\n--{boundary}--\r\n".encode()
        headers = {"Origin": ORIGIN, "X-CSRF-Token": csrf, "Cookie": cookie,
                   "Content-Type": f"multipart/form-data; boundary={boundary}"}
        self.assertEqual(call("POST", "/api/nami/sources", body, headers)[0], 201)
        body = (f'--{boundary}\r\nContent-Disposition: form-data; name="kind"\r\n\r\nmatrix\r\n'
                f'--{boundary}\r\nContent-Disposition: form-data; name="file"; filename="m.xlsx"\r\n'
                'Content-Type: application/octet-stream\r\n\r\n').encode() + data + f"\r\n--{boundary}--\r\n".encode()
        status, _, content = call("POST", "/api/nami/sources", body, headers)
        self.assertEqual(status, 201, content[:200])
        self.assertEqual(json.loads(content)["source"]["size"], len(data))

    def test_expanded_size_limits_differ_for_planning_workbooks_and_word_documents(self):
        from .workflow import MAX_EXPANDED_DOCUMENT, MAX_EXPANDED_WORKBOOK, check_workbook
        directory = Path(self.base.name)

        def bomb(name, megabytes):
            path = directory / name
            with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
                with archive.open("xl/worksheets/sheet1.xml", "w", force_zip64=True) as stream:
                    chunk = bytes(1024 * 1024)
                    for _ in range(megabytes):
                        stream.write(chunk)
            return path

        mid = bomb("mid.zip", 180)
        check_workbook(mid)
        with self.assertRaises(ValueError):
            check_workbook(mid, MAX_EXPANDED_DOCUMENT)
        with self.assertRaises(ValueError):
            check_workbook(bomb("huge.zip", MAX_EXPANDED_WORKBOOK // (1024 * 1024) + 20))

    def test_uploaded_files_are_private_hashed_and_kept_in_history(self):
        self.upload_all()
        first = self.upload("matrix", data=None)
        self.assertTrue(first.json["duplicate"])
        newer = self.files["matrix"][1]
        book = load_workbook(io.BytesIO(newer))
        book["Matriz"]["D5"] = "Pendiente"
        stream = io.BytesIO()
        book.save(stream)
        changed = self.upload("matrix", name="matriz nueva.xlsx", data=stream.getvalue())
        self.assertEqual(changed.status_code, 201)
        self.assertFalse(changed.json["duplicate"])
        listing = self.client.get("/api/nami/sources").json
        self.assertEqual(listing["active"]["matrix"]["filename"], "matriz nueva.xlsx")
        self.assertEqual(len([item for item in listing["history"] if item["kind"] == "matrix"]), 2)
        for stored in (self.root / "sources").iterdir():
            self.assertEqual(stored.stat().st_mode & 0o777, 0o600)
        self.assertEqual((self.root / "sources").stat().st_mode & 0o777, 0o700)
        self.assertNotIn(str(self.root), json.dumps(listing))
        db = connect(self.root)
        try:
            with self.assertRaises(Exception):
                db.execute("DELETE FROM sources")
        finally:
            db.close()

    def test_plan_requires_sources_and_reports_shifts_and_pending(self):
        missing = self.client.get("/api/nami/plan")
        self.assertEqual((missing.status_code, missing.json["missing"]), (409, ["tarja", "crews", "matrix"]))
        self.upload_all()
        plan = self.client.get("/api/nami/plan?courses=RF-01").json
        crews = {crew["id"]: crew for crew in plan["crews"]}
        self.assertEqual((crews["1"]["shift"], crews["2"]["shift"]), ("Turno Día", "Turno Noche"))
        self.assertEqual((crews["1"]["pending"], crews["2"]["pending"]), (2, 1))
        self.assertTrue(crews["1"]["foreman_found"])
        self.assertEqual(plan["warnings"]["excluidos"], 1)
        courses = {course["code"]: course["evaluation"] for course in plan["courses"]}
        self.assertEqual(courses, {"RF-01": True, "RF-02": True, "RF-03": False})
        self.assertEqual(self.client.get("/api/nami/plan?courses=RF-99").status_code, 400)

    @unittest.skipUnless(LIBREOFFICE, "LibreOffice requerido")
    def test_generated_record_has_foreman_evaluations_and_is_registered_for_scanning(self):
        self.upload_all()
        response = self.post("/records/word", {"crew_id": "1", "date": "2026-10-12", "courses": ["RF-01", "RF-02"]})
        self.assertEqual(response.status_code, 201, response.json)
        summary = response.json["summary"]
        self.assertEqual((summary["included"], summary["already_approved"], summary["evaluation_sheets"]), (3, 1, 4))
        self.assertEqual(response.json["shift"], "Turno Día")
        db = connect(self.root)
        try:
            manifest = json.loads(db.execute("SELECT manifest FROM records WHERE id=?",
                                             (response.json["id"],)).fetchone()[0])
        finally:
            db.close()
        self.assertEqual(manifest["layout"], "word")
        self.assertEqual(len(manifest["sources"]), 3)
        self.assertTrue(all(len(item["sha256"]) == 64 for item in manifest["sources"].values()))
        self.assertEqual({p["rut"] for p in manifest["participants"]},
                         {RUTS["a1"], RUTS["a2"], RUTS["foreman_day"]})
        page = manifest["pages"][0]
        self.assertEqual(page["size"], [612.0, 792.0])
        self.assertEqual(len(page["rows"]), 3)
        self.assertEqual(valid_payload(page["qr"], SECRET)[0], manifest["id"])
        register = self.client.get(response.json["pdf_url"])
        with fitz.open(stream=register.data, filetype="pdf") as document:
            text = "".join(p.get_text() for p in document)
            self.assertIn("CAPATAZ: CAPATAZ DIA UNO", text)
            self.assertIn("12-10-2026", text)
            self.assertNotIn("CARLA", text)
            digits = re.sub(r"\D", "", text)
            for key in ("a1", "a2", "foreman_day"):
                self.assertNotIn(RUTS[key], text, "the RUT must stay blank for the worker to write")
                self.assertNotIn(RUTS[key].split("-")[0], digits)
        register.close()
        bundle = self.client.get(response.json["bundle_url"])
        self.assertIn("CUADRILLA 1 - CAPATAZ DIA UNO.pdf", bundle.headers["Content-Disposition"])
        with fitz.open(stream=bundle.data, filetype="pdf") as document:
            self.assertEqual(len(document), 1 + 4)
            sheets = [p.get_text() for p in list(document)[1:]]
            import cv2
            import numpy as np
            scale = 150 / 72
            decoded = []
            for number in range(1, len(document)):
                pix = document[number].get_pixmap(matrix=fitz.Matrix(scale, scale), colorspace=fitz.csGRAY, alpha=False)
                gray = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width)
                decoded.append(valid_evaluation_payload(cv2.QRCodeDetector().detectAndDecode(gray)[0], SECRET))
        bundle.close()
        self.assertEqual([item[2] for item in decoded], [0, 1, 2, 3])
        self.assertEqual({item[1] for item in decoded}, {manifest["id"]})
        self.assertEqual(len(manifest["evaluation_pages"]), 4)
        self.assertEqual({entry["rut"] for entry in manifest["evaluation_pages"]},
                         {RUTS["a1"], RUTS["a2"], RUTS["foreman_day"]})
        self.assertTrue(all("Capataz: CAPATAZ DIA UNO" in sheet and "12-10-2026" in sheet for sheet in sheets))
        self.assertEqual(sum("CAPATAZ DIA UNO]" in sheet for sheet in sheets), 1)
        self.assertEqual(sum("BEA SEGUNDA SINTETICA]" in sheet for sheet in sheets), 2)
        self.assertEqual(sum("ANA PRIMERA SINTETICA]" in sheet for sheet in sheets), 1)
        self.assertEqual(os.listdir(self.root / "work"), ["profile"] if "profile" in os.listdir(self.root / "work") else [])

    @unittest.skipUnless(LIBREOFFICE, "LibreOffice requerido")
    def test_scan_analyzer_reads_the_word_layout(self):
        import numpy as np
        import cv2
        from . import scanning
        self.upload_all()
        response = self.post("/records/word", {"crew_id": "1", "date": "2026-10-12", "courses": ["RF-01", "RF-02"]})
        pdf = self.client.get(response.json["pdf_url"]).data
        db = connect(self.root)
        try:
            manifest = json.loads(db.execute("SELECT manifest FROM records").fetchone()[0])
        finally:
            db.close()
        page = manifest["pages"][0]
        out = fitz.open()
        with fitz.open(stream=pdf, filetype="pdf") as source:
            pix = source[0].get_pixmap(matrix=fitz.Matrix(150 / 72, 150 / 72), colorspace=fitz.csRGB, alpha=False)
            image = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width, 3).copy()
            row = page["rows"][1]
            y = int((row["box"][1] + row["box"][3] / 2) * 150 / 72)
            cv2.line(image, (int((row["rut_x"] - 3) * 150 / 72), y),
                     (int((row["signature_x"] + 30) * 150 / 72), y), (220, 30, 30), 3)
            matrix = cv2.getRotationMatrix2D((image.shape[1] / 2, image.shape[0] / 2), -0.6, 0.97)
            image = cv2.warpAffine(image, matrix, (image.shape[1], image.shape[0]), borderValue=(255, 255, 255))
            encoded = cv2.imencode(".jpg", cv2.cvtColor(image, cv2.COLOR_RGB2BGR), [cv2.IMWRITE_JPEG_QUALITY, 85])[1]
            scan = out.new_page(width=source[0].rect.width, height=source[0].rect.height)
            scan.insert_image(scan.rect, stream=encoded.tobytes())
        scan_path = self.root / "scans" / "synthetic.pdf"
        out.save(str(scan_path))
        result = scanning.analyze(self.root, "synthetic", scan_path, {manifest["id"]: manifest}, SECRET)
        self.assertEqual(result[0]["record_id"], manifest["id"])
        self.assertEqual([r["suggested"] for r in result[0]["rows"]], ["FIRMO", "NO_FIRMO", "FIRMO"])

    @unittest.skipUnless(LIBREOFFICE, "LibreOffice requerido")
    def test_scan_identifies_evaluations_by_qr_alongside_the_register(self):
        from . import scanning
        self.upload_all()
        response = self.post("/records/word", {"crew_id": "1", "date": "2026-10-12", "courses": ["RF-01", "RF-02"]})
        bundle = self.client.get(response.json["bundle_url"]).data
        db = connect(self.root)
        try:
            manifest = json.loads(db.execute("SELECT manifest FROM records").fetchone()[0])
        finally:
            db.close()
        path = self.root / "scans" / "bundle.pdf"
        path.write_bytes(bundle)
        result = scanning.analyze(self.root, "bundle", path, {manifest["id"]: manifest}, SECRET)
        self.assertEqual([page["kind"] if "kind" in page else "registro" for page in result],
                         ["registro"] + ["evaluation"] * 4)
        self.assertTrue(all(page["record_id"] == manifest["id"] for page in result))
        self.assertEqual({page["evaluation"]["method"] for page in result[1:]}, {"qr"})
        self.assertEqual(sorted((p["evaluation"]["code"], p["evaluation"]["nombre"]) for p in result[1:]),
                         sorted((e["code"], e["nombre"]) for e in manifest["evaluation_pages"]))
        self.assertTrue((self.root / "rows" / f'{result[1]["preview_id"]}.png').is_file())
        repeated = fitz.open()
        repeated.insert_pdf(fitz.open(stream=bundle, filetype="pdf"), from_page=1, to_page=1)
        repeated.insert_pdf(fitz.open(stream=bundle, filetype="pdf"), from_page=1, to_page=1)
        path.write_bytes(repeated.tobytes())
        again = scanning.analyze(self.root, "again", path, {manifest["id"]: manifest}, SECRET)
        self.assertEqual([page["record_id"] is None for page in again], [False, True])

    @unittest.skipUnless(tesseract_path(), "Tesseract requerido")
    def test_scan_identifies_printed_evaluations_without_qr_by_title_and_name(self):
        import cv2
        import numpy as np
        from . import scanning
        record_id = str(uuid.uuid4())
        manifest = {"id": record_id, "version": 3, "pages": [{}],
                    "courses": [{"id": "RF-02", "code": "RF-02"}],
                    "participants": [{"nombre": "BEA SEGUNDA SINTETICA", "rut": "2", "pending": ["RF-02"]}]}
        document = fitz.open()
        page = document.new_page(width=612, height=792)
        page.insert_text((150, 50), "EVALUACION DE ENTENDIMIENTO", fontsize=16)
        page.insert_text((190, 75), "RF-02 Trabajo en Altura vs 03", fontsize=16)
        page.insert_text((40, 760), "[BEA SEGUNDA SINTETICA]      Capataz: CAPATAZ DIA UNO", fontsize=14)
        pix = page.get_pixmap(matrix=fitz.Matrix(2, 2), colorspace=fitz.csGRAY, alpha=False)
        gray = np.frombuffer(pix.samples, np.uint8).reshape(pix.height, pix.width)
        scan = fitz.open()
        sheet = scan.new_page(width=612, height=792)
        sheet.insert_image(sheet.rect, stream=cv2.imencode(".png", gray)[1].tobytes())
        path = self.root / "scans" / "printed.pdf"
        scan.save(str(path))
        result = scanning.analyze(self.root, "printed", path, {record_id: manifest}, SECRET)
        self.assertEqual(result[0]["record_id"], record_id, result)
        self.assertEqual((result[0]["evaluation"]["code"], result[0]["evaluation"]["method"]), ("RF-02", "texto"))

    @unittest.skipUnless(LIBREOFFICE, "LibreOffice requerido")
    def test_night_crew_flags_people_missing_from_the_matrix_and_rejects_bad_requests(self):
        self.upload_all()
        response = self.post("/records/word", {"crew_id": "2", "date": "2026-10-12", "courses": ["RF-01"]})
        self.assertEqual(response.status_code, 201, response.json)
        self.assertEqual(response.json["shift"], "Turno Noche")
        self.assertEqual(response.json["summary"]["without_matrix"], 1)
        self.assertEqual(response.json["participants"], 1)
        self.assertEqual(self.post("/records/word", {"crew_id": "99", "date": "2026-10-12",
                                                      "courses": ["RF-01"]}).status_code, 422)
        self.assertEqual(self.post("/records/word", {"crew_id": "1", "date": "2026-13-45",
                                                      "courses": ["RF-01"]}).status_code, 400)
        self.assertEqual(self.post("/records/word", {"crew_id": "1", "date": "2026-10-12",
                                                      "courses": ["RF-99"]}).status_code, 422)
        self.assertEqual(self.post("/records/word", {"crew_id": "1", "date": "2026-10-12",
                                                      "courses": ["RF-01"] * 5}).status_code, 400)

    def test_generation_needs_the_private_register_template(self):
        app, _ = self.make_app(prepare_templates=False)
        client = self.login(app)
        for kind in ("tarja", "crews", "matrix"):
            self.assertEqual(self.upload(kind, client=client).status_code, 201)
        response = client.post("/api/nami/records/word", headers=self.headers(),
                               json={"crew_id": "1", "date": "2026-10-12", "courses": ["RF-01"]})
        self.assertEqual(response.status_code, 422)
        self.assertIn("plantilla", response.json["error"])

    @unittest.skipUnless(LIBREOFFICE, "LibreOffice requerido")
    def test_batch_downloads_a_zip_with_day_and_night_folders(self):
        self.upload_all()
        started = self.post("/batches", {"crew_ids": ["1", "2"], "date": "2026-10-12", "courses": ["RF-01"]})
        self.assertEqual(started.status_code, 202)
        deadline = time.time() + 180
        while time.time() < deadline:
            status = self.client.get(f"/api/nami/batches/{started.json['id']}").json
            if status["status"] != "running":
                break
            time.sleep(1)
        self.assertEqual((status["status"], status["done"], status["errors"]), ("done", 2, []))
        archive = self.client.get(status["zip_url"])
        self.assertEqual(archive.status_code, 200)
        with zipfile.ZipFile(io.BytesIO(archive.data)) as zipped:
            self.assertEqual(sorted(zipped.namelist()), ["Turno Día/CUADRILLA 1 - CAPATAZ DIA UNO.pdf",
                                                         "Turno Noche/CUADRILLA 2 - CAPATAZ NOCHE DOS.pdf"])
            for name in zipped.namelist():
                with fitz.open(stream=zipped.read(name), filetype="pdf") as document:
                    self.assertGreaterEqual(len(document), 2)
        archive.close()
        self.assertFalse([name for name in os.listdir(self.root / "work") if name.endswith(".zip")])
        self.assertEqual(self.post("/batches", {"crew_ids": ["1", "1"], "date": "2026-10-12",
                                                "courses": ["RF-01"]}).status_code, 400)
        self.assertEqual(self.client.get("/api/nami/batches/" + str(uuid.uuid4())).status_code, 404)

    def wait_for_batch(self, identifier):
        deadline = time.time() + 240
        while time.time() < deadline:
            status = self.client.get(f"/api/nami/batches/{identifier}").json
            if status["status"] != "running":
                return status
            time.sleep(1)
        self.fail("El lote no terminó a tiempo")

    @unittest.skipUnless(LIBREOFFICE, "LibreOffice requerido")
    def test_grouped_batch_shares_a_sheet_between_small_crews_of_the_same_shift(self):
        self.upload_all()
        plan = self.client.get("/api/nami/plan?courses=RF-01&group=1").json
        self.assertEqual(plan["groups"], [["1", "3"], ["2"]])
        self.assertIsNone(self.client.get("/api/nami/plan?courses=RF-01").json["groups"])
        started = self.post("/batches", {"crew_ids": ["1", "2", "3"], "date": "2026-10-12",
                                         "courses": ["RF-01"], "group_small": True})
        self.assertEqual((started.status_code, started.json["units"]), (202, [["1", "3"], ["2"]]))
        status = self.wait_for_batch(started.json["id"])
        self.assertEqual((status["status"], status["done"], status["errors"]), ("done", 2, []))
        archive = self.client.get(status["zip_url"])
        with zipfile.ZipFile(io.BytesIO(archive.data)) as zipped:
            self.assertEqual(sorted(zipped.namelist()), ["Turno Día/CUADRILLAS 1, 3.pdf",
                                                         "Turno Noche/CUADRILLA 2 - CAPATAZ NOCHE DOS.pdf"])
            with fitz.open(stream=zipped.read("Turno Día/CUADRILLAS 1, 3.pdf"), filetype="pdf") as document:
                self.assertEqual(len(document), 1 + 4)
                register = document[0].get_text()
                sheets = [page.get_text() for page in list(document)[1:]]
        archive.close()
        self.assertIn("CAPATACES: CAPATAZ DIA UNO / CAPATAZ DIA TRES", register)
        self.assertIn("Cuadrilla 1", register)
        self.assertIn("Cuadrilla 3", register)
        self.assertEqual(sum("Capataz: CAPATAZ DIA UNO" in sheet for sheet in sheets), 2)
        self.assertEqual(sum("Capataz: CAPATAZ DIA TRES" in sheet for sheet in sheets), 2)
        db = connect(self.root)
        try:
            manifests = [json.loads(row[0]) for row in db.execute("SELECT manifest FROM records")]
        finally:
            db.close()
        group = next(m for m in manifests if m["crew"].get("group"))
        self.assertEqual([item["id"] for item in group["crew"]["group"]], ["1", "3"])
        self.assertEqual([p["crew_id"] for p in group["participants"]], ["1", "1", "3", "3"])
        self.assertEqual(len(group["pages"]), 1)
        self.assertEqual(len(group["pages"][0]["rows"]), 4)
        single = next(m for m in manifests if not m["crew"].get("group"))
        self.assertFalse(any("crew_label" in p for p in single["participants"]))

    def test_plan_lists_every_cargo_in_the_tarja_with_pending_counts(self):
        self.upload_all()
        plain = self.client.get("/api/nami/plan").json
        self.assertEqual([(c["label"], c["total"], c["pending"]) for c in plain["cargos"]],
                         [("CAPATAZ", 3, None), ("CARPINTERO", 1, None), ("MAESTRO PRIMERA", 5, None)])
        one = self.client.get("/api/nami/plan?courses=RF-01").json
        self.assertEqual([(c["key"], c["pending"]) for c in one["cargos"]],
                         [("CAPATAZ", 2), ("CARPINTERO", 1), ("MAESTRO PRIMERA", 2)])
        two = self.client.get("/api/nami/plan?courses=RF-01,RF-02").json
        self.assertEqual([c["pending"] for c in two["cargos"]], [2, 1, 3])

    def test_cargo_and_relator_requests_are_validated(self):
        self.upload_all()
        base = {"date": "2026-10-12", "courses": ["RF-01"]}
        invalid = [
            {**base, "cargos": ["CAPATAZ"], "crew_ids": ["1"]},
            dict(base),
            {**base, "cargos": []}, {**base, "cargos": "CAPATAZ"}, {**base, "cargos": [1]},
            {**base, "cargos": ["NO EXISTE"]},
            {**base, "cargos": ["CAPATAZ", "capataz"]},
            {**base, "cargos": ["CAPATAZ"], "group_small": True},
            {**base, "cargos": ["x" * 101]},
            {**base, "cargos": ["CAPATAZ"], "relator": "x" * 101},
            {**base, "cargos": ["CAPATAZ"], "relator": "línea\nnueva"},
            {**base, "cargos": ["CAPATAZ"], "relator": 5},
            {**base, "crew_ids": ["1"], "relator": ["a"]},
        ]
        for payload in invalid:
            self.assertEqual(self.post("/batches", payload).status_code, 400, payload)
        self.assertEqual(self.post("/records/word", {"crew_id": "1", **base, "relator": "x" * 101}).status_code, 400)

    @unittest.skipUnless(LIBREOFFICE, "LibreOffice requerido")
    def test_relator_is_written_in_the_register_and_kept_in_the_record(self):
        self.upload_all()
        request = {"crew_id": "1", "date": "2026-10-12", "courses": ["RF-01"]}
        with_relator = self.post("/records/word", {**request, "relator": "  MARÍA   PÉREZ "})
        self.assertEqual(with_relator.status_code, 201, with_relator.json)
        self.assertEqual(with_relator.json["relator"], "MARÍA PÉREZ")
        without = self.post("/records/word", request)
        self.assertEqual((without.status_code, without.json["relator"]), (201, ""))
        for response, expected in ((with_relator, True), (without, False)):
            document = self.client.get(response.json["pdf_url"])
            with fitz.open(stream=document.data, filetype="pdf") as pdf:
                self.assertEqual("MARÍA PÉREZ" in pdf[0].get_text(), expected)
            document.close()
        db = connect(self.root)
        try:
            stored = json.loads(db.execute("SELECT manifest FROM records WHERE id=?",
                                           (with_relator.json["id"],)).fetchone()[0])
        finally:
            db.close()
        self.assertEqual(stored["relator"], "MARÍA PÉREZ")

    @unittest.skipUnless(LIBREOFFICE, "LibreOffice requerido")
    def test_cargo_batch_makes_one_register_per_cargo_with_the_relator(self):
        self.upload_all()
        started = self.post("/batches", {"cargos": ["capataz", "Maestro  Primera", "carpintero"],
                                         "date": "2026-10-12", "courses": ["RF-01"], "relator": "MARÍA PÉREZ"})
        self.assertEqual(started.status_code, 202, started.json)
        self.assertEqual(started.json["units"], [["cargo:CAPATAZ"], ["cargo:MAESTRO PRIMERA"], ["cargo:CARPINTERO"]])
        status = self.wait_for_batch(started.json["id"])
        self.assertEqual((status["status"], status["done"], status["errors"]), ("done", 3, []))
        archive = self.client.get(status["zip_url"])
        texts = {}
        with zipfile.ZipFile(io.BytesIO(archive.data)) as zipped:
            self.assertEqual(sorted(zipped.namelist()), ["Por cargo/CARGO CAPATAZ.pdf", "Por cargo/CARGO CARPINTERO.pdf",
                                                         "Por cargo/CARGO MAESTRO PRIMERA.pdf"])
            for name in zipped.namelist():
                with fitz.open(stream=zipped.read(name), filetype="pdf") as document:
                    texts[name.split("/")[1]] = (" ".join(document[0].get_text().split()), len(document))
        archive.close()
        for register, _ in texts.values():
            self.assertIn("MARÍA PÉREZ", register)
            self.assertNotIn("CAPATAZ:", register, "cargo registers have no foreman line")
        chiefs, chiefs_pages = texts["CARGO CAPATAZ.pdf"]
        self.assertIn("CAPATAZ DIA UNO", chiefs)
        self.assertIn("CAPATAZ DIA TRES", chiefs)
        self.assertNotIn("CAPATAZ NOCHE DOS", chiefs, "already approved")
        self.assertEqual(chiefs_pages, 1 + 2)
        masters, _ = texts["CARGO MAESTRO PRIMERA.pdf"]
        self.assertIn("BEA SEGUNDA SINTETICA", masters)
        self.assertIn("DANA CUARTA SINTETICA", masters)
        self.assertNotIn("ANA PRIMERA", masters)
        self.assertNotIn("FANY SEXTA SINTETICA", masters)
        self.assertIn("FANY SEXTA SINTETICA", texts["CARGO CARPINTERO.pdf"][0], "blank CATEGORIA falls back to ESPECIALIDAD")
        listed = {r["crew"]: (r["shift"], r["relator"]) for r in self.client.get("/api/nami/records").json["records"]}
        self.assertEqual(listed["CARGO CAPATAZ"], ("Por cargo", "MARÍA PÉREZ"))
        db = connect(self.root)
        try:
            manifests = [json.loads(row[0]) for row in db.execute("SELECT manifest FROM records")]
        finally:
            db.close()
        chief = next(m for m in manifests if m["crew"]["name"] == "CARGO CAPATAZ")
        self.assertEqual((chief["crew"]["cargo"], chief["crew"]["foreman"]), ("CAPATAZ", ""))
        self.assertEqual({p["rut"] for p in chief["participants"]}, {RUTS["foreman_day"], RUTS["foreman_day3"]})
        self.assertTrue(all(len(page["rows"]) == len(chief["participants"]) for page in chief["pages"]))
    def test_group_small_must_be_a_boolean(self):
        self.upload_all()
        self.assertEqual(self.post("/batches", {"crew_ids": ["1"], "date": "2026-10-12", "courses": ["RF-01"],
                                                "group_small": "si"}).status_code, 400)

    def test_excel_export_lists_confirmed_diffusion_without_formulas(self):
        self.upload_all()
        manifest = {"id": str(uuid.uuid4()), "version": 3, "date": "2026-10-12", "layout": "word",
                    "relator": "MARÍA PÉREZ",
                    "crew": {"id": "1", "name": "CUADRILLA 1 - =CAPATAZ", "foreman": "=CAPATAZ", "shift": "NOCHE"},
                    "participants": [{"rut": RUTS["a1"], "nombre": "=1+1 PRUEBA"}],
                    "courses": [{"id": "RF-034", "code": "RF-01", "title": "@Curso"}], "pages": []}
        scan = str(uuid.uuid4())
        db = connect(self.root)
        try:
            db.execute("INSERT INTO records(id,manifest,pdf) VALUES (?,?,?)",
                       (manifest["id"], json.dumps(manifest), "x.pdf"))
            db.execute("INSERT INTO scans(id,sha256,pdf,pages) VALUES (?,?,?,?)", (scan, "0" * 64, "x.pdf", "[]"))
            db.execute("""INSERT INTO results(scan_id,row_id,record_id,rut,course_id,status,confirmed_by)
                          VALUES (?,?,?,?,?,?,?)""", (scan, "r1", manifest["id"], RUTS["a1"], "RF-034", "NO_FIRMO", "tester"))
        finally:
            db.close()
        response = self.client.get("/api/nami/results/export.xlsx")
        self.assertEqual(response.status_code, 200)
        book = load_workbook(io.BytesIO(response.data))
        sheet = book["Difusión confirmada"]
        values = [cell.value for cell in sheet[2]]
        self.assertEqual(values[3:9], ["Turno Noche", RUTS["a1"], "=1+1 PRUEBA", "RF-01", "@Curso", "No firmó"])
        self.assertEqual(values[12], "MARÍA PÉREZ")
        self.assertEqual(sheet["F2"].data_type, "s")
        self.assertEqual(sheet["C2"].data_type, "s")
        self.assertIn("NO equivale a aprobar", book["Leyenda"]["A1"].value)


if __name__ == "__main__":
    unittest.main()
