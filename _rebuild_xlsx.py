"""Regenera el consolidado .xlsx aplicando sobre el anterior los cambios reales del CSV.
Se usa el xlsx como base porque el CSV exportado perdio todos los acentos (U+FFFD).
"""
import copy
import datetime
import openpyxl

BASE = "Consolidado Incidentes y Medidas Correctivas (Formato Interno) 05.08.2026.xlsx"

wb = openpyxl.load_workbook(BASE)
ws = wb[wb.sheetnames[0]]

MEDIDA, RESP, CIERRE, ESTATUS, VERIFM = 7, 8, 9, 10, 11

CELL_EDITS = [
    # item 51 - medida reescrita y cerrada
    (175, MEDIDA, "Se aplicaron dos medidas; Difusión del evento y pavimentación "),
    (175, ESTATUS, "CERRADO"),
    (175, VERIFM, "CERRADO"),
    # item 67 - typo "por pate" -> "por parte" + detalle de verificacion
    (212, MEDIDA, "Capacitación por parte del departamento de calidad al personal asignado al pre armado "
                  "de soportes en ambos turnos.\u200b\n* Registro asistencia personal eléctrico ambos turnos"),
    (212, VERIFM, "PENDIENTE Capacitación por parte del departamento de calidad a turno A y B "),
    # item 71 - medida cerrada
    (222, ESTATUS, "CERRADO"),
    (222, VERIFM, None),
    # item 78 - medida 1 cerrada
    (237, ESTATUS, "CERRADO"),
    (237, VERIFM, "Comentario: Se emitieron 2 boletines. Falta difusión de boletin Turno A y B"),
    # item 94
    (274, MEDIDA, "Implementar en proyecto Procedimiento de Gestión de Cambio.\n"
                  "Entregable: Procedimiento generado, difundido a línea de supervisión"),
    (275, ESTATUS, "CERRADO"),
    (278, MEDIDA, "\nTaller de aplicabilidad del Plan Local de Emergencia Obra, dirigido a capataces "
                  "y supervisores.\n\nR: Registro de asistencia"),
    (278, ESTATUS, "PENDIENTE TA"),
    (280, ESTATUS, "PENDIENTE TA \nTB NOCHE"),
    # item 96
    (284, VERIFM, "Generar boletin Pendiente difusión 29-09"),
    # item 99 - fecha corregida
    (290, 2, datetime.datetime(2026, 9, 15)),
]

for row, col, value in CELL_EDITS:
    ws.cell(row, col).value = value

# ---- item 100 (nuevo) ----
NEW_ROW = 292
TEMPLATE = 290
NEW_ITEM = {
    1: 100,
    2: datetime.datetime(2026, 9, 24),
    3: "10X10 B",
    4: "GOLPE EN CANASTILLO DE EQUIPO MANIPULADOR",
    5: "DAÑO MATERIAL",
    6: "Durante el traslado del manipulador telescópico MAGNI 4.5 por la Calle 42, el operador "
       "realizó una maniobra de giro para ingresar a la Zanja 13. En dicha maniobra, el canastillo "
       "del equipo impactó contra la caja de la calle, dañándose en el costado izquierdo.\n"
       "El área se encontraba debidamente iluminada y segregada, sin exposición de personas.\n",
    MEDIDA: "PENDIENTE REVISIÓN DE ICAM POR JAIME GALLEGOS ",
}

for col in range(1, 14):
    src, dst = ws.cell(TEMPLATE, col), ws.cell(NEW_ROW, col)
    dst._style = copy.copy(src._style)
    dst.value = NEW_ITEM.get(col)
ws.row_dimensions[NEW_ROW].height = ws.row_dimensions[TEMPLATE].height

wb.save(BASE)
print(f"Guardado {BASE}: item 100 agregado en fila {NEW_ROW}, {len(CELL_EDITS)} celdas actualizadas.")
