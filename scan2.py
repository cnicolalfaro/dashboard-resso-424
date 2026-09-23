from openpyxl import load_workbook
from pathlib import Path
p=Path('Consolidado Incidentes y Medidas Correctivas (Formato Interno) 05.08.2026.xlsx')
wb=load_workbook(p,read_only=True,data_only=True)
def s(x): return '' if x is None else str(x).strip()
def n(x): return ' '.join(s(x).upper().split())
F=['ITEM','NOMBRE DEL EVENTO','MEDIDA CORRECTIVA','RESPONSABLE','FECHA CIERRE ACCION','ESTATUS','VERIFICACION MEDIDAS','VERIFICACION EFICACIA','COMENTARIOS']
for ws in wb.worksheets:
 rows=list(ws.iter_rows(values_only=True)); hi=None; h=None
 for i,r in enumerate(rows):
  z=[n(x) for x in r]
  if 'ITEM' in z and any('MEDIDA' in x for x in z): hi=i; h=z; break
 if hi is None: continue
 ci={f:next((j for j,x in enumerate(h) if x==f or (f=='MEDIDA CORRECTIVA' and 'MEDIDA' in x)),None) for f in F}
 print('SHEET',ws.title,'HEADER',hi+1,'COLS',ci)
 def get(r,f): return s(r[ci[f]]) if ci[f] is not None and ci[f]<len(r) else ''
 def row(r): return '\t'.join(get(r,f) for f in F)
 i30=[]; bad=[]
 for r in rows[hi+1:]:
  item=n(get(r,'ITEM')); st=n(get(r,'ESTATUS')); me=n(get(r,'MEDIDA CORRECTIVA')); re=n(get(r,'RESPONSABLE'))
  if item=='30': i30.append(row(r))
  if (not st or any(q in st for q in ('ABIERTO','PENDIENTE','FALTA'))) and (re in ('','SIN RESPONSABLE','NO APLICA','SIN INFORME FINAL') or 'SIN INFORME FINAL' in me or 'SIN INFORME FINAL' in re): bad.append(row(r))
 print('ITEM30_COUNT',len(i30)); print(*i30,sep='\n')
 print('FLAGGED_COUNT',len(bad)); print(*bad,sep='\n')
wb.close()
