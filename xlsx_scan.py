from openpyxl import load_workbook
from pathlib import Path
p=Path('Consolidado Incidentes y Medidas Correctivas (Formato Interno) 05.08.2026.xlsx')
wb=load_workbook(p, read_only=True, data_only=True)

def s(x): return '' if x is None else str(x).strip()
def n(x): return ' '.join(s(x).upper().split())
fields=['ITEM','NOMBRE DEL EVENTO','MEDIDA CORRECTIVA','RESPONSABLE','FECHA CIERRE ACCION','ESTATUS','VERIFICACION MEDIDAS','VERIFICACION EFICACIA','COMENTARIOS']
for ws in wb.worksheets:
    allrows=list(ws.iter_rows(values_only=True)); hi=None; hdr=None
    for i,row in enumerate(allrows):
        h=[n(x) for x in row]
        if 'ITEM' in h and ('MEDIDA CORRECTIVA' in h or 'MEDIDA' in h): hi=i; hdr=h; break
    if hi is None: continue
    ci={f:(hdr.index(f) if f in hdr else hdr.index('MEDIDA')) for f in fields if f in hdr or (f=='MEDIDA CORRECTIVA' and 'MEDIDA' in hdr)}
    def rec(row): return [s(row[ci[f]]) if ci[f]<len(row) else '' for f in fields]
    item30=[]; flagged=[]
    for row in allrows[hi+1:]:
        if 'ITEM' not in ci: continue
        item=n(row[ci['ITEM']]) if ci['ITEM']<len(row) else ''
        measure=n(row[ci['MEDIDA CORRECTIVA']]) if ci['MEDIDA CORRECTIVA']<len(row) else ''
        resp=n(row[ci['RESPONSABLE']]) if ci['RESPONSABLE']<len(row) else ''
        status=n(row[ci['ESTATUS']]) if ci['ESTATUS']<len(row) else ''
        if item=='30': item30.append(rec(row))
        if (not status or any(x in status for x in ('ABIERTO','PENDIENTE','FALTA'))) and (resp in ('','SIN RESPONSABLE','NO APLICA','SIN INFORME FINAL') or 'SIN INFORME FINAL' in measure or 'SIN INFORME FINAL' in resp): flagged.append(rec(row))
    print('SHEET',ws.title,'HEADER',hi+1)
    print('ITEM30',len(item30))
    for r in item30: print('\t'.join(r))
    print('FLAGGED',len(flagged))
    for r in flagged: print('\t'.join(r))
wb.close()
