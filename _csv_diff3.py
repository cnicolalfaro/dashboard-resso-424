import csv, re
import openpyxl

OLD = "Consolidado Incidentes y Medidas Correctivas (Formato Interno) 05.08.2026.xlsx"
NEW = "Consolidado Incidentes y Medidas Correctivas (Formato Interno) 05.08.csv"


def tidy(v):
    s = str(v if v is not None else "")
    return re.sub(r"\s+", " ", s).strip().upper()


def key(v):
    s = tidy(v).lower().translate(str.maketrans("áéíóúüñ", "aeiouun"))
    return re.sub(r"[^a-z0-9]", "", s)


def norm_date(v):
    s = tidy(v)
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
    m = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})(.*)$", s)
    if m:
        y = m.group(3)
        y = "20" + y if len(y) == 2 else y
        return f"{int(m.group(1)):02d}-{int(m.group(2)):02d}-{y}{m.group(4)}"
    return s


def same(old, new):
    """True when the only difference is a lost character (U+FFFD acts as wildcard)."""
    if old == new:
        return True
    pattern = "".join("." if ch == "\ufffd" else re.escape(ch) for ch in new)
    return re.fullmatch(pattern, old, re.DOTALL) is not None


def build(rows):
    hidx = next(i for i, r in enumerate(rows)
                if "item" in [key(c) for c in r] and "nombredelevento" in [key(c) for c in r])
    headers = [key(c) for c in rows[hidx]]

    def col(*names):
        for n in names:
            if n in headers:
                return headers.index(n)
        return -1

    C = {"item": col("item"), "fecha": col("fecha"), "turno": col("turno"),
         "nombre": col("nombredelevento"), "categoria": col("categoria"),
         "incidente": col("incidente"), "medida": col("medidacorrectiva"),
         "responsable": col("responsable"),
         "cierre": col("fechacierreaccion", "fechacierreaccin", "fechacierre"),
         "estatus": col("estatus"),
         "verifm": col("verificacionmedidas", "verificacinmedidas"),
         "verife": col("verificaciondeeficacia", "verificacindeeficacia", "verificacioneficacia"),
         "comentarios": col("comentarios")}

    def val(r, c):
        return tidy(r[c]) if 0 <= c < len(r) and r[c] is not None else ""

    out, cur = [], None
    for r in rows[hidx + 1:]:
        item = val(r, C["item"])
        if item:
            cur = {"item": item.split(".")[0].strip(), "fecha": norm_date(val(r, C["fecha"])),
                   "turno": val(r, C["turno"]), "nombre": val(r, C["nombre"]),
                   "categoria": val(r, C["categoria"]), "incidente": val(r, C["incidente"]),
                   "verifm": val(r, C["verifm"]), "verife": val(r, C["verife"]),
                   "comentarios": val(r, C["comentarios"]), "medidas": []}
            out.append(cur)
        if cur is None:
            continue
        m = val(r, C["medida"])
        if m:
            cur["medidas"].append({"medida": m, "responsable": val(r, C["responsable"]),
                                   "cierre": norm_date(val(r, C["cierre"])), "estatus": val(r, C["estatus"])})
        else:
            for f, c in (("verifm", C["verifm"]), ("verife", C["verife"])):
                v = val(r, c)
                if v and not cur[f]:
                    cur[f] = v
    return {i["item"]: i for i in out}


wb = openpyxl.load_workbook(OLD, data_only=True)
old = build([list(row) for row in wb[wb.sheetnames[0]].iter_rows(values_only=True)])
with open(NEW, "r", encoding="utf-8", errors="replace", newline="") as fh:
    new = build(list(csv.reader(fh, delimiter=";")))

order = lambda s: sorted(s, key=lambda x: int(re.match(r"\d+", x).group()))
print(f"ANTERIOR {len(old)} incidentes | NUEVO {len(new)} incidentes")
print("NUEVOS:", order(set(new) - set(old)))
print("ELIMINADOS:", order(set(old) - set(new)))

print("\n=== CAMBIOS DE CONTENIDO (no por acentos) ===")
FIELDS = ["fecha", "turno", "nombre", "categoria", "incidente", "verifm", "verife", "comentarios"]
hits = 0
for k in order(set(old) & set(new)):
    a, b = old[k], new[k]
    diffs = []
    for f in FIELDS:
        if not same(a[f], b[f]):
            diffs.append(f"   [{f}]\n      antes: {a[f][:150]}\n      ahora: {b[f][:150]}")
    if len(a["medidas"]) != len(b["medidas"]):
        diffs.append(f"   [n medidas] {len(a['medidas'])} -> {len(b['medidas'])}")
    else:
        for n, (ma, mb) in enumerate(zip(a["medidas"], b["medidas"]), 1):
            for f in ("medida", "responsable", "cierre", "estatus"):
                if not same(ma[f], mb[f]):
                    diffs.append(f"   medida #{n} [{f}]: {ma[f][:90]!r}\n                 -> {mb[f][:90]!r}")
    if diffs:
        hits += 1
        print(f"\nITEM {k}  ({b['nombre'][:65]})")
        print("\n".join(diffs))
print(f"\n({hits} items con cambios reales)")

print("\n=== ITEM(S) NUEVO(S) ===")
for k in order(set(new) - set(old)):
    i = new[k]
    print(f"ITEM {k} | fecha {i['fecha']} | turno {i['turno']} | categoria {i['categoria']}")
    print(f"  Evento: {i['nombre']}")
    print(f"  Descripcion: {i['incidente'][:600]}")
    print(f"  Verif medidas: {i['verifm']} | Verif eficacia: {i['verife']}")
    print(f"  Comentarios: {i['comentarios'][:300]}")
    for n, m in enumerate(i["medidas"], 1):
        print(f"   {n}. {m['medida'][:140]}\n      resp: {m['responsable']} | cierre: {m['cierre']} | estatus: {m['estatus']}")
