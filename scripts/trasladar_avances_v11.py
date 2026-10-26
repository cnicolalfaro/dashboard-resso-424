"""Traslada avances v10 a una copia v11 conservando niveles de preguntas modificadas.

Uso: python scripts/trasladar_avances_v11.py <v10.xlsx> <v11.xlsx> <salida.xlsx> --date DD-MM-AAAA
Requiere openpyxl y lxml. Conserva las partes originales del paquete Excel.
"""
import argparse
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from zipfile import ZipFile

import openpyxl
from lxml import etree

NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
TAG = "{" + NS["s"] + "}"
EQUIVALENT = {"Sin cambio", "Redacción", "Ponderación"}


def convert_level(value):
    if value == "N/A":
        return value
    if not isinstance(value, (int, float)) or not 0 <= value <= 1:
        raise ValueError(f"Nivel v10 invalido: {value!r}")
    return float((Decimal(str(value)) * 4 // 1) / 4)


def prepare_updates(source, target, date):
    previous = {row[0]: row for row in source.worksheets[0].iter_rows(min_row=30, values_only=True)
                if isinstance(row[0], int)}
    current = {row[0]: (index, row) for index, row in
               enumerate(target.worksheets[0].iter_rows(min_row=30, values_only=True), 30)
               if isinstance(row[0], int)}
    updates = {"M29": f"Observación {date}"}
    seen = set()
    for mapping in target["Cambios v10 → v11"].iter_rows(values_only=True):
        if len(mapping) < 9 or not isinstance(mapping[0], int) or not isinstance(mapping[3], (int, float)):
            continue
        id_, kind = mapping[0], mapping[7]
        if id_ in seen or id_ not in current:
            raise ValueError(f"Equivalencia duplicada o desconocida: {id_}")
        seen.add(id_)
        index, record = current[id_]
        if kind == "Nueva":
            continue
        if kind not in EQUIVALENT | {"Modificada", "Modificada / reubicada"}:
            raise ValueError(f"Tipo desconocido: {kind}")
        ids = [int(part.strip()) for part in str(mapping[4]).split("+")]
        rows = [previous[n] for n in ids]
        if kind in EQUIVALENT:
            if len(rows) != 1:
                raise ValueError(f"Equivalencia multiple no permitida para {id_}")
            updates[f"F{index}"] = convert_level(rows[0][5])
        observations = []
        actions = []
        responsible = []
        for n, row in zip(ids, rows):
            observation = row[12] or row[11] or "Sin observación registrada."
            observations.append(f"V10 #{n} al {date} (nivel {row[5]}): {observation}")
            if row[10]:
                observations.append(f"Criterio registrado en v10 #{n}: {row[10]}")
            if row[13]:
                actions.append(f"V10 #{n}: {row[13]}")
            if row[14] and row[14] not in responsible:
                responsible.append(row[14])
        if kind not in EQUIVALENT:
            observations.append(
                f"V11: se conserva el nivel previo ({record[5]}). Avance heredado, "
                f"no validación de exigencias modificadas. Cambio: {mapping[8]}")
            actions.append(f"Validar y complementar evidencia específica de v11: {mapping[8]}")
        else:
            observations.append("V11: nivel trasladado a la escala inferior de 25 puntos.")
        updates[f"M{index}"] = "\n\n".join(observations)
        updates[f"N{index}"] = "\n\n".join(actions)
        updates[f"O{index}"] = " / ".join(str(value) for value in responsible)
    if seen != set(current):
        raise ValueError("La hoja de equivalencias no cubre todos los requisitos.")
    return updates


def write_copy(template, output, updates):
    with ZipFile(template) as source, ZipFile(output, "x") as dest:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename.startswith("xl/worksheets/") and item.filename.endswith(".xml"):
                tree = etree.fromstring(data)
                if item.filename == "xl/worksheets/sheet1.xml":
                    for address, value in updates.items():
                        matches = tree.xpath('.//s:sheetData/s:row/s:c[@r=$address]', namespaces=NS, address=address)
                        if len(matches) != 1:
                            raise ValueError(f"Celda ausente o duplicada: {address}")
                        cell = matches[0]
                        for child in list(cell):
                            cell.remove(child)
                        if isinstance(value, (int, float)):
                            cell.attrib.pop("t", None)
                            etree.SubElement(cell, TAG + "v").text = str(value)
                        else:
                            cell.set("t", "inlineStr")
                            node = etree.SubElement(etree.SubElement(cell, TAG + "is"), TAG + "t")
                            node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
                            node.text = str(value)
                for cell in tree.xpath(".//s:c[s:f]", namespaces=NS):
                    for cached in cell.findall(TAG + "v"):
                        cell.remove(cached)
                data = etree.tostring(tree, encoding="UTF-8", xml_declaration=True)
            elif item.filename == "xl/workbook.xml":
                tree = etree.fromstring(data)
                calc = tree.find(TAG + "calcPr")
                if calc is None:
                    calc = etree.SubElement(tree, TAG + "calcPr")
                calc.set("calcMode", "auto")
                calc.set("fullCalcOnLoad", "1")
                calc.set("forceFullCalc", "1")
                data = etree.tostring(tree, encoding="UTF-8", xml_declaration=True)
            dest.writestr(item, data)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("template", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("--date", required=True)
    args = parser.parse_args()
    datetime.strptime(args.date, "%d-%m-%Y")
    if args.output.exists():
        raise FileExistsError(f"No se sobrescribira {args.output}")
    source = openpyxl.load_workbook(args.source, data_only=True)
    target = openpyxl.load_workbook(args.template, data_only=True)
    updates = prepare_updates(source, target, args.date)
    write_copy(args.template, args.output, updates)
    print(f"Creada copia v11: {args.output}")


if __name__ == "__main__":
    main()
