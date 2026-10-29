#!/usr/bin/env python3
"""Actualiza pages/evidencias (index, plan, do, check, act, auditor) desde la
planilla "Revision_Documental_RESSO_DD-MM-AAAA.xlsx".

Uso:  python scripts/actualizar_evidencias.py <planilla.xlsx> [--dry-run]
"""
import html
import json
import os
import re
import sys
import warnings
from collections import OrderedDict

import openpyxl

warnings.filterwarnings("ignore")
EV = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "pages", "evidencias")
PHASES = [("A. PLAN", "plan"), ("B. DO", "do"), ("C. CHECK", "check"), ("D. ACT", "act")]
FIRST_ROW, HEADER_ROW = 30, 29


def num(v):
    return v if isinstance(v, (int, float)) else None


def clean(v):
    return None if v is None or str(v).strip() == "" else str(v).strip()


def load(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    if not str(ws.cell(HEADER_ROW, 1).value or "").startswith("N"):
        sys.exit("No se encontro la tabla de requisitos en la fila 29 de la primera hoja.")
    recs = []
    final_observations = clean(ws.cell(HEADER_ROW, 14).value) == "Observaciones Finales"
    for r in ws.iter_rows(min_row=FIRST_ROW, values_only=True):
        if not isinstance(r[0], int):
            continue
        status = r[8]
        recs.append(OrderedDict(
            id=r[0], phase=r[1], section=r[2], question=r[3],
            weight=r[4], level=r[5] if num(r[5]) is not None else "N/A",
            result=num(r[6]), gap=num(r[7]), status=status,
            priority=r[9] if status != "No aplica" else "—",
            criterion=clean(r[10]) or "",
            observation=clean(r[12]) or clean(r[11]) or "",
            action="" if final_observations else clean(r[13]) or "", responsible=clean(r[14]) or "",
            evidence_ref=clean(r[15]),
        ))
        if final_observations:
            recs[-1]["final_observation"] = clean(r[13]) or ""
    return recs


def pct(res, wt):
    return round(res / wt * 100, 2) if wt else None


def fmt(p, comma=True):
    if p is None:
        return "N/A"
    s = f"{p:.2f}".rstrip("0").rstrip(".")
    return s.replace(".", ",") if comma else s


def stats(items):
    app = [r for r in items if r["status"] != "No aplica"]
    wt = sum(r["weight"] for r in app)
    res = sum(r["result"] or 0 for r in app)
    return dict(
        n=len(items), napp=len(app), na=len(items) - len(app),
        ok=sum(r["status"] == "Cumple" for r in items),
        menor=sum(r["status"] == "Brecha menor" for r in items),
        rel=sum(r["status"] == "Brecha relevante" for r in items),
        gap=sum(r["status"] in ("Brecha menor", "Brecha relevante") for r in items),
        alta=sum(r["priority"] == "Alta" for r in items), pct=pct(res, wt),
    )


def read(name):
    with open(os.path.join(EV, name), encoding="utf-8", newline="") as f:
        return f.read()


def write(name, text):
    with open(os.path.join(EV, name), "w", encoding="utf-8", newline="") as f:
        f.write(text)


def sub1(pattern, repl, text, label, flags=re.S):
    new, n = re.subn(pattern, repl if callable(repl) else lambda m: repl, text, count=1, flags=flags)
    if n == 0:
        print(f"  ! sin coincidencia: {label}")
    return new


DATA_RE = r"const DATA ?= ?\[.*?\];(?=\r?\n)"


def replace_dates(text, old_dates, new):
    """Cambia fechas de corte en el texto sin tocar la linea DATA."""
    m = re.search(DATA_RE, text, re.S)
    head, data, tail = (text[:m.start()], m.group(0), text[m.end():]) if m else (text, "", "")
    cutoff_dates = re.findall(
        r"(?:Revisión documental · |Estado real al |RESSO del )(\d{2}-\d{2}-\d{4})",
        head + tail,
    )
    for d in set(old_dates) | set(cutoff_dates):
        head, tail = head.replace(d, new), tail.replace(d, new)
    return head + data + tail


def group_of(r):
    return r["section"].split(".")[0]


def group_name(r):
    return r["section"].split(". ", 1)[1]


def groups_of(items):
    g = OrderedDict()
    for r in items:
        g.setdefault(group_of(r), []).append(r)
    return g


def update_group_cards(text, items):
    for code, gi in groups_of(items).items():
        s = stats(gi)
        p = "N/A" if s["pct"] is None else fmt(s["pct"]) + "%"
        width = "0" if s["pct"] is None else fmt(s["pct"], False)
        meta = "No aplican" if s["pct"] is None else f'{s["gap"]} con brecha'

        def block(m, p=p, width=width, meta=meta, n=s["n"]):
            b = m.group(0)
            b = re.sub(r"(</h3></div><strong>)[^<]*(</strong>)", lambda x: x.group(1) + p + x.group(2), b, 1)
            b = re.sub(r"<span>\d+ requisitos</span><span>[^<]*</span>",
                       f"<span>{n} requisitos</span><span>{meta}</span>", b, 1)
            b = re.sub(r'(<div class="bar"><span style="width:)[^"]*(")', lambda x: x.group(1) + width + "%" + x.group(2), b, 1)
            return b

        text = sub1(rf"<div class=\"group-card\" onclick=\"setGroup\('{code}'\)\">.*?<div class=\"bar\">.*?</div>",
                    block, text, f"tarjeta {code}")
    return text


def update_phase_kpis(text, s, key):
    text = re.sub(r'(<div class="kpi-icon">)\d+(</div><div><small>Requisitos \w+</small><strong>)\d+(</strong><span>)[^<]*',
                  lambda m: f'{m.group(1)}{s["n"]}{m.group(2)}{s["n"]}{m.group(3)}{s["napp"]} aplicables · {s["na"]} N/A', text, 1)
    text = re.sub(r"(<small>Cumplen</small><strong>)\d+", lambda m: m.group(1) + str(s["ok"]), text, 1)
    text = re.sub(r"(<small>Con brecha</small><strong>)\d+(</strong><span>)[^<]*",
                  lambda m: f'{m.group(1)}{s["gap"]}{m.group(2)}{s["menor"]} menor · {s["rel"]} relevante', text, 1)
    text = re.sub(r"(<small>Cumplimiento</small><strong>)[^<]*", lambda m: m.group(1) + fmt(s["pct"]) + "%", text, 1)
    return text


def update_hero(text, s, date):
    p = fmt(s["pct"]) + "%"
    text = re.sub(r"(cumplimiento ponderado de <strong>)[^<]*(</strong>, con )\d+( requisitos aplicables)[^<]*",
                  lambda m: f'{m.group(1)}{p}{m.group(2)}{s["napp"]}{m.group(3)} y {s["gap"]} requisitos con brecha.', text, 1)
    text = re.sub(r"(Cumplimiento ponderado actual: <strong>)[^<]*", lambda m: m.group(1) + p, text, 1)
    return text


def filters(text, items):
    opts = "".join(f'<option value="{c}">{c} · {html.escape(group_name(gi[0]))}</option>'
                   for c, gi in groups_of(items).items())
    return sub1(r'(<select id="groupFilter">\s*<option value="all">Todas las sub-secciones</option>).*?(</select>)',
                lambda m: m.group(1) + opts + m.group(2), text, "filtro sub-secciones")


def phase_record(r, check=False):
    if check:
        lv = num(r["level"])
        record = OrderedDict(
            id=r["id"], phase=r["phase"], subsection=r["section"], question=r["question"],
            weight=r["weight"], level=r["level"], weighted=r["result"], not_achieved=r["gap"],
            status=r["status"], priority=r["priority"], criterion=r["criterion"],
            observation=r["observation"], action=r["action"], responsible=r["responsible"],
            link=r["evidence_ref"], level_pct=None if lv is None else round(lv * 100, 2),
            weight_pct=round(r["weight"] * 100, 2), usable_link=r["evidence_ref"] or "",
            link_label="Abrir SharePoint" if r["evidence_ref"] else "")
        if "final_observation" in r:
            record["final_observation"] = r["final_observation"]
        return record
    lv = num(r["level"])
    record = OrderedDict(
        id=r["id"], group=group_of(r), groupName=group_name(r), question=r["question"],
        weight=round(r["weight"] * 100, 2), level="N/A" if lv is None else round(lv * 100),
        status=r["status"], priority=r["priority"], criterion=r["criterion"],
        observation=r["observation"], action=r["action"], responsible=r["responsible"],
        link=r["evidence_ref"] or "")
    if "final_observation" in r:
        record["final_observation"] = r["final_observation"]
    return record


def preserve_previous_actions(records, previous):
    """La columna de observaciones finales no sustituye acciones anteriores."""
    actions = {r["id"]: r["action"] for r in previous}
    for record in records:
        if "final_observation" in record:
            record["action"] = actions[record["id"]]


def js_data(rows, compact):
    sep = (",", ":") if compact else (", ", ": ")
    return json.dumps(rows, ensure_ascii=False, separators=sep).replace("</", "<\\/")


def update_phase_page(phase, name, items, date):
    text = read(name + ".html")
    s = stats(items)
    rows = [phase_record(r, check=name == "check") for r in items]
    head = "const DATA=" if name == "check" else "const DATA = "
    text = sub1(DATA_RE, lambda m: head + js_data(rows, name == "check") + ";", text, "DATA " + name)
    text = update_hero(text, s, date)
    text = update_phase_kpis(text, s, name)
    text = update_group_cards(text, items)
    text = filters(text, items)
    text = replace_dates(text, ["30-09-2026", "14-10-2026"], date)
    write(name + ".html", text)
    return s


def update_auditor(recs, date, tot):
    text = read("auditor.html")
    text = sub1(DATA_RE, lambda m: "const DATA = " + js_data(recs, False) + ";", text, "DATA auditor")
    n, app = tot["n"], tot["napp"]
    kpi = lambda label, val, extra=None: None
    text = re.sub(r'(<div class="kpi-icon">)\d+(</div>\s*<div><small>Requisitos totales</small><strong>)\d+',
                  lambda m: f'{m.group(1)}{n}{m.group(2)}{n}', text, 1)
    text = re.sub(r"(<small>Cumplen</small><strong>)\d+", lambda m: m.group(1) + str(tot["ok"]), text, 1)
    text = re.sub(r"(<small>Con brecha</small><strong>)\d+", lambda m: m.group(1) + str(tot["gap"]), text, 1)
    text = re.sub(r"(<small>Prioridad alta</small><strong>)\d+", lambda m: m.group(1) + str(tot["alta"]), text, 1)
    text = re.sub(r"(<small>Cumplimiento total</small><strong>)[^<]*", lambda m: m.group(1) + fmt(tot["pct"], False) + "%", text, 1)

    high = [r for r in recs if r["priority"] == "Alta"]
    cards = "".join(
        '\n        <div class="finding high">\n'
        f'          <strong>#{r["id"]} · {html.escape(group_name(r))}</strong>\n'
        f'          <span>{r["status"]} · Prioridad alta · {html.escape(r["responsible"].title() if r["responsible"].isupper() else r["responsible"])}</span>\n'
        '        </div>\n' for r in high) or '\n        <p class="req-sub">Sin hallazgos de prioridad alta.</p>\n'
    text = sub1(r'(<h3>Hallazgos prioritarios</h3>\s*).*?(\s*</section>)',
                lambda m: m.group(1) + cards.strip() + m.group(2), text, "hallazgos")

    text = sub1(r'(<span>Requisitos aplicables</span><strong>)[^<]*(</strong></div>\s*<div class="progress"><span style="width:)[^"]*',
                lambda m: f'{m.group(1)}{app}/{n}{m.group(2)}{round(app / n * 100, 1)}%', text, "progreso aplicables")
    text = sub1(r'(<span>Cumplimiento ponderado</span><strong>)[^<]*(</strong></div>\s*<div class="progress green"><span style="width:)[^"]*',
                lambda m: f'{m.group(1)}{fmt(tot["pct"], False)}%{m.group(2)}{fmt(tot["pct"], False)}%', text, "progreso ponderado")
    text = sub1(r'(<span>Requisitos con brecha</span><strong>)[^<]*(</strong></div>\s*<div class="progress amber"><span style="width:)[^"]*',
                lambda m: f'{m.group(1)}{tot["gap"]}/{app}{m.group(2)}{round(tot["gap"] / app * 100, 1)}%', text, "progreso brechas")
    for phase, name in PHASES:
        k = len([r for r in recs if r["phase"] == phase])
        text = re.sub(rf'(href="{name}\.html">[^<]*<small>)\d+( requisitos)', lambda m: f"{m.group(1)}{k}{m.group(2)}", text, 1)
    text = replace_dates(text, ["14-10-2026", "30-09-2026"], date)
    write("auditor.html", text)


def update_index(recs, per_phase, tot, date):
    text = read("index.html")
    text = re.sub(r"(<small>Requisitos</small><strong>)\d+", lambda m: m.group(1) + str(tot["n"]), text, 1)
    text = re.sub(r"(<small>Cumplen</small><strong>)\d+", lambda m: m.group(1) + str(tot["ok"]), text, 1)
    text = re.sub(r"(<small>Con brecha</small><strong>)\d+(</strong><span>)[^<]*",
                  lambda m: f'{m.group(1)}{tot["gap"]}{m.group(2)}{tot["menor"]} menores · {tot["rel"]} relevantes', text, 1)
    text = re.sub(r"(<small>Cumplimiento global</small><strong>)[^<]*", lambda m: m.group(1) + fmt(tot["pct"], False) + "%", text, 1)
    text = re.sub(r'(<div class="donut-center"><strong>)[^<]*', lambda m: m.group(1) + f'{round(tot["pct"])}%', text, 1)

    # El donut se reparte por cantidad de requisitos (cumple / menor / relevante / no aplica)
    n = tot["n"]
    cuts, acc = [], 0.0
    for v in (tot["ok"], tot["menor"], tot["rel"]):
        acc += v / n * 100
        cuts.append(round(acc, 2))
    text = sub1(r"(\.donut\{[^}]*?background:).*?(;display:grid)",
                lambda m: m.group(1) +
                f"conic-gradient(\n  var(--green) 0 {cuts[0]:.2f}%,\n  var(--amber) {cuts[0]:.2f}% {cuts[1]:.2f}%,\n"
                f"  var(--red) {cuts[1]:.2f}% {cuts[2]:.2f}%,\n  #b9c2cf {cuts[2]:.2f}% 100%\n)" +
                m.group(2), text, "donut")
    for label, key in (("Cumple", "ok"), ("Brecha menor", "menor"), ("Brecha relevante", "rel"), ("No aplica", "na")):
        text = re.sub(rf"(</i>{label}</span><strong>)\d+", lambda m: m.group(1) + str(tot[key]), text, 1)

    for phase, name in PHASES:
        s = per_phase[phase]
        text = re.sub(rf'(<span class="bar-label">{re.escape(phase)}</span><div class="bar-track"><div class="bar-fill" style="width:)[^;]*(;background:[^"]*"></div></div><span class="bar-val">)[^<]*',
                      lambda m: f'{m.group(1)}{s["pct"]:.2f}%{m.group(2)}{s["pct"]:.2f}%', text, 1)
        text = re.sub(rf'(href="{name}\.html">.*?<div class="phase-meta"><span>)\d+( requisitos</span><strong>)[^<]*',
                      lambda m: f'{m.group(1)}{s["n"]}{m.group(2)}{s["pct"]:.1f}%', text, 1)
    text = replace_dates(text, ["14-10-2026", "30-09-2026"], date)
    write("index.html", text)


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    if not args:
        sys.exit(__doc__)
    path = args[0]
    m = re.search(r"(\d{2})-(\d{2})-(\d{4})", os.path.basename(path))
    date = f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else sys.exit("El nombre de la planilla debe incluir la fecha DD-MM-AAAA.")

    recs = load(path)
    if any("final_observation" in r for r in recs):
        previous = re.search(DATA_RE, read("auditor.html"), re.S)
        if previous is None:
            sys.exit("No se encontraron las acciones anteriores; no se actualizaron las paginas.")
        preserve_previous_actions(recs, json.loads(previous.group(0).split("=", 1)[1].strip().removesuffix(";")))
    tot = stats(recs)
    per_phase = {p: stats([r for r in recs if r["phase"] == p]) for p, _ in PHASES}
    print(f"{len(recs)} requisitos | global {tot['pct']}% | " +
          " | ".join(f"{p} {s['pct']}%" for p, s in per_phase.items()))
    if "--dry-run" in sys.argv:
        return
    for phase, name in PHASES:
        update_phase_page(phase, name, [r for r in recs if r["phase"] == phase], date)
    update_auditor(recs, date, tot)
    update_index(recs, per_phase, tot, date)
    print("Paginas actualizadas con fecha de corte", date)


if __name__ == "__main__":
    main()
