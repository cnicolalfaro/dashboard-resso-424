import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook
from pyxlsb import open_workbook


FIELDS = ("nombre", "rut", "cargo", "especialidad", "rol", "estado")


def normalize(value):
    text = unicodedata.normalize("NFKD", str(value or "").strip().lower())
    return "".join(c for c in text if not unicodedata.combining(c))


def rut(value):
    text = re.sub(r"[.\s-]", "", str(value or "")).upper()
    if not re.fullmatch(r"\d{7,8}[\dK]", text):
        raise ValueError("Invalid RUT")
    body, check = text[:-1], text[-1]
    total = sum(int(c) * (2 + i % 6) for i, c in enumerate(reversed(body)))
    expected = 11 - total % 11
    expected = "K" if expected == 10 else "0" if expected == 11 else str(expected)
    if check != expected:
        raise ValueError("Invalid RUT checksum")
    return f"{int(body)}-{check}"


def workbook_rows(path):
    if path.suffix.lower() == ".xlsb":
        with open_workbook(str(path)) as book:
            for name in book.sheets:
                with book.get_sheet(name) as sheet:
                    yield name, ([cell.v for cell in row] for row in sheet.rows())
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in book:
                yield sheet.title, sheet.iter_rows(values_only=True)
        finally:
            book.close()
    else:
        raise ValueError("TARJA must be xlsb, xlsx or xlsm")


def import_tarja(path):
    if not path.is_absolute() or not path.is_file():
        raise ValueError("NAMI_TARJA_PATH must identify an absolute readable workbook")
    if path.stat().st_size > 40 * 1024 * 1024:
        raise ValueError("TARJA exceeds size limit")
    people, terminated = {}, set()
    aliases = {"name": "nombre", "nombres": "nombre", "nombre completo": "nombre"}
    count = 0
    for sheet, rows in workbook_rows(path):
        header = None
        for values in rows:
            count += 1
            if count > 100000:
                raise ValueError("TARJA exceeds row limit")
            if header is None:
                names = [aliases.get(normalize(v), normalize(v)) for v in values]
                if all(field in names for field in FIELDS):
                    header = {field: names.index(field) for field in FIELDS}
                elif count > 100 and not people:
                    raise ValueError("TARJA requires name/rut/cargo/especialidad/rol/estado headers")
                continue
            item = {key: str(values[index] or "").strip() if index < len(values) else ""
                    for key, index in header.items()}
            if not item["rut"]:
                continue
            item["rut"] = rut(item["rut"])
            if "finiquit" in normalize(sheet) or "finiquit" in normalize(item["estado"]):
                terminated.add(item["rut"])
                continue
            if not item["nombre"] or any(len(item[field]) > 200 for field in FIELDS):
                raise ValueError("Invalid TARJA participant")
            if item["rut"] in people and people[item["rut"]] != item:
                raise ValueError("Ambiguous duplicate TARJA RUT")
            people[item["rut"]] = item
    return {key: {f: value[f] for f in FIELDS if f != "estado"}
            for key, value in people.items() if key not in terminated}


def load_catalog(root, tarja):
    source = root / "catalog.json"
    if source.stat().st_size > 1024 * 1024:
        raise ValueError("Catalog exceeds size limit")
    data = json.loads(source.read_text(encoding="utf-8"))
    people = import_tarja(Path(tarja))
    courses = data["courses"]
    if not isinstance(courses, list) or not 1 <= len(courses) <= 100:
        raise ValueError("Configure private RF courses")
    course_ids = set()
    for item in courses:
        if not isinstance(item, dict) or set(item) != {"id", "code", "title", "content"}:
            raise ValueError("Courses require id/code/title/content")
        if any(not isinstance(v, str) or not v.strip() or len(v) > 4000 for v in item.values()):
            raise ValueError("Invalid course")
        if not re.fullmatch(r"[A-Za-z0-9_-]{1,64}", item["id"]) or item["id"] in course_ids:
            raise ValueError("Invalid/duplicate course id")
        if not item["code"].upper().startswith("RF"):
            raise ValueError("Only configured RF courses are supported")
        course_ids.add(item["id"])
    crews, ids = [], set()
    for crew in data["crews"]:
        identifier = str(crew["id"])
        if identifier not in {"1", "2", "3", "4"} or identifier in ids:
            raise ValueError("Only crews 1-4 may be configured")
        ids.add(identifier)
        memberships = [rut(value) for value in crew["ruts"]]
        if len(memberships) != len(set(memberships)) or not memberships:
            raise ValueError("Crew membership must be explicit and unique")
        if any(value not in people for value in memberships):
            raise ValueError("Crew includes unknown or terminated RUT")
        if len(memberships) > 500 or not isinstance(crew["name"], str) or len(crew["name"]) > 100:
            raise ValueError("Invalid crew")
        crews.append({"id": identifier, "name": crew["name"],
                      "participants": [people[value] for value in memberships]})
    if not crews:
        raise ValueError("Configure at least one crew")
    return {"crews": crews, "courses": courses}
