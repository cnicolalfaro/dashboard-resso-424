import datetime
import hashlib
import io
import json
import os
import re
import threading
import uuid
import zipfile
from contextlib import closing
from pathlib import Path

from docx import Document
from docx.table import Table
from flask import abort, jsonify, request, send_file
from openpyxl import Workbook

from .evaluations import has_placeholder
from .catalog import import_tarja
from .planillas import matrix_courses, plan_groups, read_crews, read_matrix, read_shifts, select_pending
from .prepare import (CARGO_PREFIX, OBJECTIVES, SHIFT_FOLDERS, cargo_key, compose_cargo, folder_name, list_cargos,
                      load_sources, merge_pdfs, prepare, resolve_courses, safe_name, unit_label)
from .storage import connect, transaction, write_private
from .templates import FolderTemplates, course_slug
from .word_register import RenderError, check_template


KINDS = {"tarja": {".xlsb", ".xlsx", ".xlsm"}, "crews": {".xlsb", ".xlsx", ".xlsm"}, "matrix": {".xlsx", ".xlsm"}}
MAX_SOURCE = 20 * 1024 * 1024
MAX_TEMPLATE = 10 * 1024 * 1024
# Planning workbooks keep thousands of formatted rows: 3 MB zipped can be ~180 MB once expanded.
MAX_EXPANDED_WORKBOOK = 400 * 1024 * 1024
MAX_EXPANDED_DOCUMENT = 100 * 1024 * 1024
MAX_BATCH = 100
SCHEMA = """
CREATE TABLE IF NOT EXISTS sources (
    id TEXT PRIMARY KEY, kind TEXT NOT NULL CHECK(kind IN ('tarja','crews','matrix')),
    filename TEXT NOT NULL, stored TEXT NOT NULL, sha256 TEXT NOT NULL, size INTEGER NOT NULL,
    uploaded_by TEXT NOT NULL, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
    UNIQUE(kind, sha256)
);
CREATE TRIGGER IF NOT EXISTS immutable_sources_update BEFORE UPDATE ON sources
BEGIN SELECT RAISE(ABORT, 'immutable source history'); END;
CREATE TRIGGER IF NOT EXISTS immutable_sources_delete BEFORE DELETE ON sources
BEGIN SELECT RAISE(ABORT, 'immutable source history'); END;
CREATE TABLE IF NOT EXISTS source_active (
    kind TEXT PRIMARY KEY, source_id TEXT NOT NULL REFERENCES sources(id),
    changed_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
CREATE TABLE IF NOT EXISTS batches (
    id TEXT PRIMARY KEY, params TEXT NOT NULL, status TEXT NOT NULL, total INTEGER NOT NULL,
    done INTEGER NOT NULL DEFAULT 0, records TEXT NOT NULL DEFAULT '[]', errors TEXT NOT NULL DEFAULT '[]',
    created_by TEXT NOT NULL, created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
);
"""


class Missing(Exception):
    pass


def check_workbook(path, expanded=MAX_EXPANDED_WORKBOOK):
    if path.stat().st_size > 40 * 1024 * 1024:
        raise ValueError("El archivo excede el tamaño permitido")
    with zipfile.ZipFile(path) as archive:
        entries = archive.infolist()
        if len(entries) > 5000 or sum(entry.file_size for entry in entries) > expanded:
            raise ValueError("El archivo expandido excede el tamaño permitido")


class Workspace:
    def __init__(self, root, roster_sheet):
        self.root, self.sheet = root, roster_sheet
        self.lock, self.cache = threading.Lock(), (None, None)

    def active(self):
        with closing(connect(self.root)) as db:
            rows = db.execute("""SELECT s.* FROM source_active a JOIN sources s ON s.id=a.source_id""").fetchall()
        return {row["kind"]: dict(row) for row in rows}

    def path(self, row):
        return self.root / "sources" / row["stored"]

    def load(self):
        rows = self.active()
        missing = [kind for kind in KINDS if kind not in rows]
        if missing:
            raise Missing(missing)
        key = tuple(rows[kind]["id"] for kind in KINDS)
        with self.lock:
            if self.cache[0] != key:
                data = load_sources(self.path(rows["tarja"]), self.path(rows["crews"]),
                                    self.path(rows["matrix"]), self.sheet)
                self.cache = (key, data)
            return self.cache[1], rows

    def validate(self, kind, path):
        check_workbook(path)
        if kind == "tarja":
            import_tarja(path, include_known=True, roster_sheet=self.sheet)
            read_shifts(path, self.sheet)
        elif kind == "matrix":
            read_matrix(path)
        else:
            active = self.active().get("tarja")
            if not active:
                raise ValueError("Suba primero la TARJA: las cuadrillas se validan contra ella")
            people, known = import_tarja(self.path(active), include_known=True, roster_sheet=self.sheet)
            read_crews(path, people, known, read_shifts(self.path(active), self.sheet))


def public_source(row):
    return {"id": row["id"], "kind": row["kind"], "filename": row["filename"], "sha256": row["sha256"][:16],
            "size": row["size"], "uploaded_by": row["uploaded_by"], "uploaded_at": row["created_at"]}


def display_name(name):
    return re.sub(r"[^\w .()\-]", "", Path(name or "").name)[:120] or "archivo"


def valid_date(value):
    try:
        if isinstance(value, str) and datetime.date.fromisoformat(value).isoformat() == value:
            return value
    except ValueError:
        pass
    abort(400)


def request_parameters(value):
    codes = value.get("courses")
    if (not isinstance(codes, list) or not 1 <= len(codes) <= 4 or
            any(not isinstance(code, str) or not 2 <= len(code) <= 20 for code in codes)):
        abort(400)
    objectives = value.get("objectives", OBJECTIVES)
    if not isinstance(objectives, str) or not objectives.strip() or len(objectives) > 500:
        abort(400)
    relator = value.get("relator", "")
    if not isinstance(relator, str) or re.search(r"[\x00-\x1f\x7f]", relator):
        abort(400)
    relator = " ".join(relator.split())
    if len(relator) > 100:
        abort(400)
    return codes, valid_date(value.get("date")), objectives.strip(), relator


def register(app, root, secret, protected, username, prefix, roster_sheet="TARJA", publish=lambda: None):
    for name in ("sources", "work", "templates"):
        (root / name).mkdir(mode=0o700, exist_ok=True)
    with closing(connect(root)) as db:
        db.executescript(SCHEMA)
        db.execute("UPDATE batches SET status='failed' WHERE status='running'")
    workspace = Workspace(root, roster_sheet)
    templates, work = root / "templates", root / "work"
    stock = FolderTemplates(templates)
    upload_lock, batch_lock, template_lock = threading.Lock(), threading.BoundedSemaphore(1), threading.Lock()

    def need_sources():
        try:
            data, rows = workspace.load()
            return data, rows, None
        except Missing as error:
            return None, None, (jsonify(error="Faltan planillas activas", missing=error.args[0]), 409)

    def generate(unit, codes, date, objectives, relator=""):
        data, rows = workspace.load()
        result = prepare(data, unit, codes, date, templates, work, secret, objectives, relator)
        manifest, register_pdf, evaluations = result["manifest"], result["register"], result["evaluations"]
        manifest["sources"] = {kind: {"id": row["id"], "filename": row["filename"], "sha256": row["sha256"]}
                               for kind, row in rows.items()}
        path = root / "records" / f'{manifest["id"]}.pdf'
        written = [path]
        write_private(path, register_pdf)
        try:
            if evaluations:
                extra = root / "records" / f'{manifest["id"]}-evaluaciones.pdf'
                write_private(extra, evaluations)
                written.append(extra)
                manifest["evaluations_pdf"] = extra.name
            with transaction(root) as db:
                db.execute("INSERT INTO records(id,manifest,pdf) VALUES (?,?,?)",
                           (manifest["id"], json.dumps(manifest, ensure_ascii=False), path.name))
        except BaseException:
            for item in written:
                item.unlink(missing_ok=True)
            raise
        publish()  # carpeta compartida: el registro queda disponible para las otras instalaciones
        return manifest

    def record_summary(identifier, manifest, created):
        return {"id": identifier, "date": manifest.get("date"), "created_at": created,
                "crew": manifest["crew"]["name"], "foreman": manifest["crew"].get("foreman", ""),
                "relator": manifest.get("relator", ""),
                "shift": folder_name(manifest["crew"]),
                "pages": len(manifest["pages"]), "participants": len(manifest["participants"]),
                "courses": [course["code"] for course in manifest["courses"]],
                "evaluations": bool(manifest.get("evaluations_pdf")),
                "summary": manifest.get("summary", {}),
                "pdf_url": f"{prefix}/records/{identifier}/pdf",
                "bundle_url": f"{prefix}/records/{identifier}/bundle"}

    @app.get(prefix + "/sources")
    @protected("read")
    def sources():
        with closing(connect(root)) as db:
            history = [public_source(row) for row in db.execute(
                "SELECT * FROM sources ORDER BY created_at DESC, rowid DESC")]
        active = workspace.active()
        return jsonify(active={kind: public_source(row) for kind, row in active.items()}, history=history,
                       templates={"registro": stock.register() is not None})

    def validate_template(kind, path):
        check_workbook(path, MAX_EXPANDED_DOCUMENT)
        document = Document(str(path))
        if kind == "evaluation":
            if not has_placeholder(document):
                raise ValueError("La evaluación debe tener [NOMBRETRABAJADOR] en el pie de página")
        elif len(document.tables) < 2 or len(document.tables[1].rows) < 2:
            raise ValueError("El registro debe tener la tabla de datos y la tabla de participantes")
        else:
            try:
                check_template(Table(document.tables[0]._tbl, document))
            except (RenderError, IndexError):
                raise ValueError("El registro no tiene la estructura esperada (FECHA, OBJETIVOS, MATERIA, PARTICIPANTES)")

    @app.post(prefix + "/templates")
    @protected("create")
    def upload_template():
        if request.mimetype != "multipart/form-data" or len(request.files) != 1 or "file" not in request.files:
            abort(400)
        kind = request.form.get("kind")
        if kind not in {"evaluation", "register"}:
            abort(400)
        upload = request.files["file"]
        if Path(display_name(upload.filename)).suffix.lower() != ".docx":
            abort(400)
        data = upload.stream.read(MAX_TEMPLATE + 1)
        if not data or len(data) > MAX_TEMPLATE or not data.startswith(b"PK\x03\x04"):
            abort(400)
        if kind == "evaluation":
            code = request.form.get("course", "")
            data_sources, _, problem = need_sources()
            if problem:
                return problem
            course = next((item for item in matrix_courses(data_sources["matrix"]["legend"])
                           if item["code"].upper() == code.strip().upper()), None)
            if course is None:
                return jsonify(error="El curso no existe en la matriz activa"), 422
            target = templates / f"eval-{course_slug(course)}.docx"
            label = course["code"]
        else:
            target, label = templates / "registro.docx", "Registro"
        if not template_lock.acquire(blocking=False):
            abort(429)
        temporary = work / f"{uuid.uuid4().hex}.docx"
        try:
            write_private(temporary, data)
            try:
                validate_template(kind, temporary)
            except ValueError as error:
                return jsonify(error=f"No se puede usar el archivo: {str(error)[:180]}"), 422
            except Exception:
                return jsonify(error="El archivo no es un documento Word válido"), 422
            replaced = target.is_file()
            if replaced:
                archive = templates / "anteriores"
                archive.mkdir(mode=0o700, exist_ok=True)
                stamp = datetime.datetime.now(datetime.timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")
                backup = archive / f"{target.stem}-{stamp}.docx"
                os.replace(target, backup)
                os.chmod(backup, 0o600)
            os.replace(temporary, target)
            return jsonify(kind=kind, course=label, replaced=replaced, size=len(data),
                           sha256=hashlib.sha256(data).hexdigest()[:16]), 201
        finally:
            temporary.unlink(missing_ok=True)
            template_lock.release()

    @app.post(prefix + "/sources")
    @protected("create")
    def upload_source():
        if request.mimetype != "multipart/form-data" or len(request.files) != 1 or "file" not in request.files:
            abort(400)
        kind = request.form.get("kind")
        upload = request.files["file"]
        name = display_name(upload.filename)
        suffix = Path(name).suffix.lower()
        if kind not in KINDS or suffix not in KINDS[kind]:
            abort(400)
        data = upload.stream.read(MAX_SOURCE + 1)
        if not data or len(data) > MAX_SOURCE or not data.startswith(b"PK\x03\x04"):
            abort(400)
        digest = hashlib.sha256(data).hexdigest()
        if not upload_lock.acquire(blocking=False):
            abort(429)
        path, duplicate = None, False
        try:
            with closing(connect(root)) as db:
                existing = db.execute("SELECT * FROM sources WHERE kind=? AND sha256=?", (kind, digest)).fetchone()
            if existing:
                row, duplicate = dict(existing), True
            else:
                row = {"id": str(uuid.uuid4())}
                row["stored"] = f'{row["id"]}{suffix}'
                path = root / "sources" / row["stored"]
                write_private(path, data)
                try:
                    workspace.validate(kind, path)
                except ValueError as error:
                    path.unlink(missing_ok=True)
                    return jsonify(error=f"No se pudo leer la planilla: {str(error)[:160]}"), 422
                except Exception:
                    path.unlink(missing_ok=True)
                    return jsonify(error="El archivo no tiene un formato de planilla válido"), 422
            with transaction(root) as db:
                if not duplicate:
                    db.execute("INSERT INTO sources(id,kind,filename,stored,sha256,size,uploaded_by) VALUES (?,?,?,?,?,?,?)",
                               (row["id"], kind, name, row["stored"], digest, len(data), username))
                db.execute("""INSERT INTO source_active(kind,source_id) VALUES (?,?)
                              ON CONFLICT(kind) DO UPDATE SET source_id=excluded.source_id, changed_at=CURRENT_TIMESTAMP""",
                           (kind, row["id"]))
            return jsonify(kind=kind, duplicate=duplicate, source=public_source(workspace.active()[kind])), 201
        except BaseException:
            if path and not duplicate:
                path.unlink(missing_ok=True)
            raise
        finally:
            upload_lock.release()

    @app.get(prefix + "/plan")
    @protected("read")
    def plan():
        data, rows, problem = need_sources()
        if problem:
            return problem
        codes = [code for code in request.args.get("courses", "").split(",") if code][:4]
        try:
            chosen = resolve_courses(data["matrix"], codes) if codes else []
        except ValueError:
            abort(400)
        available = stock.available(matrix_courses(data["matrix"]["legend"]))
        crews, items = [], []
        for crew in data["crews"]["crews"]:
            pending = select_pending(crew, chosen, data["matrix"])[0] if chosen else []
            items.append({"id": crew["id"], "shift": crew.get("shift"), "pending": len(pending)})
            crews.append({"id": crew["id"], "name": crew["name"], "foreman": crew.get("foreman", ""),
                          "foreman_found": crew.get("foreman_found", False),
                          "shift": SHIFT_FOLDERS[crew.get("shift")], "shift_mixed": crew.get("shift_mixed", False),
                          "members": len(crew["participants"]), "excluded": len(crew["excluded"]),
                          "pending": len(pending) if chosen else None})
        wanted = {item for item in request.args.get("crews", "").split(",") if item}
        group_items = [item for item in items if item["id"] in wanted] if wanted else items
        cargos = []
        for item in list_cargos(data):
            pending = len(compose_cargo(data, item["key"], chosen)[1]) if chosen else None
            cargos.append({**item, "pending": pending})
        return jsonify(ready=True, sources={kind: public_source(row) for kind, row in rows.items()},
                       courses=[{"code": course["code"], "title": course["title"],
                                 "evaluation": course["id"] in available}
                                for course in matrix_courses(data["matrix"]["legend"])],
                       crews=crews, cargos=cargos, warnings=data["crews"]["warnings"],
                       groups=plan_groups(group_items) if chosen and request.args.get("group") == "1" else None,
                       templates={"registro": stock.register() is not None})

    @app.post(prefix + "/records/word")
    @protected("create")
    def word_record():
        value = request.get_json(silent=True) if request.mimetype == "application/json" else None
        if not isinstance(value, dict) or not isinstance(value.get("crew_id"), (str, int)) or isinstance(value.get("crew_id"), bool):
            abort(400)
        codes, date, objectives, relator = request_parameters(value)
        try:
            manifest = generate(str(value["crew_id"]), codes, date, objectives, relator)
        except Missing as error:
            return jsonify(error="Faltan planillas activas", missing=error.args[0]), 409
        except (ValueError, RenderError) as error:
            return jsonify(error=str(error)[:200]), 422
        created = datetime.datetime.now(datetime.timezone.utc).isoformat()
        return jsonify(record_summary(manifest["id"], manifest, created)), 201

    @app.get(prefix + "/records")
    @protected("read")
    def list_records():
        with closing(connect(root)) as db:
            rows = db.execute("SELECT id,manifest,created_at FROM records ORDER BY created_at DESC, rowid DESC LIMIT 300").fetchall()
        return jsonify(records=[record_summary(row["id"], json.loads(row["manifest"]), row["created_at"]) for row in rows])

    def bundle_bytes(identifier):
        with closing(connect(root)) as db:
            row = db.execute("SELECT manifest,pdf FROM records WHERE id=?", (identifier,)).fetchone()
        if not row:
            abort(404)
        manifest = json.loads(row["manifest"])
        if not (root / "records" / row["pdf"]).is_file():
            abort(404)  # registro importado de otra instalación sin su PDF en la carpeta compartida
        parts = [(root / "records" / row["pdf"]).read_bytes()]
        if manifest.get("evaluations_pdf") and (root / "records" / manifest["evaluations_pdf"]).is_file():
            parts.append((root / "records" / manifest["evaluations_pdf"]).read_bytes())
        return merge_pdfs(*parts), manifest

    @app.get(prefix + "/records/<identifier>/bundle")
    @protected("read")
    def record_bundle(identifier):
        try:
            uuid.UUID(identifier)
        except ValueError:
            abort(404)
        data, manifest = bundle_bytes(identifier)
        return send_file(io.BytesIO(data), mimetype="application/pdf", as_attachment=True,
                         download_name=f'{safe_name(manifest["crew"]["name"])}.pdf', conditional=False)

    def run_batch(batch_id, params):
        records, errors, status = [], [], "failed"
        try:
            for index, unit in enumerate(params["units"], 1):
                try:
                    manifest = generate(unit, params["courses"], params["date"], params["objectives"],
                                        params.get("relator", ""))
                    records.append(manifest["id"])
                except (ValueError, RenderError) as error:
                    errors.append({"crew_id": unit_label(unit), "error": str(error)[:160]})
                except Exception:
                    errors.append({"crew_id": unit_label(unit), "error": "No fue posible generar el documento"})
                with transaction(root) as db:
                    db.execute("UPDATE batches SET done=?, records=?, errors=? WHERE id=?",
                               (index, json.dumps(records), json.dumps(errors), batch_id))
            status = "done"
        finally:
            try:
                with transaction(root) as db:
                    db.execute("UPDATE batches SET status=? WHERE id=?", (status, batch_id))
            finally:
                batch_lock.release()

    @app.post(prefix + "/batches")
    @protected("create")
    def start_batch():
        value = request.get_json(silent=True) if request.mimetype == "application/json" else None
        if not isinstance(value, dict):
            abort(400)
        crew_ids, cargos = value.get("crew_ids"), value.get("cargos")
        if (crew_ids is None) == (cargos is None):
            abort(400)
        if cargos is not None:
            if (not isinstance(cargos, list) or not 1 <= len(cargos) <= MAX_BATCH or
                    any(not isinstance(item, str) or len(item) > 100 for item in cargos)):
                abort(400)
            crew_ids = [CARGO_PREFIX + cargo_key({"cargo": item}) for item in cargos]
        elif (not isinstance(crew_ids, list) or not 1 <= len(crew_ids) <= MAX_BATCH or
                any(not isinstance(item, (str, int)) or isinstance(item, bool) for item in crew_ids)):
            abort(400)
        if len({str(item) for item in crew_ids}) != len(crew_ids):
            abort(400)
        group_small = value.get("group_small", False)
        if not isinstance(group_small, bool) or (group_small and cargos is not None):
            abort(400)
        codes, date, objectives, relator = request_parameters(value)
        data, _, problem = need_sources()
        if problem:
            return problem
        ids = [str(item) for item in crew_ids]
        if cargos is not None:
            known = {item["key"] for item in list_cargos(data)}
            if any(item[len(CARGO_PREFIX):] not in known for item in ids):
                abort(400)
        if group_small:
            try:
                courses = resolve_courses(data["matrix"], codes)
            except ValueError:
                abort(400)
            by_id = {crew["id"]: crew for crew in data["crews"]["crews"]}
            units = plan_groups([{"id": item, "shift": by_id[item].get("shift") if item in by_id else None,
                                  "pending": len(select_pending(by_id[item], courses, data["matrix"])[0])
                                  if item in by_id else 0} for item in ids])
        else:
            units = [[item] for item in ids]
        if not batch_lock.acquire(blocking=False):
            abort(429)
        batch_id = str(uuid.uuid4())
        params = {"units": units, "courses": codes, "date": date, "objectives": objectives, "relator": relator}
        try:
            with transaction(root) as db:
                db.execute("INSERT INTO batches(id,params,status,total,created_by) VALUES (?,?,?,?,?)",
                           (batch_id, json.dumps(params), "running", len(units), username))
            threading.Thread(target=run_batch, args=(batch_id, params), daemon=True).start()
        except BaseException:
            batch_lock.release()
            raise
        return jsonify(id=batch_id, status="running", total=len(units), units=units), 202

    def batch_row(identifier):
        try:
            uuid.UUID(identifier)
        except ValueError:
            abort(404)
        with closing(connect(root)) as db:
            row = db.execute("SELECT * FROM batches WHERE id=?", (identifier,)).fetchone()
        if not row:
            abort(404)
        return row

    @app.get(prefix + "/batches/<identifier>")
    @protected("read")
    def batch_status(identifier):
        row = batch_row(identifier)
        return jsonify(id=row["id"], status=row["status"], total=row["total"], done=row["done"],
                       errors=json.loads(row["errors"]), records=json.loads(row["records"]),
                       zip_url=f"{prefix}/batches/{row['id']}/zip")

    @app.get(prefix + "/batches/<identifier>/zip")
    @protected("read")
    def batch_zip(identifier):
        row = batch_row(identifier)
        record_ids = json.loads(row["records"])
        if not record_ids:
            abort(404)
        stream = io.BytesIO()
        used = set()
        with zipfile.ZipFile(stream, "w", zipfile.ZIP_STORED) as archive:
            for record_id in record_ids:
                data, manifest = bundle_bytes(record_id)
                folder = folder_name(manifest["crew"])
                base = safe_name(manifest["crew"]["name"])
                name, counter = f"{folder}/{base}.pdf", 1
                while name in used:
                    counter += 1
                    name = f"{folder}/{base} ({counter}).pdf"
                used.add(name)
                archive.writestr(name, data)
        stream.seek(0)
        return send_file(stream, mimetype="application/zip", as_attachment=True,
                         download_name=f"NAMI-{json.loads(row['params'])['date']}.zip", conditional=False)

    @app.get(prefix + "/results/export.xlsx")
    @protected("read")
    def export_results():
        with closing(connect(root)) as db:
            results = db.execute("SELECT * FROM results ORDER BY id").fetchall()
            manifests = {row["id"]: json.loads(row["manifest"]) for row in db.execute("SELECT id,manifest FROM records")}
        book = Workbook()
        sheet = book.active
        sheet.title = "Difusión confirmada"
        headers = ["Fecha del registro", "Cuadrilla", "Capataz", "Turno", "RUT", "Nombre", "Curso", "Título del curso",
                   "Difusión", "Confirmado por", "Fecha de confirmación", "ID del registro", "Relator"]
        rows = [headers]
        for item in results:
            manifest = manifests.get(item["record_id"], {})
            person = next((p for p in manifest.get("participants", []) if p["rut"] == item["rut"]), {})
            course = next((c for c in manifest.get("courses", []) if c["id"] == item["course_id"]), {})
            crew = manifest.get("crew", {})
            rows.append([manifest.get("date", ""), person.get("crew_name") or crew.get("name", ""),
                         person.get("foreman") or crew.get("foreman", ""),
                         folder_name(crew), item["rut"], person.get("nombre", ""),
                         course.get("code", item["course_id"]), course.get("title", ""),
                         "Firmó" if item["status"] == "FIRMO" else "No firmó", item["confirmed_by"],
                         item["created_at"], item["record_id"], manifest.get("relator", "")])
        for number, values in enumerate(rows, 1):
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row=number, column=column, value=value)
                if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
                    cell.data_type = "s"
        note = book.create_sheet("Leyenda")
        note.append(["Firmar acredita la difusión; NO equivale a aprobar una evaluación."])
        note.append(["Este archivo no modifica la matriz de capacitación ni Avance RESSO."])
        stream = io.BytesIO()
        book.save(stream)
        stream.seek(0)
        return send_file(stream, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                         as_attachment=True, download_name="difusion-confirmada.xlsx", conditional=False)

    return workspace
