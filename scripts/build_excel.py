#!/usr/bin/env python3
"""
Reconstruye las hojas "Matriz" y "Resumen por curso" del Excel RESSO a partir
de un snapshot new_data.json, preservando intacta la hoja
"Lineas de Mando - Controles" del libro existente (no se toca).

Uso:
    python3 build_excel.py <NEW_DATA_JSON> <PREV_XLSX> <OUT_XLSX>
"""
import sys, json
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def main():
    if len(sys.argv) != 4:
        print("RESULT_ERROR uso: build_excel.py <NEW_DATA_JSON> <PREV_XLSX> <OUT_XLSX>")
        sys.exit(1)
    new_data_path, prev_xlsx_path, out_xlsx_path = sys.argv[1:4]

    data = json.load(open(new_data_path, encoding='utf-8'))
    docs = data['documentos']
    personas = data['personas']
    n_docs = len(docs)

    FONT = "Arial"
    header_font = Font(name=FONT, bold=True, size=9, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="B8420F")
    title_font = Font(name=FONT, bold=True, size=14)
    sub_font = Font(name=FONT, size=9, italic=True, color="5C5440")
    base_font = Font(name=FONT, size=10)
    thin = Side(style="thin", color="D9D9D9")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    try:
        wb = openpyxl.load_workbook(prev_xlsx_path)
    except Exception as e:
        print(f"RESULT_ERROR no se pudo abrir el Excel existente {prev_xlsx_path}: {e}")
        sys.exit(1)

    preserved_sheets = [s for s in wb.sheetnames if s not in ('Matriz', 'Resumen por curso')]
    if 'Líneas de Mando - Controles' not in preserved_sheets:
        print(f"RESULT_ERROR la hoja 'Líneas de Mando - Controles' no existe en {prev_xlsx_path} "
              f"(hojas encontradas: {wb.sheetnames}) - abortando para no perderla")
        sys.exit(1)

    for name in ('Matriz', 'Resumen por curso'):
        if name in wb.sheetnames:
            del wb[name]

    # ---------------- Sheet 1: Matriz (se inserta primera) ----------------
    ws = wb.create_sheet("Matriz", 0)

    ws["A1"] = "Matriz de Capacitación RESSO — Dotación Activa"
    ws["A1"].font = title_font
    ws["A2"] = (f"Generado: {data['generado']}  ·  Dotación activa: {data['total_dotacion_activa']}"
                f"  ·  Avance promedio: {data['avance_promedio_pct']}%")
    ws["A2"].font = sub_font

    HEADER_ROW = 4
    FIXED_COLS = ["RUT", "Nombre", "Especialidad", "Categoría", "Turno", "Avance %"]
    for i, h in enumerate(FIXED_COLS, start=1):
        c = ws.cell(row=HEADER_ROW, column=i, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border

    first_doc_col = len(FIXED_COLS) + 1
    for j, d in enumerate(docs):
        col = first_doc_col + j
        c = ws.cell(row=HEADER_ROW, column=col, value=d['titulo'])
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True, text_rotation=90)
        c.border = border
        ws.column_dimensions[get_column_letter(col)].width = 4.2

    ws.row_dimensions[HEADER_ROW].height = 160

    last_doc_col = first_doc_col + n_docs - 1
    first_data_row = HEADER_ROW + 1
    last_data_row = first_data_row + len(personas) - 1

    aprob_fill = PatternFill("solid", fgColor="D9EAD3")
    pend_fill = PatternFill("solid", fgColor="F4CCCC")

    for i, p in enumerate(personas):
        r = first_data_row + i
        ws.cell(row=r, column=1, value=p['rut']).font = base_font
        ws.cell(row=r, column=2, value=p['nombre']).font = base_font
        ws.cell(row=r, column=3, value=p['especialidad']).font = base_font
        ws.cell(row=r, column=4, value=p['categoria']).font = base_font
        ws.cell(row=r, column=5, value=p.get('turno', '')).font = base_font
        for c in range(1, 6):
            ws.cell(row=r, column=c).border = border

        for j, bit in enumerate(p['aprob']):
            col = first_doc_col + j
            val = "Aprobado" if bit == '1' else "Pendiente"
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = Font(name=FONT, size=8)
            cell.alignment = Alignment(horizontal="center")
            cell.fill = aprob_fill if bit == '1' else pend_fill
            cell.border = border

        rng = f"{get_column_letter(first_doc_col)}{r}:{get_column_letter(last_doc_col)}{r}"
        fcell = ws.cell(row=r, column=6, value=f'=COUNTIF({rng},"Aprobado")/{n_docs}')
        fcell.number_format = "0.0%"
        fcell.font = Font(name=FONT, size=10, bold=True)
        fcell.border = border

    ws.freeze_panes = ws.cell(row=first_data_row, column=first_doc_col)
    for i, w in enumerate([14, 34, 24, 16, 10, 10], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # ---------------- Sheet 2: Resumen por curso ----------------
    ws2 = wb.create_sheet("Resumen por curso", 1)

    ws2["A1"] = "Resumen por curso — RESSO"
    ws2["A1"].font = title_font
    ws2["A2"] = f"Generado: {data['generado']}  ·  Dotación activa: {data['total_dotacion_activa']}"
    ws2["A2"].font = sub_font

    hdrs = ["#", "Curso", "Modo", "Fuentes", "Aprobados", "Avance %"]
    HR = 4
    for i, h in enumerate(hdrs, start=1):
        c = ws2.cell(row=HR, column=i, value=h)
        c.font = header_font
        c.fill = header_fill
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = border
    ws2.row_dimensions[HR].height = 20

    total_dot = data['total_dotacion_activa']
    for j, d in enumerate(docs):
        r = HR + 1 + j
        col_letter = get_column_letter(first_doc_col + j)
        ws2.cell(row=r, column=1, value=j + 1).font = base_font
        tc = ws2.cell(row=r, column=2, value=d['titulo'])
        tc.font = base_font
        tc.alignment = Alignment(wrap_text=True)
        ws2.cell(row=r, column=3, value=d['modo']).font = base_font
        ws2.cell(row=r, column=4, value=", ".join(d.get('fuentes') or []) or "—").font = base_font
        acell = ws2.cell(row=r, column=5,
                          value=f'=COUNTIF(Matriz!{col_letter}{first_data_row}:{col_letter}{last_data_row},"Aprobado")')
        acell.font = base_font
        pcell = ws2.cell(row=r, column=6, value=f"=E{r}/{total_dot}")
        pcell.number_format = "0.0%"
        pcell.font = Font(name=FONT, bold=True)
        for c in range(1, 7):
            ws2.cell(row=r, column=c).border = border

    for i, w in enumerate([5, 55, 18, 34, 12, 11], start=1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    ws2.freeze_panes = "A5"

    wb.save(out_xlsx_path)
    print(f"RESULT_OK filas_matriz={last_data_row} cols_matriz={last_doc_col} hojas={wb.sheetnames}")


if __name__ == '__main__':
    main()
