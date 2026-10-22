#!/usr/bin/env python3
"""Genera pages/evidencias-v11 sin modificar RESSO v10.

Uso: python scripts/actualizar_evidencias_v11.py <planilla.xlsx> --date DD-MM-AAAA
"""
import argparse
import html
import os
import re
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from urllib.parse import quote

import openpyxl

import actualizar_evidencias as shared

ROOT = Path(__file__).resolve().parents[1]
DEST = ROOT / "pages" / "evidencias-v11"
SP_ROOT = (
    "/sites/ICSK-HSEC/Documentos compartidos/05 - Respaldo HSEC faenas/"
    "424 - CC 101/0.- RESSO Ver. 10/05 RESSO V11"
)
SP_LIBRARY = (
    "https://empresassk.sharepoint.com/sites/ICSK-HSEC/"
    "Documentos%20compartidos/Forms/AllItems.aspx"
)
SP_URL = SP_LIBRARY + "?id=" + quote(SP_ROOT, safe="")


def load(path):
    workbook = openpyxl.load_workbook(path, data_only=True)
    ws = workbook.worksheets[0]
    if not str(ws["A29"].value).startswith("N"):
        raise ValueError("No se encontro la cabecera de requisitos en A29.")
    records = []
    for cells in ws.iter_rows(min_row=30):
        r = [c.value for c in cells]
        if not isinstance(r[0], int):
            continue
        weight = Decimal(str(r[4]))
        level = r[5]
        if level is None or level == "" or level == "N/A":
            status = "No aplica" if level == "N/A" else "Pendiente"
            result = gap = None
            priority = "—"
        else:
            if not isinstance(level, (int, float)) or not 0 <= level <= 1:
                raise ValueError(f"Nivel invalido en requisito {r[0]}: {level!r}")
            result = weight * Decimal(str(level))
            gap = weight - result
            status = ("Cumple" if level >= 1 else "Brecha menor" if level >= .75
                      else "Brecha relevante" if level >= .5 else "Crítico")
            rounded_gap = gap.quantize(Decimal(".0001"), rounding=ROUND_HALF_UP)
            priority = "Alta" if rounded_gap >= Decimal(".01") else "Media" if rounded_gap > 0 else "Baja"
        link = cells[15].hyperlink
        evidence = link.target if link else shared.clean(r[15])
        if evidence and not evidence.startswith("https://"):
            raise ValueError(f"Enlace invalido en requisito {r[0]}: {evidence!r}")
        records.append(dict(
            id=r[0], phase=r[1], section=r[2], question=r[3], weight=float(weight),
            level=None if status == "Pendiente" else level,
            result=None if result is None else float(result),
            gap=None if gap is None else float(gap), status=status, priority=priority,
            criterion=shared.clean(r[10]) or "", observation=shared.clean(r[12]) or shared.clean(r[11]) or "",
            action=shared.clean(r[13]) or "", responsible=shared.clean(r[14]) or "",
            evidence_ref=evidence,
        ))
    if not records or len({r["id"] for r in records}) != len(records):
        raise ValueError("La planilla debe contener requisitos con identificadores unicos.")
    if any(r["phase"] not in dict(shared.PHASES) for r in records):
        raise ValueError("La planilla contiene una fase desconocida.")
    if "Cambios v10 → v11" not in workbook.sheetnames:
        raise ValueError("Falta la hoja Cambios v10 → v11 para clasificar las preguntas.")
    changes = {}
    for row in workbook["Cambios v10 → v11"].iter_rows(values_only=True):
        if len(row) >= 9 and isinstance(row[0], int) and isinstance(row[3], (int, float)):
            if row[0] in changes:
                raise ValueError(f"Clasificacion duplicada del requisito {row[0]}.")
            changes[row[0]] = row
    for record in records:
        if record["id"] not in changes:
            raise ValueError(f"Falta la comparacion del requisito {record['id']}.")
        change = changes[record["id"]]
        record.update(classify_change(change[7], record["level"]))
        record["change_type"] = change[7]
        record["change_detail"] = shared.clean(change[8]) or ""
    return records


def classify_change(change_type, level):
    allowed = ("Nueva", "Modificada", "Modificada / reubicada", "Sin cambio", "Redacción", "Ponderación")
    if change_type not in allowed:
        raise ValueError(f"Tipo de cambio desconocido: {change_type!r}")
    if level == "N/A":
        category, label = "na", "No aplica"
    elif change_type == "Nueva":
        category, label = "new", "Totalmente nueva en v11"
    elif change_type.startswith("Modificada"):
        category, label = "review", "Modificada: revisar y complementar respaldo"
    elif level == 1:
        category, label = "complete", "Equivalente a v10: completa según planilla"
    else:
        category, label = "review", "Equivalente a v10: falta completar evaluación o respaldo"
    return dict(migration=category, migration_label=label)


def migration_ui(text, records, auditor=False):
    css = """
.migration-new{--migration-bg:#fce7ef}
.migration-review{--migration-bg:#fff4c7}
.migration-complete{--migration-bg:#e4f4e8}
.migration-na{--migration-bg:#eef1f5}
.req[class*="migration-"],tbody tr[class*="migration-"]>td,
tbody tr[class*="migration-"]:hover>td{background:var(--migration-bg)}
.migration-badge{display:inline-block;background:var(--migration-bg);color:#334155;
border:1px solid #cbd5e1;border-radius:8px;padding:5px 8px;font-size:11px;margin-top:7px;line-height:1.5}
.migration-legend{display:flex;flex-wrap:wrap;gap:8px;margin:12px 0}
.migration-detail{background:var(--migration-bg)}
"""
    legend = '<aside aria-label="Adaptación de v10 a v11"><div class="migration-legend">'
    for category, label in (("new", "Nueva"), ("review", "Revisar / complementar"),
                            ("complete", "Equivalente y completa"), ("na", "No aplica")):
        count = sum(r["migration"] == category for r in records)
        legend += f'<span class="migration-badge migration-{category}">{label}: {count}</span>'
    legend += ('</div><p>Colores de adaptación a v11, no de cumplimiento. Amarillo: pregunta modificada '
               'o equivalente con brechas. Verde: equivalente (incluye solo redacción/ponderación) '
               'y nivel 100 % en la planilla; no implica una nueva validación de evidencia.</p></aside>')
    text = text.replace("</style>", css + "</style>", 1)
    text = text.replace('<main class="page">', '<main class="page">' + legend, 1)
    if 'id="mQuestion"' not in text:
        return text
    helper = """
function migrationBadge(r) {
  const badge = document.createElement("span");
  badge.className = `migration-badge migration-${r.migration}`;
  badge.textContent = r.migration_label;
  return badge.outerHTML;
}
"""
    text = text.replace("const el = id =>", helper + "\nconst el = id =>", 1)
    if auditor:
        text = text.replace("      <tr>", '      <tr class="migration-${r.migration}">')
        text = text.replace("<td>${detailLine}</td>", "<td>${detailLine}${migrationBadge(r)}</td>")
    else:
        text = text.replace('<article class="req">', '<article class="req migration-${r.migration}">')
        text = text.replace('${r.question}</div></div>', '${r.question}</div>${migrationBadge(r)}</div>')
    text = text.replace('<div class="detail">', '<div class="detail">'
                        '<div id="mMigrationCard" class="detail-card">'
                        '<h4>Adaptación de v10 a v11</h4><p id="mMigration"></p></div>', 1)
    text = text.replace('  el("mQuestion").textContent =',
                        '  el("mMigrationCard").className = `detail-card migration-detail migration-${r.migration}`;\n'
                        '  el("mMigration").textContent = `${r.migration_label}. ${r.change_type}: ${r.change_detail}`;\n'
                        '  el("mQuestion").textContent =', 1)
    return text


def stats(records):
    evaluated = [r for r in records if r["status"] != "Pendiente"]
    s = shared.stats(evaluated)
    s.update(n=len(records), pending=len(records) - len(evaluated))
    s["critical"] = sum(r["status"] == "Crítico" for r in records)
    s["gap"] += s["critical"]
    return s


def link_code():
    return f"""
const SP_ROOT_FOLDER = {shared.js_data(SP_ROOT, False)};
const SP_LIBRARY_URL = {shared.js_data(SP_LIBRARY, False)};
function specificSharePointLink(r) {{
  if (r.evidence_ref || r.link) return r.evidence_ref || r.link;
  return `${{SP_LIBRARY_URL}}?id=${{encodeURIComponent(`${{SP_ROOT_FOLDER}}/${{String(r.id).padStart(2, "0")}} versión 11`)}}`;
}}
function sharePointReadablePath(r) {{
  return r.evidence_ref || r.link ? "Evidencia indicada en la planilla" : `${{String(r.id).padStart(2, "0")}} versión 11`;
}}
"""


def template(name, date):
    text = (ROOT / "pages" / "evidencias" / f"{name}.html").read_text(encoding="utf-8")
    text = text.replace("RESSO V10", "RESSO V11")
    text = re.sub(r'href="https://empresassk\.sharepoint\.com/[^"]+"',
                  lambda m: f'href="{html.escape(SP_URL, quote=True)}"', text)
    text = shared.replace_dates(text, ["21-10-2026"], date) if name != "index" else text.replace("21-10-2026", date)
    text = text.replace("</style>", ".chip.gray{background:#edf1f6;color:#64748b}</style>")
    return text


def select_options(text, id_, label, values):
    options = f'<option value="all">{label}</option>' + "".join(
        f'<option value="{html.escape(v, quote=True)}">{html.escape(v)}</option>' for v in values)
    return re.sub(rf'(<select id="{id_}">).*?(</select>)',
                  lambda m: m[1] + options + m[2], text, flags=re.S)


def phase_page(phase, name, records, date):
    # Una misma plantilla evita divergencias de esquema entre las cuatro fases.
    text = template("plan", date)
    text = text.replace("A. PLAN", phase).replace("PLAN", phase.split(". ")[1])
    text = text.replace("Planificación Preventiva Integral · Programa de trabajo, evaluación de riesgos y cumplimiento regulatorio.",
                        "Revisión documental RESSO v11 · Requisitos y resultados por fase.")
    rows = [shared.phase_record(r) for r in records]
    for row, source in zip(rows, records):
        row["level"] = source["level"] * 100 if isinstance(source["level"], (int, float)) else source["level"]
        for key in ("migration", "migration_label", "change_type", "change_detail"):
            row[key] = source[key]
    text = re.sub(shared.DATA_RE, lambda m: "const DATA = " + shared.js_data(rows, False) + ";", text, flags=re.S)
    s = stats(records)
    text = shared.update_phase_kpis(text, s, name)
    text = shared.update_hero(text, s, date)
    text = re.sub(r'<section class="groups">.*?</section>',
                  lambda m: '<section class="groups">' + group_cards(records) + '</section>', text, flags=re.S)
    text = shared.filters(text, records)
    text = re.sub(r'const SP_LIBRARY_URL =.*?(?=const el =)', lambda m: link_code(), text, flags=re.S)
    text = re.sub(r'function statusClass\(status\) \{.*?\n\}', '''function statusClass(status) {
  if (status === "Cumple") return "green";
  if (status === "Brecha menor") return "amber";
  if (status === "Pendiente" || status === "No aplica") return "gray";
  return "red";
}
function levelText(level) {
  return typeof level === "number" ? `${level}%` : level === "N/A" ? "N/A" : "Pendiente";
}''', text, count=1, flags=re.S)
    text = text.replace("${r.level}%", "${levelText(r.level)}")
    text = text.replace("${r.group} → Pregunta ${r.id}", "${sharePointReadablePath(r)}")
    text = add_filters(text, records)
    text = text.replace('Resultado ponderado</span>', 'Solo requisitos evaluados</span>')
    text = text.replace(f'{s["napp"]} requisitos aplicables', f'{s["napp"]} requisitos evaluados')
    text = text.replace(f'{s["napp"]} aplicables · {s["na"]} N/A',
                        f'{s["napp"]} evaluados · {s["pending"]} pendientes · {s["na"]} N/A')
    return migration_ui(text, records)


def group_cards(records):
    cards = []
    for code, rows in shared.groups_of(records).items():
        s = stats(rows)
        percent = shared.fmt(s["pct"])
        cards.append(
            f'<div class="group-card" onclick="setGroup(\'{code}\')"><div class="group-top"><div>'
            f'<span class="group-code">{code}</span><h3>{html.escape(shared.group_name(rows[0]))}</h3></div>'
            f'<strong>{percent + "%" if s["pct"] is not None else "N/A"}</strong></div>'
            f'<div class="group-meta"><span>{s["n"]} requisitos</span>'
            f'<span>{s["gap"]} con brecha · {s["pending"]} pendientes · {s["na"]} N/A</span></div>'
            f'<div class="bar"><span style="width:{s["pct"] or 0}%"></span></div></div>')
    return "".join(cards)


def add_filters(text, records):
    text = select_options(text, "statusFilter", "Todos los estados", dict.fromkeys(r["status"] for r in records))
    return select_options(text, "priorityFilter", "Todas las prioridades", dict.fromkeys(r["priority"] for r in records))


def auditor_page(records, date):
    text = template("auditor", date)
    text = re.sub(r'const SP_LIBRARY_URL =.*?(?=const EVIDENCE_STORAGE_KEY)', lambda m: link_code(), text, flags=re.S)
    text = text.replace("resso424_auditor_evidence_links_v1", "resso424_v11_auditor_evidence_links_v1")
    text = re.sub(r'<div class="evidence-upload">.*?(?=<p id="evidenceMsg")',
                  '<p class="evidence-msg">Visor local y subida pendientes de configurar para v11. '
                  'Puedes agregar varios enlaces de SharePoint y usar «ábrelo en SharePoint» si bloquea el visor.</p>\n',
                  text, count=1, flags=re.S)
    text = re.sub(r'const LOCAL_VIEWER =.*?(?=async function openEvidenceFile)',
                  '', text, count=1, flags=re.S)
    text = re.sub(r'  const rel = localRelativePath\(file\);.*?(?=  el\("evidenceViewerFrame"\).src = evidenceEmbedUrl)',
                  '', text, count=1, flags=re.S)
    text = re.sub(r'  msg.textContent = rel.*?msg.className = "evidence-msg";',
                  '  msg.textContent = "Si SharePoint bloquea el visor, usa el enlace para abrir el archivo.";\n'
                  '  msg.className = "evidence-msg";', text, count=1, flags=re.S)
    text = re.sub(r'async function uploadEvidenceFile\(\).*?(?=function saveEvidenceSelection)',
                  '', text, count=1, flags=re.S)
    return migration_ui(add_filters(text, records), records, auditor=True)


def finalize(name, records, date):
    path = DEST / f"{name}.html"
    text = path.read_text(encoding="utf-8")
    s = stats(records)
    text = text.replace("Requisitos aplicables", "Requisitos evaluados")
    note = (f'{s["napp"]} evaluados · {s["pending"]} pendientes · {s["na"]} no aplican. '
            'Los pendientes y N/A se excluyen del porcentaje. Evaluación heredada de v10; '
            f'preguntas nuevas pendientes de evaluar. Actualización: {date}.')
    text = text.replace('<main class="page">', f'<main class="page"><p class="note">{note}</p>', 1)
    if name == "index":
        cuts, acc = [], 0
        for key in ("ok", "menor", "rel", "critical", "na", "pending"):
            start = acc
            acc += s[key] / s["n"] * 100
            color = dict(ok="#22a96b", menor="#e6a022", rel="#e34f5f", critical="#9f1239",
                         na="#b9c2cf", pending="#7185a5")[key]
            cuts.append(f'{color} {start}% {acc}%')
        text = re.sub(r'(\.donut\{.*?background:).*?(;display:grid)',
                      lambda m: m[1] + "conic-gradient(" + ",".join(cuts) + ")" + m[2], text, count=1, flags=re.S)
        text = text.replace('<div class="legend">', '<div class="legend">'
                            f'<div class="legend-row"><span>Pendiente</span><strong>{s["pending"]}</strong></div>'
                            f'<div class="legend-row"><span>Crítico</span><strong>{s["critical"]}</strong></div>', 1)
    path.write_text(text, encoding="utf-8")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook")
    parser.add_argument("--date", required=True)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    datetime.strptime(args.date, "%d-%m-%Y")
    records = load(args.workbook)
    total = stats(records)
    phases = {p: stats([r for r in records if r["phase"] == p]) for p, _ in shared.PHASES}
    print("RESSO v11:", total, "| Fases:", phases)
    if args.dry_run:
        return
    DEST.mkdir(exist_ok=True)
    shared.EV = os.fspath(DEST)
    for phase, name in shared.PHASES:
        rows = [r for r in records if r["phase"] == phase]
        shared.write(name + ".html", phase_page(phase, name, rows, args.date))
    shared.write("auditor.html", auditor_page(records, args.date))
    shared.write("index.html", migration_ui(template("index", args.date), records))
    shared.update_auditor(records, args.date, total)
    shared.update_index(records, phases, total, args.date)
    for phase, name in shared.PHASES:
        finalize(name, [r for r in records if r["phase"] == phase], args.date)
    for name in ("index", "auditor"):
        finalize(name, records, args.date)


if __name__ == "__main__":
    main()
