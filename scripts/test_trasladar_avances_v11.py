import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

import openpyxl

import trasladar_avances_v11 as transfer


class TransferTests(unittest.TestCase):
    def test_levels_use_lower_step(self):
        for before, after in ((.6, .5), (.8, .75), (.9, .75), (1, 1), (.25, .25), (0, 0), ("N/A", "N/A")):
            self.assertEqual(transfer.convert_level(before), after)
        with self.assertRaises(ValueError):
            transfer.convert_level(None)

    def test_conservative_mapping_and_excel_preservation(self):
        source = openpyxl.Workbook()
        target = openpyxl.Workbook()
        changes = target.create_sheet("Cambios v10 → v11")
        for n, kind in enumerate(("Sin cambio", "Modificada", "Nueva"), 1):
            for workbook, level in ((source, .9), (target, 1 if n == 3 else .5)):
                sheet = workbook.worksheets[0]
                sheet.cell(n + 29, 1, n)
                sheet.cell(n + 29, 6, level)
                sheet.cell(n + 29, 7, f"=F{n + 29}*0.1")
                for col, value in ((11, "Criterio"), (12, "Anterior"), (13, "Actual"),
                                   (14, "Acción"), (15, "Responsable"), (16, "Enlace")):
                    sheet.cell(n + 29, col, value)
            changes.append([n, "A1", "Pregunta", .1, n, "Anterior", .1, kind, "Detalle"])
        target.worksheets[0]["M29"] = "Observación anterior"
        updates = transfer.prepare_updates(source, target, "26-10-2026")
        self.assertEqual(updates["F30"], .75)
        self.assertNotIn("F31", updates)
        self.assertNotIn("F32", updates)
        self.assertIn("no validación", updates["M31"])
        with tempfile.TemporaryDirectory() as folder:
            original, output = Path(folder) / "original.xlsx", Path(folder) / "output.xlsx"
            target.save(original)
            transfer.write_copy(original, output, updates)
            result = openpyxl.load_workbook(output)
            self.assertEqual([result.worksheets[0].cell(i, 6).value for i in (30, 31, 32)], [.75, .5, 1])
            self.assertEqual(result.worksheets[0]["G30"].value, "=F30*0.1")
            self.assertEqual(result.worksheets[0]["P30"].value, "Enlace")
            with ZipFile(original) as before, ZipFile(output) as after:
                for name in before.namelist():
                    if not name.startswith("xl/worksheets/") and name != "xl/workbook.xml":
                        self.assertEqual(before.read(name), after.read(name), name)


if __name__ == "__main__":
    unittest.main()
