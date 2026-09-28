import csv, re, json, datetime
import openpyxl

OLD = "Consolidado Incidentes y Medidas Correctivas (Formato Interno) 05.08.2026.xlsx"
NEW = "Consolidado Incidentes y Medidas Correctivas (Formato Interno) 05.08.csv"


def key(v):
    v = str(v or "").strip().lower()
    for a, b in zip("áéíóúñ", "aeioun"):
        v = v.replace(a, b)
    return re.sub(r"[^a-z0-9]", "", v)


def norm_txt(v):
    v = str(v or "")
    v = v.replace("\ufffd", "?")
    return re.sub(r"\s+", " ", v).strip().upper()


def build(rows):
    hidx = None
    for i, r in enumerate(rows):
        ks = [key(c) for c in r]
        if "item" in ks and "nombredelevento" in ks:
            hidx = i
            break
    headers = [key(c) for c in rows[hidx]]

    def col(*names):
        for n in names:
            if n in headers:
                return headers.index(n)
        return -1

    C = {
        "item": col("item"), "fecha": col("fecha"), "nombre": col("nombredelevento"),
        "categoria": col("categoria"),
        "cierre": col("fechacierreaccion", "fechacierreaccin", "fechacierre"),
        "medida": col("medidacorrectiva"), "estatus": col("estatus"),
    }

    def val(r, c):
        return str(r[c]).strip() if 0 <= c < len(r) and r[c] is not None else ""

    out, cur = [], None
    for r in rows[hidx + 1:]:
        item = val(r, C["item"])
        if item:
            cur = {"item": item, "nombre": norm_txt(val(r, C["nombre"])),
                   "fecha": val(r, C["fecha"]), "categoria": norm_txt(val(r, C["categoria"])),
                   "medidas": []}
            out.append(cur)
        if cur is None:
            continue
        m = val(r, C["medida"])
        if m:
            cur["medidas"].append({"cierre": val(r, C["cierre"]), "estatus": norm_txt(val(r, C["estatus"]))})
    return out


def fmt_date(v):
    v = str(v).strip()
    if re.match(r"^\d{4}-\d{2}-\d{2}", v):
        y, m, d = v[:10].split("-")
        return f"{d}-{m}-{y}"
    m = re.match(r"^(\d{1,2})[-/.](\d{1,2})[-/.](\d{2,4})", v)
    if m:
        y = m.group(3)
        y = "20" + y if len(y) == 2 else y
        return f"{int(m.group(1)):02d}-{int(m.group(2)):02d}-{y}"
    return v


wb = openpyxl.load_workbook(OLD, data_only=True)
old_rows = [[c if c is not None else "" for c in row] for row in wb[wb.sheetnames[0]].iter_rows(values_only=True)]
old = build(old_rows)

with open(NEW, "r", encoding="latin-1", newline="") as fh:
    new_rows = list(csv.reader(fh, delimiter=";"))
new = build(new_rows)

print(f"ANTERIOR: {len(old)} incidentes | NUEVO: {len(new)} incidentes")
oi = {i["item"]: i for i in old}
ni = {i["item"]: i for i in new}
print("ITEMS NUEVOS:", sorted(set(ni) - set(oi), key=lambda x: int(re.match(r'\d+', x).group())))
print("ITEMS ELIMINADOS:", sorted(set(oi) - set(ni)))

print("\n--- ITEMS QUE CAMBIARON DE EVENTO (riesgo para links) ---")
for k in sorted(set(oi) & set(ni), key=lambda x: int(re.match(r"\d+", x).group())):
    a, b = oi[k], ni[k]
    if a["nombre"] != b["nombre"]:
        print(f"ITEM {k}\n   antes: {a['nombre'][:80]}\n   ahora: {b['nombre'][:80]}")

print("\n--- ITEMS QUE CAMBIARON DE FECHA ---")
for k in sorted(set(oi) & set(ni), key=lambda x: int(re.match(r"\d+", x).group())):
    a, b = fmt_date(oi[k]["fecha"]), fmt_date(ni[k]["fecha"])
    if a != b:
        print(f"ITEM {k}: {a}  ->  {b}   ({ni[k]['nombre'][:55]})")

print("\n--- CAMBIOS EN N° DE MEDIDAS ---")
for k in sorted(set(oi) & set(ni), key=lambda x: int(re.match(r"\d+", x).group())):
    a, b = len(oi[k]["medidas"]), len(ni[k]["medidas"])
    if a != b:
        print(f"ITEM {k}: {a} -> {b}   ({ni[k]['nombre'][:55]})")
