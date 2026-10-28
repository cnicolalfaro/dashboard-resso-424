import tempfile
import subprocess
import unittest
from pathlib import Path

import openpyxl

import actualizar_evidencias_v11 as v11


class V11Tests(unittest.TestCase):
    def test_numbered_folder_paths(self):
        code = v11.link_code() + r"""
for (const id of [1, 9, 10, 40]) {
  const r = {id};
  const folder = String(id).padStart(2, "0") + " versi\u00f3n 11";
  const actual = new URL(specificSharePointLink(r)).searchParams.get("id");
  if (actual !== SP_ROOT_FOLDER + "/" + folder) throw Error(actual);
  if (sharePointReadablePath(r) !== folder) throw Error(sharePointReadablePath(r));
}
const explicit = "https://empresassk.sharepoint.com/explicit";
if (specificSharePointLink({id: 36, link: explicit}) !== explicit) throw Error("Link changed");
"""
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "paths.js"
            path.write_text(code, encoding="utf-8")
            subprocess.run(["node", str(path)], check=True, capture_output=True)

    def test_formula_results_and_pending_are_calculated_from_levels(self):
        with tempfile.TemporaryDirectory() as folder:
            path = Path(folder) / "v11.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet["A29"] = "N°"
            for n, level in enumerate((1, .75, .5, None, "N/A", 0), start=1):
                row = n + 29
                values = [n, "A. PLAN", "A1. Programa", f"Pregunta {n}", .04, level,
                          f"=E{row}*F{row}", None, None, None, "Criterio",
                          "Observación anterior", "Observación nueva", "Acción", "Responsable"]
                for col, value in enumerate(values, start=1):
                    sheet.cell(row, col, value)
            sheet["P30"] = "Abrir SharePoint"
            sheet["P30"].hyperlink = "https://empresassk.sharepoint.com/evidencia"
            changes = workbook.create_sheet("Cambios v10 → v11")
            for n in range(1, 7):
                changes.append([n, "A1", f"Pregunta {n}", .04, n, "Anterior", .04,
                                "Nueva" if n == 4 else "Sin cambio", "Detalle"])
            workbook.save(path)
            rows = v11.load(path)
        self.assertEqual([r["status"] for r in rows],
                         ["Cumple", "Brecha menor", "Brecha relevante", "Pendiente", "No aplica", "Crítico"])
        self.assertEqual([r["priority"] for r in rows], ["Baja", "Alta", "Alta", "—", "—", "Alta"])
        self.assertEqual(rows[1]["result"], .03)
        self.assertEqual(rows[1]["gap"], .01)
        self.assertIsNone(rows[3]["level"])
        self.assertIsNone(rows[3]["result"])
        self.assertEqual(rows[0]["observation"], "Observación nueva")
        self.assertEqual(rows[0]["evidence_ref"], "https://empresassk.sharepoint.com/evidencia")
        s = v11.stats(rows)
        self.assertEqual((s["n"], s["napp"], s["pending"], s["na"], s["gap"]), (6, 4, 1, 1, 3))
        self.assertEqual(s["pct"], 56.25)

    def test_change_colors_follow_agreed_rules(self):
        cases = [
            ("Nueva", None, "new"), ("Nueva", 1, "new"),
            ("Modificada", 1, "review"), ("Modificada / reubicada", .5, "review"),
            ("Sin cambio", 1, "complete"), ("Redacción", 1, "complete"),
            ("Ponderación", 1, "complete"), ("Sin cambio", .75, "review"),
            ("Sin cambio", None, "review"), ("Modificada", "N/A", "na"),
        ]
        for change_type, level, expected in cases:
            with self.subTest(change_type=change_type, level=level):
                self.assertEqual(v11.classify_change(change_type, level)["migration"], expected)
        with self.assertRaises(ValueError):
            v11.classify_change("Desconocida", 1)

    def test_v11_template_isolates_storage_and_local_service(self):
        text = v11.auditor_page([], "22-10-2026")
        self.assertIn('"resso424_v11_auditor_evidence_links_v1"', text)
        self.assertNotIn('"resso424_auditor_evidence_links_v1"', text)
        self.assertNotIn("127.0.0.1:8765", text)
        self.assertNotIn("/upload?", text)
        self.assertNotIn("PHASE_FOLDER", text)
        self.assertIn('id="evidenceMsg"', text)
        self.assertIn("05 RESSO V11", text)
        self.assertIn('${String(r.id).padStart(2, "0")} versión 11', text)
        self.assertIn("RessoObservations.open(r.id)", text)
        self.assertIn('src="../../assets/resso-observations.js"', text)

    def test_phase_generator_preserves_shared_history(self):
        record = dict(id=1, phase="A. PLAN", section="A1. Programa",
                      question="Pregunta", weight=.1, level=1, result=.1, gap=0,
                      status="Cumple", priority="Baja", criterion="", observation="",
                      action="", responsible="", evidence_ref=None,
                      migration="complete", migration_label="Equivalente",
                      change_type="Sin cambio", change_detail="Igual")
        for phase, name in v11.shared.PHASES:
            with self.subTest(phase=phase):
                text = v11.phase_page(phase, name, [record], "28-10-2026")
                self.assertIn("RessoObservations.open(r.id)", text)
                self.assertIn('src="../../assets/resso-observations.js"', text)


if __name__ == "__main__":
    unittest.main()
