"""Replica la logica de incidentStatus() de incidentes-v2.js sobre el xlsx antes/despues."""
import re, subprocess, unicodedata
import openpyxl

BASE = "Consolidado Incidentes y Medidas Correctivas (Formato Interno) 05.08.2026 (1).xlsx"
LATEST = "Consolidado Incidentes y Medidas Correctivas (Formato Interno) 05.08.2026 hoy 05-10-26.xlsx"


def norm(v):
    s = str(v if v is not None else "")
    s = re.sub(r"[\u200B-\u200D\u2060\uFEFF]", "", s).strip().lower()
    return "".join(c for c in unicodedata.normalize("NFD", s) if unicodedata.category(c) != "Mn")


def key(v):
    return re.sub(r"[^a-z0-9]", "", norm(v))


def build(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    hidx = next(i for i, r in enumerate(rows)
                if "item" in [key(c) for c in r] and "nombredelevento" in [key(c) for c in r])
    out, cur = [], None
    for r in rows[hidx + 1:]:
        item = str(r[0]).strip() if r[0] not in (None, "") else ""
        if item:
            cur = {"item": item.split(".")[0], "nombre": re.sub(r"\s+", " ", str(r[3] or "")).strip(), "medidas": []}
            out.append(cur)
        if cur is None:
            continue
        if r[6] not in (None, "") and str(r[6]).strip():
            cur["medidas"].append({"medida": norm(r[6]), "resp": norm(r[7]), "estatus": norm(r[9])})
    return {i["item"]: i for i in out}


def status(inc):
    if any("sin informe final" in " ".join([m["medida"], m["resp"]]) for m in inc["medidas"]):
        return "pendiente"
    if not inc["medidas"]:
        return "pendiente"
    for m in inc["medidas"]:
        s = m["estatus"]
        if s in ("n/a", "na") or "no aplica" in s:
            return "pendiente"
        if "abiert" in s or "pendient" in s or "proceso" in s:
            return "pendiente"
    return "cerrado" if all("cerrad" in m["estatus"] for m in inc["medidas"]) else "sin estado"


subprocess.run(["git", "stash", "list"], capture_output=True)
subprocess.run(["git", "show", f"HEAD:{BASE}"], stdout=open("_prev.xlsx", "wb"), check=True)

prev, new = build("_prev.xlsx"), build(LATEST)
sp = {k: status(v) for k, v in prev.items()}
sn = {k: status(v) for k, v in new.items()}

order = lambda s: sorted(s, key=lambda x: int(re.match(r"\d+", x).group()))
print("ANTES  ->", {s: list(sp.values()).count(s) for s in set(sp.values())}, f"({len(sp)} incidentes)")
print("AHORA  ->", {s: list(sn.values()).count(s) for s in set(sn.values())}, f"({len(sn)} incidentes)")
print("\nINCIDENTES QUE CAMBIARON DE ESTADO GLOBAL:")
for k in order(set(sp) & set(sn)):
    if sp[k] != sn[k]:
        print(f"  ITEM {k}: {sp[k]} -> {sn[k]}   ({new[k]['nombre'][:55]})")
for k in order(set(sn) - set(sp)):
    print(f"  ITEM {k}: (nuevo) -> {sn[k]}   ({new[k]['nombre'][:55]})")

print("\nMEDIDAS QUE CAMBIARON DE ESTATUS:")
for k in order(set(sp) & set(sn)):
    a, b = prev[k]["medidas"], new[k]["medidas"]
    if len(a) != len(b):
        print(f"  ITEM {k}: cantidad de medidas {len(a)} -> {len(b)}")
        continue
    for n, (ma, mb) in enumerate(zip(a, b), 1):
        if ma["estatus"] != mb["estatus"]:
            print(f"  ITEM {k} medida #{n}: {ma['estatus'].upper()!r} -> {mb['estatus'].upper()!r}")
