import re
import unicodedata
from pathlib import Path

import openpyxl

JS = Path("incidentes-v2.js").read_text(encoding="utf-8")
SOURCE = "Consolidado Incidentes y Medidas Correctivas (Formato Interno) 05.08.2026 (1).xlsx"
QUERY = "query_incidentes.xlsx"


def norm(value):
    text = re.sub(r"[\u200b-\u200d\u2060\ufeff]", "", str(value or "")).strip().lower()
    text = unicodedata.normalize("NFD", text)
    return "".join(char for char in text if not unicodedata.combining(char))


def tokens(value):
    ignored = {"dano", "material", "incidente", "medida", "medidas", "correctivas", "evento",
               "nivel", "para", "con", "del", "los", "las", "una", "uno", "entre", "sector",
               "skic", "camioneta"}
    return [word for word in re.split(r"[^a-z0-9]+", norm(value)) if len(word) >= 3 and word not in ignored]


def map_block(name):
    match = re.search(rf"const {name} = \{{(.*?)\n  \}};", JS, re.S)
    return match.group(1) if match else ""


static_block = map_block("EVIDENCE_STATIC_LINKS")
static = {key: (label, path) for key, label, path in
          re.findall(r'"(\d+)": \["([^"]*)", "([^"]*)"\]', static_block)}
override_block = map_block("EVIDENCE_OVERRIDES")
overrides = dict(re.findall(r'"(\d+)": "([^"]*)"', override_block))

book = openpyxl.load_workbook(SOURCE, data_only=True, read_only=True)
sheet = book.worksheets[0]
rows = list(sheet.iter_rows(values_only=True))
header_index = next(i for i, row in enumerate(rows)
                    if "item" in [norm(value).replace(" ", "") for value in row]
                    and "nombredelevento" in [norm(value).replace(" ", "") for value in row])
headers = [norm(value).replace(" ", "") for value in rows[header_index]]
item_col, name_col, date_col, category_col = (headers.index(column) for column in
                                               ("item", "nombredelevento", "fecha", "categoria"))
incidents, current = [], None
for row in rows[header_index + 1:]:
    item = str(row[item_col] or "").strip()
    if item:
        match = re.match(r"\d+", item)
        current = {"item": match.group() if match else item,
                   "name": str(row[name_col] or "").strip(),
                   "date": str(row[date_col] or "").strip(),
                   "category": str(row[category_col] or "").strip()}
        incidents.append(current)

query = openpyxl.load_workbook(QUERY, data_only=True, read_only=True)
query_sheet = query.worksheets[0]
query_rows = list(query_sheet.iter_rows(values_only=True))
query_headers = [str(value or "").strip() for value in query_rows[0]]
name_idx = query_headers.index("Nombre")
path_idx = query_headers.index("Ruta de acceso")
type_idx = query_headers.index("Tipo de elemento")
folders = []
for row in query_rows[1:]:
    if "carpeta" not in norm(row[type_idx]):
        continue
    name = str(row[name_idx] or "").strip()
    parent = str(row[path_idx] or "").strip()
    path = "/".join(part for part in (parent, name) if part)
    folder_tokens = set(tokens(f"{name} {path}"))
    if name and path and tokens(name):
        folders.append({"name": name, "path": path, "tokens": folder_tokens})

by_category = {}
for incident in incidents:
    by_category[incident["category"]] = by_category.get(incident["category"], 0) + 1

linked, missing, item67_matches = [], [], []
for incident in incidents:
    item, name = incident["item"], incident["name"]
    target = static.get(item)
    found = None
    reason = "static"
    if target:
        found = next((folder for folder in folders if norm(folder["path"]) == norm(target[1])), None)
        if not found:
            found = next((folder for folder in folders if norm(folder["name"]) == norm(target[0])), None)
        if not found:
            found = next((folder for folder in folders if norm(target[0]) in norm(folder["name"])
                          or norm(folder["name"]) in norm(target[0])), None)
        linked.append((item, "static exact" if found else "static path not found"))
    else:
        reason = "automatic"
        override = overrides.get(item)
        if override:
            found = next((folder for folder in folders if norm(folder["name"]) == norm(override)), None)
            if not found:
                found = next((folder for folder in folders if norm(override) in norm(folder["name"])
                              or norm(folder["name"]) in norm(override)), None)
        if not found:
            words = tokens(name)
            scored = []
            for folder in folders:
                hits = [word for word in words if word in folder["tokens"]]
                score = len(hits) / len(words) if words else 0
                if ((score >= 0.6 and len(hits) >= min(2, len(words)))):
                    scored.append((score, folder))
            if scored:
                found = sorted(scored, key=lambda pair: (-pair[0], len(pair[1]["path"])))[0][1]
        if found:
            linked.append((item, reason))
        else:
            missing.append((item, name))
    if item == "67":
        item67_matches.append((target, found, incident))

print(f"INCIDENTES {len(incidents)} | enlaces estáticos {len(static)} | carpetas indexadas {len(folders)}")
for target, folder, incident in item67_matches:
    print("ITEM 67:", incident["name"])
    print("  ESTÁTICO:", target)
    print("  COINCIDENCIA EN INVENTARIO:", folder["path"] if folder else "NO ENCONTRADA")
    similar = [folder for folder in folders if any(word in norm(folder["name"]) for word in ("tensor", "soporte"))]
    for candidate in similar:
        print("  CARPETA RELACIONADA:", candidate["name"], "|", candidate["path"])
print("SIN ENLACE ASIGNADO POR MAPA NI INVENTARIO:", len(missing))
for item, name in missing:
    incident = next(incident for incident in incidents if incident["item"] == item)
    print(f"  {item}: {name} | fecha {incident['date']} | categoría {incident['category']}")
    words = tokens(name)
    candidates = sorted(((sum(word in folder["tokens"] for word in words), folder)
                         for folder in folders), key=lambda pair: (-pair[0], len(pair[1]["path"])))
    for hits, folder in [candidate for candidate in candidates if candidate[0]][:5]:
        print(f"    POSIBLE CARPETA ({hits} coincidencias): {folder['name']} | {folder['path']}")
print("ESTÁTICOS SIN RUTA/PATH QUE COINCIDA EN INVENTARIO:")
for item, state in linked:
    if state == "static path not found":
        print(f"  {item}: {static[item][0]}")