import copy
import datetime
import io
import os
import re
import shutil
import subprocess
import threading
import uuid
from pathlib import Path

import pymupdf as fitz
import qrcode
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

from .documents import qr_payload


WORD_VERSION = 3
CONVERT_SECONDS = 120
MARKER = 10
QR_SIDES = (112, 96)
COMPACT_ROW = 24.5
BODY_ELEMENTS = 9
EXPECTED_LABELS = ((0, 0, "NOMBRE RELATOR"), (1, 0, "FECHA"), (3, 3, "OBJETIVOS"), (4, 2, "MATERIA"),
                   (8, 0, "PARTICIPANTES"))
BLUE = (0, 0.459, 0.741)
_converter = threading.BoundedSemaphore(1)


class RenderError(ValueError):
    pass


class NoRoom(RenderError):
    pass


def soffice_path():
    found = os.environ.get("NAMI_SOFFICE") or shutil.which("soffice") or shutil.which("libreoffice")
    if not found:
        raise RenderError("LibreOffice no está instalado")
    return found


def convert_to_pdf(sources, outdir, workdir):
    command = [soffice_path(), "--headless", "--norestore", "--nolockcheck", "--nodefault", "--nologo",
               f"-env:UserInstallation={(Path(workdir) / 'profile').as_uri()}",
               "--convert-to", "pdf:writer_pdf_Export", "--outdir", str(outdir), *map(str, sources)]
    environment = {"PATH": os.environ.get("PATH", ""), "HOME": str(workdir), "LANG": "C.UTF-8"}
    with _converter:
        try:
            subprocess.run(command, capture_output=True, timeout=CONVERT_SECONDS, check=False,
                           env=environment, cwd=workdir)
        except subprocess.TimeoutExpired as error:
            raise RenderError("La conversión a PDF excedió el tiempo permitido") from error
    results = [Path(outdir) / f"{Path(source).stem}.pdf" for source in sources]
    if any(not result.is_file() for result in results):
        raise RenderError("No fue posible convertir el documento a PDF")
    return results


def set_text(cell, text, size=None):
    for paragraph in cell.paragraphs[1:]:
        paragraph._element.getparent().remove(paragraph._element)
    paragraph = cell.paragraphs[0]
    runs = paragraph.runs
    if not runs:
        runs = [paragraph.add_run(text)]
    else:
        runs[0].text = text
        for extra in runs[1:]:
            extra.text = ""
    for run in paragraph.runs:
        if size is not None:
            run.font.size = Pt(size)


def compact(paragraph):
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.line_spacing = 1.0


def check_template(info):
    for row, column, label in EXPECTED_LABELS:
        if not info.cell(row, column).text.strip().upper().startswith(label):
            raise RenderError("La plantilla del registro no tiene la estructura esperada")
    if len(info.cell(4, 2).paragraphs) < 2:
        raise RenderError("La plantilla del registro no tiene la estructura esperada")


def fill_info(info, header):
    check_template(info)
    if header.get("relator"):
        set_text(info.cell(0, 6), header["relator"])
    set_text(info.cell(1, 0), f"FECHA: {header['date']}")
    objectives = info.cell(3, 3)
    set_text(objectives, f"OBJETIVOS: {header['objectives']}")
    if header.get("foreman"):
        cell = info.cell(8, 4).merge(info.cell(8, 13))
        set_text(cell, f"{header.get('foreman_label', 'CAPATAZ')}: {header['foreman']}",
                 size=header.get("foreman_size", 10))
        for run in cell.paragraphs[0].runs:
            run.bold = True
    materials = info.cell(4, 2)
    model = materials.paragraphs[-1]
    for paragraph in materials.paragraphs[1:-1]:
        paragraph._element.getparent().remove(paragraph._element)
    for number, title in enumerate(header["subjects"], 1):
        paragraph = model if number == 1 else Paragraph(copy.deepcopy(model._p), model._parent)
        if number > 1:
            materials._tc.append(paragraph._p)
        for extra in list(paragraph.runs)[1:]:
            extra._element.getparent().remove(extra._element)
        if paragraph.runs:
            paragraph.runs[0].text = f"{number}. {title}"
        else:
            paragraph.add_run(f"{number}. {title}")


def fill_people(table, people, first_number):
    rows = table.rows
    for row in list(rows)[len(people) + 1:]:
        row._tr.getparent().remove(row._tr)
    for index, person in enumerate(people):
        cells = table.rows[index + 1].cells
        set_text(cells[0], str(first_number + index))
        name = cells[1].paragraphs[0]
        if name.runs:
            name.runs[0].text = person["nombre"]
            for extra in name.runs[1:]:
                extra.text = ""
        else:
            name.add_run(person["nombre"]).font.size = Pt(8)
        label = person.get("crew_label")
        if (person.get("rol") or label) and len(cells[1].paragraphs) > 1:
            role = cells[1].paragraphs[1]
            text = f"ROL:[ {person['rol']} ]" if person.get("rol") else "ROL:[        ]"
            if label:
                text += f"    {label}"
            if role.runs:
                role.runs[0].text = text
                for extra in role.runs[1:]:
                    extra.text = ""
        # The RUT cell stays blank: each worker writes it by hand when signing. The RUT is kept only in the manifest.
        set_text(cells[2], "", size=8)
        for position, key in ((3, "cargo"), (4, "especialidad")):
            value = person.get(key, "")
            set_text(cells[position], "" if str(value).strip() == "." else value, size=8)
        set_text(cells[5], "")


def build_docx(template, people_pages, header, output, row_height=None):
    document = Document(str(template))
    body = document.element.body
    children = list(body)
    if len(children) != BODY_ELEMENTS + 1:
        raise RenderError("La plantilla del registro no tiene la estructura esperada")
    block, section = children[:-1], children[-1]
    for element in block:
        body.remove(element)
    first_number = 1
    for page_number, people in enumerate(people_pages):
        copies = [copy.deepcopy(element) for element in block]
        for element in copies:
            section.addprevious(element)
        info, table = Table(copies[6], document), Table(copies[7], document)
        fill_info(info, header)
        fill_people(table, people, first_number)
        if row_height:
            for row in table.rows[1:]:
                row.height = Pt(row_height)
                row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        first_number += len(people)
        if page_number:
            Paragraph(copies[0], document).paragraph_format.page_break_before = True
        if page_number < len(people_pages) - 1:
            section_parent = copies[8].getparent()
            section_parent.remove(copies[8])
    document.save(str(output))


def locate_rows(page, count):
    for table in page.find_tables().tables:
        header = table.extract()[0] if table.row_count else []
        labels = [" ".join(str(value or "").split()).upper() for value in header]
        if "FIRMA" not in labels or "RUT" not in labels or table.row_count != count + 1:
            continue
        cells = table.rows[0].cells
        rut_cell, signature_cell = cells[labels.index("RUT")], cells[labels.index("FIRMA")]
        if rut_cell is None or signature_cell is None:
            continue
        x0, _, x1, _ = table.bbox
        rows = [{"box": [x0, row.bbox[1], x1 - x0, row.bbox[3] - row.bbox[1]]} for row in table.rows[1:]]
        return rows, rut_cell[0], signature_cell[0]
    raise RenderError("No se pudo ubicar la tabla de participantes en el PDF")


def occupied_areas(page):
    areas = [fitz.Rect(block[:4]) for block in page.get_text("blocks")]
    areas += [fitz.Rect(drawing["rect"]) for drawing in page.get_drawings()]
    areas += [fitz.Rect(image["bbox"]) for image in page.get_image_info()]
    return [area + (-3, -3, 3, 3) for area in areas if area.width > 0 or area.height > 0]


def free(rect, areas, page_rect):
    margin = fitz.Rect(page_rect) + (8, 8, -8, -8)
    return rect in margin and not any(rect.intersects(area) for area in areas)


def place_markers(page, areas):
    width, height = page.rect.width, page.rect.height
    markers = []
    for sx, sy in ((1, 1), (-1, 1), (-1, -1), (1, -1)):
        for distance in range(15, 90, 3):
            cx = distance if sx > 0 else width - distance
            cy = distance if sy > 0 else height - distance
            rect = fitz.Rect(cx - MARKER / 2, cy - MARKER / 2, cx + MARKER / 2, cy + MARKER / 2)
            if free(rect, areas, page.rect):
                markers.append((cx, cy))
                areas.append(rect + (-3, -3, 3, 3))
                break
        else:
            raise NoRoom("No hay espacio libre para las marcas de registro")
    return markers


def find_spot(page, areas, width, height):
    y = page.rect.height - height - 12
    while y > 12:
        x = page.rect.width - width - 12
        while x > 12:
            rect = fitz.Rect(x, y, x + width, y + height)
            if free(rect, areas, page.rect):
                return rect
            x -= 6
        y -= 6
    return None


def stamp_page(page, payload, number):
    areas = occupied_areas(page)
    markers = place_markers(page, areas)
    qr = qrcode.QRCode(box_size=8, border=4)
    qr.add_data(payload)
    qr.make(fit=True)
    image = qr.make_image(fill_color="black", back_color="white").convert("RGB")
    stream = io.BytesIO()
    image.save(stream, format="PNG")
    for side in QR_SIDES:
        spot = find_spot(page, areas, side, side)
        if spot:
            break
    else:
        raise NoRoom("No hay espacio libre para el código QR")
    areas.append(spot + (-3, -3, 3, 3))
    patch = find_spot(page, areas, 14, 14)
    if patch is None:
        raise NoRoom("No hay espacio libre para el control de color")
    for cx, cy in markers:
        page.draw_rect(fitz.Rect(cx - MARKER / 2, cy - MARKER / 2, cx + MARKER / 2, cy + MARKER / 2),
                       color=None, fill=(0, 0, 0))
    page.insert_image(spot, stream=stream.getvalue())
    page.draw_rect(patch, color=None, fill=BLUE)
    border = spot.width * 4 / (qr.modules_count + 8)
    corners = [[spot.x0 + border, spot.y0 + border], [spot.x1 - border, spot.y0 + border],
               [spot.x1 - border, spot.y1 - border], [spot.x0 + border, spot.y1 - border]]
    return {"qr": payload, "qr_corners": corners, "markers": [list(point) for point in markers],
            "color_box": [patch.x0, patch.y0, patch.width, patch.height]}


def split_pages(people, capacity):
    return [people[index:index + capacity] for index in range(0, len(people), capacity)]


def rows_on_first_page(document):
    tables = document[0].find_tables().tables
    return max((table.row_count - 1 for table in tables if table.row_count > 1), default=0)


def make_word_document(template, crew, people, courses, date, secret, workdir, objectives, relator=""):
    if not people:
        raise RenderError("No hay participantes pendientes en la cuadrilla seleccionada")
    template = Path(template)
    if not template.is_file():
        raise RenderError("Falta la plantilla privada del registro")
    work = Path(workdir) / uuid.uuid4().hex
    work.mkdir(mode=0o700)
    try:
        capacity = len(Document(str(template)).tables[1].rows) - 1
        header = {"date": datetime.date.fromisoformat(date).strftime("%d-%m-%Y"),
                  "objectives": objectives, "foreman": crew.get("foreman", ""), "relator": relator,
                  "foreman_label": "CAPATACES" if crew.get("group") else "CAPATAZ",
                  "foreman_size": 9 if crew.get("group") else 10,
                  "subjects": [course["title"] for course in courses]}
        compact = False
        while capacity >= 1:
            pages = split_pages(people, capacity)
            source = work / "registro.docx"
            build_docx(template, pages, header, source, COMPACT_ROW if compact else None)
            pdf_path = convert_to_pdf([source], work, work)[0]
            next_capacity, next_compact = capacity - 1, False
            with fitz.open(pdf_path) as document:
                if len(document) == len(pages):
                    try:
                        return finish(document, crew, pages, courses, date, secret, relator)
                    except NoRoom:
                        if not compact:
                            next_capacity, next_compact = capacity, True
                else:
                    # An overflow shows exactly how many rows fit on the first page.
                    next_capacity = min(capacity - 1, rows_on_first_page(document))
            capacity, compact = next_capacity, next_compact
            source.unlink(missing_ok=True)
            pdf_path.unlink(missing_ok=True)
        raise RenderError("El encabezado no deja espacio para participantes")
    finally:
        shutil.rmtree(work, ignore_errors=True)


def finish(document, crew, pages, courses, date, secret, relator=""):
    identifier = str(uuid.uuid4())
    manifest = {"id": identifier, "version": WORD_VERSION, "date": date, "layout": "word", "relator": relator,
                "crew": {"id": crew["id"], "name": crew["name"], "foreman": crew.get("foreman", ""),
                         "group": crew.get("group"), "cargo": crew.get("cargo")},
                "participants": [dict(person) for page in pages for person in page],
                "courses": courses, "pages": []}
    for number, (page, people) in enumerate(zip(document, pages), 1):
        located, rut_x, signature_x = locate_rows(page, len(people))
        stamp = stamp_page(page, qr_payload(identifier, WORD_VERSION, number, len(pages), secret), number)
        rows = []
        for index, (person, geometry) in enumerate(zip(people, located)):
            rows.append({"id": f"{identifier}:{number}:{index}", "rut": person["rut"], "nombre": person["nombre"],
                         "box": geometry["box"], "rut_x": rut_x + 6, "signature_x": signature_x})
        manifest["pages"].append({"page": number, "rows": rows, "markers": stamp["markers"],
                                  "qr": stamp["qr"], "qr_corners": stamp["qr_corners"],
                                  "color_box": stamp["color_box"],
                                  "size": [page.rect.width, page.rect.height]})
    document.set_metadata({"title": "Registro de difusión SGR-R-006.DCH"})
    return manifest, document.tobytes(garbage=3, deflate=True)
